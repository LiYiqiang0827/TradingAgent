"""Build the June-July 2026 box-breakout result and sensitivity deck."""

from __future__ import annotations

from datetime import datetime
import hashlib
import json
from pathlib import Path
import shutil
import subprocess
from types import SimpleNamespace
from typing import Any

import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

import study_aug_sep_2026_box_breakout_ppt as base


OUTPUT = base.ROOT / "outputs" / "jun_jul_2026_box_breakout" / "20260601_20260731_qfq"
EVENT_START = "20260601"
EVENT_END = "20260731"
DATA_END = "20260930"


def assign_sensitivity_band(box_range_pct: float) -> str:
    value = float(box_range_pct)
    if value <= 15.0:
        return "严格主口径"
    if value <= 18.0:
        return "18%边缘候选"
    return "20%宽松候选"


def expected_deck_pages(stock_count: int) -> int:
    return 6 + int(stock_count)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _load_inputs() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    args = SimpleNamespace(
        output_dir=OUTPUT,
        history_start="20240101",
        event_start=EVENT_START,
        event_end=EVENT_END,
        data_end=DATA_END,
    )
    return base.load_or_fetch_inputs(args)


def _prepare_tables() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    edge18_path = OUTPUT / "sensitivity_unique_box18.csv"
    edge20_path = OUTPUT / "sensitivity_unique_box20.csv"
    funnel_path = OUTPUT / "strict_main_funnel_raw.csv"
    for path in (edge18_path, edge20_path, funnel_path):
        if not path.exists():
            raise FileNotFoundError(f"missing frozen full-rescan result: {path}")

    edge18 = pd.read_csv(edge18_path, dtype={"ts_code": str, "event_date": str})
    edge = pd.read_csv(edge20_path, dtype={"ts_code": str, "event_date": str})
    raw_funnel = pd.read_csv(funnel_path)
    if edge["ts_code"].duplicated().any() or edge18["ts_code"].duplicated().any():
        raise ValueError("sensitivity stock tables must be unique by ts_code")
    if not set(edge18["ts_code"]).issubset(set(edge["ts_code"])):
        raise ValueError("18% sensitivity codes must be a subset of the 20% codes")
    funnel_values = dict(zip(raw_funnel["stage"], raw_funnel["unique_stocks"]))
    if int(funnel_values.get("full15", -1)) != 0:
        raise ValueError("strict <=15% scan is no longer empty; rebuild the main-result deck")

    edge18_codes = set(edge18["ts_code"].astype(str))
    edge["sensitivity_band"] = np.where(edge["ts_code"].isin(edge18_codes), "18%边缘候选", "20%宽松候选")
    edge["sensitivity_rank"] = np.where(edge["ts_code"].isin(edge18_codes), 1, 2)
    edge = edge.sort_values(
        ["sensitivity_rank", "similarity_score", "event_date", "ts_code"],
        ascending=[True, False, True, True],
        kind="mergesort",
    ).reset_index(drop=True)
    edge["display_order"] = np.arange(1, len(edge) + 1)
    edge["match_band"] = edge["sensitivity_band"]

    strict = edge.head(0).copy()
    strict.to_csv(OUTPUT / "all_structural_events.csv", index=False, encoding="utf-8-sig")
    strict.to_csv(OUTPUT / "ranked_unique_candidates.csv", index=False, encoding="utf-8-sig")
    edge.to_csv(OUTPUT / "sensitivity_ranked_candidates_box20.csv", index=False, encoding="utf-8-sig")

    labels = {
        "active": "6—7月实际有行情股票",
        "signal": "6—7月涨停或大阳",
        "new_high": "收盘突破此前60日最高价",
        "box15": "再满足紧邻20日以上、振幅≤15%箱体突破",
        "full15": "再满足同口径箱体前放量拉升",
    }
    funnel = raw_funnel.copy()
    funnel["stage_key"] = funnel["stage"]
    funnel["stage"] = funnel["stage_key"].map(labels)
    funnel = funnel[["stage_key", "stage", "unique_stocks", "event_rows"]]
    funnel.to_csv(OUTPUT / "screening_funnel.csv", index=False, encoding="utf-8-sig")

    strict_count = int(funnel_values["full15"])
    sensitivity = pd.DataFrame(
        [
            {"box_range_limit_pct": 12.0, "box_min_trading_days": 20, "unique_stocks": strict_count, "derivation": "15%完整扫描为0，12%为其严格子集"},
            {"box_range_limit_pct": 15.0, "box_min_trading_days": 20, "unique_stocks": strict_count, "derivation": "完整全市场重识别"},
            {"box_range_limit_pct": 18.0, "box_min_trading_days": 20, "unique_stocks": len(edge18), "derivation": "完整全市场重识别"},
            {"box_range_limit_pct": 20.0, "box_min_trading_days": 20, "unique_stocks": len(edge), "derivation": "完整全市场重识别"},
            {"box_range_limit_pct": 15.0, "box_min_trading_days": 25, "unique_stocks": strict_count, "derivation": "15%/20日完整扫描为0，25日为其严格子集"},
            {"box_range_limit_pct": 15.0, "box_min_trading_days": 30, "unique_stocks": strict_count, "derivation": "15%/20日完整扫描为0，30日为其严格子集"},
        ]
    )
    sensitivity.to_csv(OUTPUT / "box_threshold_sensitivity.csv", index=False, encoding="utf-8-sig")
    return edge, funnel, sensitivity


def _add_candidate_table(slide: Any, events: pd.DataFrame) -> None:
    columns = ["序", "代码", "名称", "突破日", "类型", "箱体", "归类", "相似度"]
    shape = slide.shapes.add_table(len(events) + 1, len(columns), Inches(0.56), Inches(2.72), Inches(12.20), Inches(3.10))
    table = shape.table
    widths = [0.48, 1.22, 1.32, 1.35, 0.85, 1.58, 1.62, 0.88]
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for index, value in enumerate(columns):
        table.cell(0, index).text = value
    for row_index, (_, event) in enumerate(events.iterrows(), start=1):
        values = [
            str(int(event["display_order"])), str(event["ts_code"]), str(event["stock_name"]),
            base._fmt_date(event["event_date"]), str(event["signal_type"]),
            f"{int(event['box_trading_days'])}日 / {float(event['box_range_pct']):.1f}%",
            str(event["sensitivity_band"]), f"{float(event['similarity_score']):.1f}",
        ]
        for column_index, value in enumerate(values):
            table.cell(row_index, column_index).text = value
    base._style_table(table, header_size=9.0, body_size=8.2)


def create_deck(
    events: pd.DataFrame,
    funnel: pd.DataFrame,
    sensitivity: pd.DataFrame,
    summary_chart: Path,
    chart_map: dict[str, Path],
    output: Path,
) -> Path:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    funnel_counts = dict(zip(funnel["stage_key"], funnel["unique_stocks"]))
    strict_count = int(funnel_counts["full15"])
    count_18 = int(sensitivity.loc[(sensitivity["box_range_limit_pct"].eq(18.0)) & (sensitivity["box_min_trading_days"].eq(20)), "unique_stocks"].iloc[0])
    count_20 = int(sensitivity.loc[(sensitivity["box_range_limit_pct"].eq(20.0)) & (sensitivity["box_min_trading_days"].eq(20)), "unique_stocks"].iloc[0])

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base._set_background(slide, base.NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(0.18), Inches(7.5))
    accent.fill.solid(); accent.fill.fore_color.rgb = RGBColor.from_string(base.BLUE); accent.line.fill.background()
    base._add_text(slide, "2026年6—7月全市场", 1.02, 1.05, 9.2, 0.55, size=18, color=base.CYAN, bold=True)
    base._add_text(slide, "箱体突破 · 60日新高", 1.02, 1.70, 10.7, 0.86, size=35, color=base.WHITE, bold=True)
    base._add_text(slide, f"严格主口径 {strict_count}只｜箱体振幅敏感性附录 {count_20}只", 1.05, 2.90, 10.9, 0.54, size=18, color=base.GOLD, bold=True)
    base._add_text(slide, "同一方法：前置放量拉升 → ≥20日箱体 → 涨停/大阳突破并创60日新高", 1.05, 3.70, 11.2, 0.48, size=13.5, color="DCEBFA")
    base._add_text(slide, "严格箱体振幅≤15%没有完整命中；个股页只展示18%/20%敏感性候选，不混入主样本", 1.05, 4.50, 11.0, 0.58, size=12.2, color=base.WHITE)
    base._add_text(slide, "冻结全市场5,672只｜前复权价格｜原始成交量｜数据截至2026-09-30", 1.05, 5.38, 10.5, 0.35, size=10.5, color="BFD7EA")
    base._add_text(slide, "历史结构筛选，仅供研究，不构成投资建议", 1.05, 6.72, 6.5, 0.24, size=8.5, color="94A3B8")

    slide = prs.slides.add_slide(prs.slide_layouts[6]); base._set_background(slide, base.LIGHT)
    base._add_header(slide, "研究口径", "沿用8—9月结构定义；修正复权收益字段")
    method = (
        "严格主口径\n"
        "① 事件日为2026-06-01至2026-07-31；涨停来自正式涨停事实表，或大阳线满足收涨≥6%、close/open实体≥5%、收盘位置≥0.70；日涨幅取pct_chg，不混算未复权pre_close。\n"
        "② 事件收盘严格高于此前60个实际交易日最高价。\n"
        "③ 紧邻事件日前箱体20—120日，最高/最低振幅≤15%，事件收盘突破箱体上沿。\n"
        f"④ {base.PREBOX_RULE_CN}。\n"
        "⑤ 均线粘合、多头排列、长均线上升、平台内开盘及突破量比只用于排序。"
    )
    sensitivity_note = (
        "本期结果\n"
        f"• 严格≤15%：{strict_count}只。\n"
        f"• 放宽至≤18%：{count_18}只。\n"
        f"• 放宽至≤20%：{count_20}只。\n\n"
        "个股页的蓝色箱体均来自对应的放宽阈值重识别；不是把旧名单事后过滤。\n\n"
        "所有事件身份仍满足涨停/大阳、60日新高和箱体前放量拉升。"
    )
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.55), Inches(1.05), Inches(7.55), Inches(5.75)); box.fill.solid(); box.fill.fore_color.rgb = RGBColor.from_string(base.WHITE); box.line.color.rgb = RGBColor.from_string(base.MID)
    box2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.35), Inches(1.05), Inches(4.43), Inches(5.75)); box2.fill.solid(); box2.fill.fore_color.rgb = RGBColor.from_string(base.WHITE); box2.line.color.rgb = RGBColor.from_string(base.MID)
    base._add_text(slide, method, 0.82, 1.30, 7.02, 5.15, size=12.6, color=base.DARK_TEXT)
    base._add_text(slide, sensitivity_note, 8.62, 1.35, 3.90, 4.85, size=12.2, color=base.DARK_TEXT)
    base._add_study_footer(slide, 2, DATA_END)

    slide = prs.slides.add_slide(prs.slide_layouts[6]); base._set_background(slide, base.LIGHT)
    base._add_header(slide, "严格主口径漏斗", "股票级唯一计数")
    cards = [
        (f"{int(funnel_counts['active']):,}", "6—7月实际有行情股票", base.BLUE),
        (f"{int(funnel_counts['signal']):,}", "6—7月涨停或大阳", base.RED),
        (f"{int(funnel_counts['new_high']):,}", "收盘突破前60日最高价", base.TEAL),
        (f"{int(funnel_counts['box15']):,}", "再满足≤15%箱体突破", base.GOLD),
        (f"{strict_count:,}", "再满足箱体前放量拉升", base.ORANGE),
    ]
    for index, (value, label, color) in enumerate(cards):
        base._add_card(slide, 0.38 + index * 2.57, 1.18, 2.28, 1.38, value, label, accent=color)
    base._add_text(slide, "本期严格样本为0不是数据缺失：主要断点发生在“≤15%紧邻箱体突破”之后，再叠加同口径箱体前放量拉升。为避免空报告，下页仅列箱体振幅敏感性候选。", 0.82, 3.40, 11.65, 1.12, size=16.0, color=base.NAVY, bold=True, align=PP_ALIGN.CENTER)
    base._add_text(slide, "放宽箱体振幅会重新识别箱体边界与前置拉升；18%和20%结果均来自冻结全市场日线重扫。", 1.05, 5.25, 11.15, 0.56, size=12.5, color=base.BLUE, align=PP_ALIGN.CENTER)
    base._add_study_footer(slide, 3, DATA_END)

    slide = prs.slides.add_slide(prs.slide_layouts[6]); base._set_background(slide, base.LIGHT)
    base._add_header(slide, "箱体振幅敏感性候选", "严格主样本仍为0只")
    base._add_card(slide, 0.66, 1.04, 2.80, 1.18, f"{strict_count}只", "主口径：振幅≤15%", accent=base.RED)
    base._add_card(slide, 3.78, 1.04, 2.80, 1.18, f"{count_18}只", "边缘：振幅≤18%", accent=base.TEAL)
    base._add_card(slide, 6.90, 1.04, 2.80, 1.18, f"{count_20}只", "宽松：振幅≤20%", accent=base.ORANGE)
    base._add_text(slide, f"20%层包含18%层；以下{count_20}只均不属于严格15%主样本", 9.88, 1.31, 2.55, 0.56, size=10.0, color=base.MUTED, align=PP_ALIGN.CENTER)
    _add_candidate_table(slide, events)
    base._add_study_footer(slide, 4, DATA_END)

    slide = prs.slides.add_slide(prs.slide_layouts[6]); base._set_background(slide, base.LIGHT)
    base._add_header(slide, "敏感性候选结构分布", "分数仅用于候选内部排序")
    slide.shapes.add_picture(str(summary_chart), Inches(0.42), Inches(0.92), width=Inches(12.48))
    base._add_study_footer(slide, 5, DATA_END)

    slide = prs.slides.add_slide(prs.slide_layouts[6]); base._set_background(slide, base.LIGHT)
    base._add_header(slide, "个股页阅读方式", f"下一页起为{count_20}只敏感性候选")
    guide = (
        "橙色阴影：箱体前放量拉升信号区间\n\n"
        "蓝色矩形：放宽阈值重新识别的紧邻箱体；箱体振幅均高于15%，因此不是严格主样本\n\n"
        "红色虚线：正式突破日；仍需满足涨停/大阳与收盘突破此前60日最高价\n\n"
        "灰色点线：突破日前60个实际交易日最高价\n\n"
        "价格与六条均线前复权，成交量为原始量；事件后行情仅作图表背景，不参与事件识别"
    )
    base._add_text(slide, guide, 1.05, 1.18, 11.15, 5.20, size=15.5, color=base.DARK_TEXT)
    base._add_study_footer(slide, 6, DATA_END)

    for slide_no, (_, event) in enumerate(events.iterrows(), start=7):
        slide = prs.slides.add_slide(prs.slide_layouts[6]); base._set_background(slide, base.LIGHT)
        base._add_header(slide, f"{int(event['display_order']):02d}｜{event['stock_name']}  {event['ts_code']}", f"{event['sensitivity_band']}｜{event['signal_type']}｜{base._fmt_date(event['event_date'])}")
        slide.shapes.add_picture(str(chart_map[str(event["ts_code"])]), Inches(0.38), Inches(1.02), width=Inches(9.25))
        color = base.TEAL if float(event["box_range_pct"]) <= 18.0 else base.ORANGE
        base._add_card(slide, 9.86, 1.02, 2.98, 1.10, f"{float(event['box_range_pct']):.1f}%", f"{event['sensitivity_band']}｜箱体振幅", accent=color)
        metric_text = (
            f"突破：{float(event['event_return_pct']):.1f}%｜量比 {float(event['event_volume_ratio']):.2f}×\n"
            f"60日新高裕度：{float(event['new_high_margin_pct']):.1f}%\n"
            f"箱体：{int(event['box_trading_days'])}日｜振幅 {float(event['box_range_pct']):.1f}%\n"
            f"超主阈值：+{float(event['box_range_pct']) - 15.0:.1f}个百分点\n"
            f"箱体前拉升：{float(event['prebox_impulse_rise_pct']):.1f}%\n"
            f"平台内开盘：{'是' if bool(event['event_open_inside_box']) else '否'}\n"
            f"六MA离差：{float(event['six_ma_spread_pct']):.1f}%\n"
            f"多头关系：{int(event['bullish_order_count'])}/5\n"
            f"MA60 / MA120上涨：{int(event['ma60_rising_streak'])} / {int(event['ma120_rising_streak'])}日"
        )
        base._add_text(slide, metric_text, 9.95, 2.38, 2.82, 3.32, size=10.8, color=base.DARK_TEXT)
        base._add_text(slide, "敏感性身份：按18%/20%阈值重识别后满足完整结构；不属于≤15%主样本", 9.95, 5.82, 2.82, 0.72, size=9.8, color=color, bold=True)
        base._add_study_footer(slide, slide_no, DATA_END)

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)
    return output


def _write_report(events: pd.DataFrame, funnel: pd.DataFrame, sensitivity: pd.DataFrame) -> None:
    strict_count = int(funnel.loc[funnel["stage_key"].eq("full15"), "unique_stocks"].iloc[0])
    count_18 = int(sensitivity.loc[(sensitivity["box_range_limit_pct"].eq(18.0)) & (sensitivity["box_min_trading_days"].eq(20)), "unique_stocks"].iloc[0])
    count_20 = int(sensitivity.loc[(sensitivity["box_range_limit_pct"].eq(20.0)) & (sensitivity["box_min_trading_days"].eq(20)), "unique_stocks"].iloc[0])
    lines = [
        "# 2026年6—7月全市场箱体突破60日新高研究",
        "",
        "## 结论",
        f"- 使用与8—9月报告完全相同的严格主口径，全市场命中 **{strict_count}只**。",
        f"- 箱体振幅从≤15%放宽至≤18%后命中{count_18}只；放宽至≤20%后命中{count_20}只。",
        f"- {count_20}只个股仅作为敏感性附录；均不计入严格主样本。",
        "",
        "## 严格主口径漏斗",
        "",
        funnel.drop(columns=["stage_key"]).to_markdown(index=False),
        "",
        "## 箱体阈值敏感性",
        "",
        sensitivity.to_markdown(index=False),
        "",
        "## 敏感性候选",
        "",
        "| 排名 | 代码 | 名称 | 突破日 | 信号 | 箱体 | 超15% | 相似度 |",
        "|---:|---|---|---|---|---|---:|---:|",
    ]
    for _, row in events.iterrows():
        lines.append(f"| {int(row['display_order'])} | {row['ts_code']} | {row['stock_name']} | {base._fmt_date(row['event_date'])} | {row['signal_type']} | {int(row['box_trading_days'])}日/{float(row['box_range_pct']):.1f}% | +{float(row['box_range_pct'])-15.0:.1f}pct | {float(row['similarity_score']):.1f} |")
    lines.extend([
        "",
        "## 口径说明",
        "- 日线经 `coreClient.data_provider`、`source=\"database_only\"` 冻结；价格和均线前复权，成交量为原始量。",
        "- 日涨幅使用数据源 `pct_chg`；禁止用前复权 `close` 除以未复权 `pre_close`。",
        "- 主口径事件窗为2026-06-01至2026-07-31；数据截止2026-09-30，事件后数据只用于图表背景。",
        "- 60日新高定义为事件收盘严格高于此前60个实际交易日最高价。",
        "- 箱体结束于事件前一交易日；前置拉升结束早于箱体开始；三段不重叠。",
        f"- 前置拉升：{base.PREBOX_RULE_CN}。",
        "- 18%与20%结果均从冻结全市场日线重新识别箱体和前置拉升，不是对15%结果表做事后过滤。",
        "- 历史结构研究，不构成投资建议。",
    ])
    (OUTPUT / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def _verify(pptx_path: Path, pdf_path: Path, events: pd.DataFrame, charts: list[Path]) -> dict[str, Any]:
    prs = Presentation(pptx_path)
    text = "\n".join(shape.text for slide in prs.slides for shape in slide.shapes if hasattr(shape, "text"))
    missing = [code for code in events["ts_code"].astype(str) if code not in text]
    hashes = [_sha256(path) for path in charts]
    pictures = sum(1 for slide in prs.slides for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)
    import fitz
    with fitz.open(pdf_path) as document:
        pdf_pages = len(document)
    expected = expected_deck_pages(len(events))
    return {
        "pptx_pages": len(prs.slides), "pdf_pages": pdf_pages, "expected_pages": expected,
        "missing_stock_codes": missing, "event_chart_count": len(charts),
        "unique_event_chart_hashes": len(set(hashes)), "pptx_picture_count": pictures,
        "strict_main_stock_count": 0, "sensitivity_stock_count": len(events),
        "passed": len(prs.slides) == pdf_pages == expected and not missing and len(charts) == len(set(hashes)) == len(events),
    }


def main() -> int:
    OUTPUT.mkdir(parents=True, exist_ok=True)
    daily, _, _ = _load_inputs()
    events, funnel, sensitivity = _prepare_tables()
    chart_dir = OUTPUT / "charts"
    rendered_dir = OUTPUT / "rendered_slides"
    contact_dir = OUTPUT / "contact_sheets"
    for directory in (chart_dir, rendered_dir, contact_dir):
        if directory.exists():
            shutil.rmtree(directory)
        directory.mkdir(parents=True, exist_ok=True)
    selected = set(events["ts_code"].astype(str))
    frames = {str(code): group.copy() for code, group in daily.loc[daily["ts_code"].isin(selected)].groupby("ts_code", sort=False)}
    chart_map: dict[str, Path] = {}
    for _, event in events.iterrows():
        code = str(event["ts_code"])
        chart = chart_dir / f"{int(event['display_order']):02d}_{code.replace('.', '_')}_{event['event_date']}.png"
        base.create_event_chart(frames[code], event, chart)
        chart_map[code] = chart
    summary = OUTPUT / "sensitivity_summary.png"
    base.create_summary_chart(events, summary)

    pptx = OUTPUT / "2026年6至7月全市场箱体突破60日新高研究.pptx"
    create_deck(events, funnel, sensitivity, summary, chart_map, pptx)
    pdf = base.convert_ppt_to_pdf(pptx, OUTPUT)
    slide_images = base.render_pdf(pdf, rendered_dir)
    contacts = base.create_contact_sheets(slide_images, contact_dir, per_sheet=12)
    _write_report(events, funnel, sensitivity)
    verification = _verify(pptx, pdf, events, list(chart_map.values()))
    (OUTPUT / "verification.json").write_text(json.dumps(verification, ensure_ascii=False, indent=2), encoding="utf-8")

    funnel_counts = dict(zip(funnel["stage_key"], funnel["unique_stocks"]))
    count_18 = int(sensitivity.loc[(sensitivity["box_range_limit_pct"].eq(18.0)) & (sensitivity["box_min_trading_days"].eq(20)), "unique_stocks"].iloc[0])
    count_20 = int(sensitivity.loc[(sensitivity["box_range_limit_pct"].eq(20.0)) & (sensitivity["box_min_trading_days"].eq(20)), "unique_stocks"].iloc[0])
    git_head = subprocess.run(
        ["git", "rev-parse", "HEAD"], cwd=base.ROOT, text=True, capture_output=True, check=True
    ).stdout.strip()
    metadata = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "scope": {"event_start": EVENT_START, "event_end": EVENT_END, "data_end": DATA_END, "description": "2026 YTD through 2026-09-30"},
        "data_governance": {"provider": "coreClient.data_provider", "source": "database_only", "price_adjustment": "qfq", "volume": "raw_unadjusted"},
        "universe": {"daily_rows": len(daily), "frozen_input_stocks": int(daily["ts_code"].nunique()), "event_period_active_stocks": int(funnel_counts["active"]), "event_period_trading_days": int(daily.loc[daily["trade_date"].between(EVENT_START, EVENT_END), "trade_date"].nunique())},
        "results": {"strict_main_stocks": int(funnel_counts["full15"]), "box_18pct_sensitivity_stocks": count_18, "box_20pct_sensitivity_stocks": count_20},
        "definitions": {"strict_main": "same as the Aug-Sep study, including box range<=15%", "daily_return": "pct_chg from the provider; qfq close/pre_close is forbidden because provider pre_close remains unadjusted", "new_high": "event close > max high of prior 60 actual trading rows", "prebox_impulse": base.PREBOX_RULE_CN, "sensitivity_display": "individual pages are <=20% box-range re-identification; none are strict-main members"},
        "code_provenance": {"git_head": git_head, "working_tree_note": "repository contains pre-existing uncommitted changes; script and test hashes are frozen below"},
        "bias_disclosure": ["No future rows are used to identify an event.", "Post-event rows through 2026-09-30 appear only as chart context.", "The 18% and 20% candidates are sensitivity results, not strict-main matches.", "Similarity scores are ranking aids, not return forecasts."],
        "verification": verification,
        "contact_sheets": [str(path) for path in contacts],
        "files": {},
    }
    paths = [
        pptx, pdf, OUTPUT / "daily_input_qfq_rawvol.csv.gz", OUTPUT / "stock_basic.csv", OUTPUT / "kpl_limit_up.csv",
        OUTPUT / "all_structural_events.csv", OUTPUT / "ranked_unique_candidates.csv",
        OUTPUT / "sensitivity_events_box18.csv", OUTPUT / "sensitivity_unique_box18.csv",
        OUTPUT / "sensitivity_events_box20.csv", OUTPUT / "sensitivity_unique_box20.csv",
        OUTPUT / "sensitivity_ranked_candidates_box20.csv", OUTPUT / "strict_main_funnel_raw.csv",
        OUTPUT / "screening_funnel.csv", OUTPUT / "box_threshold_sensitivity.csv", summary,
        OUTPUT / "report.md", OUTPUT / "verification.json",
        base.ROOT / "policyStudy/policy/题材涨停研究/scripts/study_aug_sep_2026_box_breakout_ppt.py",
        base.ROOT / "policyStudy/policy/题材涨停研究/scripts/test_study_aug_sep_2026_box_breakout_ppt.py",
        Path(__file__),
        base.ROOT / "policyStudy/policy/题材涨停研究/scripts/test_study_jun_jul_2026_box_breakout_ppt.py",
        *chart_map.values(),
    ]
    for path in paths:
        metadata["files"][path.name] = {"sha256": _sha256(path), "bytes": path.stat().st_size}
    (OUTPUT / "metadata.json").write_text(json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({"output": str(OUTPUT), "strict_main": int(funnel_counts["full15"]), "sensitivity_18": count_18, "sensitivity_20": count_20, "pptx": str(pptx), "pdf": str(pdf), "verification": verification}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

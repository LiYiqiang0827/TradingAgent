"""Build a PPT for first MA30 pullbacks with MA5/MA10/MA30 cohesion."""

from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
import shutil
import subprocess
from typing import Any, Sequence

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib import font_manager
from matplotlib.patches import Rectangle
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from study_first_ma30_pullback_cohesion_timing import compute_rising_streak


ROOT = Path(__file__).resolve().parents[5]
BASE_STUDY_DIR = (
    ROOT / "outputs" / "first_ma30_pullback_cohesion_timing" / "2026_through_20260930_qfq"
)
DEFAULT_DAILY = (
    ROOT / "outputs" / "ma60_ma120_top100_ppt" / "2026_through_20260930_qfq" / "daily_input_qfq.csv.gz"
)
DEFAULT_OUTPUT = (
    ROOT / "outputs" / "first_ma30_pullback_cohesion_ppt" / "2026_through_20260930_qfq"
)
FONT_REGULAR = Path.home() / "Library" / "Fonts" / "NotoSansCJKsc-Regular.otf"
FONT_BOLD = Path.home() / "Library" / "Fonts" / "NotoSansCJKsc-Bold.otf"
FONT_NAME = "Noto Sans CJK SC"
NAVY = "0B132B"
BLUE = "1C77C3"
CYAN = "35A7FF"
TEAL = "2A9D8F"
RED = "E63946"
ORANGE = "F4A261"
GOLD = "E9C46A"
WHITE = "FFFFFF"
LIGHT = "F5F7FA"
MID = "D8E1EA"
DARK_TEXT = "14213D"
MUTED = "5F6B7A"


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--daily-file", type=Path, default=DEFAULT_DAILY)
    parser.add_argument("--event-file", type=Path, default=BASE_STUDY_DIR / "first_pullback_event_detail.csv")
    parser.add_argument("--unique-event-file", type=Path, default=BASE_STUDY_DIR / "unique_stock_first_pullback_detail.csv")
    parser.add_argument("--all-first-file", type=Path, default=BASE_STUDY_DIR / "all_first_ma30_pullbacks.csv")
    parser.add_argument("--coverage-file", type=Path, default=BASE_STUDY_DIR / "coverage_summary.csv")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    return parser.parse_args(argv)


def _compact_date(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip().replace("-", "").split(".")[0]


def _fmt_date(value: object) -> str:
    text = _compact_date(value)
    return f"{text[:4]}-{text[4:6]}-{text[6:8]}" if len(text) == 8 else text


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clean_json(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _clean_json(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_clean_json(item) for item in value]
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        value = float(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def compute_true_streak(mask: pd.Series) -> pd.Series:
    """Count consecutive true trading-day states, resetting on false/missing."""
    values = mask.fillna(False).astype(bool)
    output: list[int] = []
    current = 0
    for flag in values:
        current = current + 1 if bool(flag) else 0
        output.append(current)
    return pd.Series(output, index=mask.index, dtype=int)


def compute_event_state_metrics(frame: pd.DataFrame, *, position: int) -> dict[str, Any]:
    """Compute MA streak, bullish-order streak, and relative MA gaps at an event."""
    if position < 0 or position >= len(frame):
        raise IndexError(position)
    ma = frame[["ma30", "ma60", "ma120"]].apply(pd.to_numeric, errors="coerce")
    rising = {period: compute_rising_streak(ma[f"ma{period}"]) for period in (30, 60, 120)}
    bull = ma["ma30"].gt(ma["ma60"]) & ma["ma60"].gt(ma["ma120"])
    bull_streak = compute_true_streak(bull)
    all_rising_bull = bull.copy()
    for period in (30, 60, 120):
        all_rising_bull &= rising[period].gt(0)
    all_rising_bull_streak = compute_true_streak(all_rising_bull)
    row = ma.iloc[position]
    return {
        "ma30_streak_at_position": int(rising[30].iloc[position]),
        "ma60_streak_at_position": int(rising[60].iloc[position]),
        "ma120_streak_at_position": int(rising[120].iloc[position]),
        "ma30_ma60_ma120_bull_streak_at_position": int(bull_streak.iloc[position]),
        "all_three_rising_bull_streak_at_position": int(all_rising_bull_streak.iloc[position]),
        "ma30_vs_ma60_pct": float(100.0 * (row["ma30"] / row["ma60"] - 1.0)),
        "ma60_vs_ma120_pct": float(100.0 * (row["ma60"] / row["ma120"] - 1.0)),
    }


def attach_event_metrics(
    events: pd.DataFrame,
    frames_by_code: dict[str, pd.DataFrame],
) -> pd.DataFrame:
    """Attach independently recomputed MA-state metrics to event rows."""
    records: list[dict[str, Any]] = []
    for event in events.to_dict("records"):
        code = str(event["ts_code"])
        if code not in frames_by_code:
            raise ValueError(f"missing daily frame for {code}")
        frame = frames_by_code[code]
        dates = frame["trade_date"].astype(str).str.replace("-", "", regex=False).str.split(".").str[0]
        target = str(event["pullback_date"]).replace("-", "").split(".")[0]
        positions = np.flatnonzero(dates.eq(target).to_numpy())
        if len(positions) != 1:
            raise ValueError(f"event date lookup failed for {code} {target}: {len(positions)}")
        record = dict(event)
        metrics = compute_event_state_metrics(frame, position=int(positions[0]))
        for period in (60, 120):
            key = f"ma{period}_streak_at_position"
            existing = record.get(key)
            if existing is not None and not pd.isna(existing) and int(existing) != int(metrics[key]):
                raise ValueError(
                    f"streak mismatch for {code} {target} {key}: "
                    f"formal={existing}, recomputed={metrics[key]}"
                )
        record.update(metrics)
        records.append(record)
    return pd.DataFrame(records)


def summarize_extended_metrics(events: pd.DataFrame) -> pd.DataFrame:
    """Summarize event-time MA states by scope, retaining zero streaks."""
    rows: list[dict[str, Any]] = []
    for scope, group in events.groupby("scope", sort=True):
        row: dict[str, Any] = {"scope": scope, "event_count": int(len(group))}
        for period in (30, 60, 120):
            values = pd.to_numeric(group[f"ma{period}_streak_at_position"], errors="coerce")
            positive = values.loc[values.gt(0)]
            row[f"ma{period}_active_count"] = int(len(positive))
            row[f"ma{period}_streak_all_median"] = float(values.median())
            row[f"ma{period}_streak_all_p25"] = float(values.quantile(0.25))
            row[f"ma{period}_streak_all_p75"] = float(values.quantile(0.75))
            row[f"ma{period}_streak_active_median"] = float(positive.median()) if len(positive) else float("nan")
            row[f"ma{period}_streak_active_p25"] = float(positive.quantile(0.25)) if len(positive) else float("nan")
            row[f"ma{period}_streak_active_p75"] = float(positive.quantile(0.75)) if len(positive) else float("nan")
        bull = pd.to_numeric(group["ma30_ma60_ma120_bull_streak_at_position"], errors="coerce")
        bull_positive = bull.loc[bull.gt(0)]
        row["bull_alignment_active_count"] = int(len(bull_positive))
        row["bull_alignment_all_median"] = float(bull.median())
        row["bull_alignment_active_median"] = float(bull_positive.median()) if len(bull_positive) else float("nan")
        row["bull_alignment_active_p25"] = float(bull_positive.quantile(0.25)) if len(bull_positive) else float("nan")
        row["bull_alignment_active_p75"] = float(bull_positive.quantile(0.75)) if len(bull_positive) else float("nan")
        simultaneous = pd.to_numeric(group["all_three_rising_bull_streak_at_position"], errors="coerce")
        row["all_three_rising_bull_active_count"] = int(simultaneous.gt(0).sum())
        row["all_three_rising_bull_all_median"] = float(simultaneous.median())
        gap = pd.to_numeric(group["ma30_vs_ma60_pct"], errors="coerce")
        row["ma30_vs_ma60_pct_median"] = float(gap.median())
        row["ma30_vs_ma60_pct_p25"] = float(gap.quantile(0.25))
        row["ma30_vs_ma60_pct_p75"] = float(gap.quantile(0.75))
        rows.append(row)
    return pd.DataFrame(rows)


def prepare_stock_frames(daily: pd.DataFrame) -> dict[str, pd.DataFrame]:
    """Recompute all MAs from qfq close on each stock's actual trading-day sequence."""
    required = {"ts_code", "trade_date", "open", "high", "low", "close", "vol"}
    missing = sorted(required - set(daily.columns))
    if missing:
        raise ValueError(f"daily input missing columns: {missing}")
    data = daily.copy()
    data["ts_code"] = data["ts_code"].astype(str)
    data["trade_date"] = data["trade_date"].map(_compact_date)
    if data.duplicated(["ts_code", "trade_date"]).any():
        raise ValueError("duplicate stock-date rows")
    frames: dict[str, pd.DataFrame] = {}
    for code, group in data.groupby("ts_code", sort=False):
        frame = group.sort_values("trade_date", kind="mergesort").reset_index(drop=True).copy()
        close = pd.to_numeric(frame["close"], errors="coerce")
        for period in (5, 10, 30, 60, 120):
            frame[f"ma{period}"] = close.rolling(period, min_periods=period).mean()
        frames[str(code)] = frame
    return frames


def _register_fonts() -> None:
    for path in (FONT_REGULAR, FONT_BOLD):
        if path.exists():
            font_manager.fontManager.addfont(str(path))
    plt.rcParams["font.family"] = FONT_NAME
    plt.rcParams["axes.unicode_minus"] = False


def _draw_candlesticks(ax: Any, frame: pd.DataFrame) -> None:
    width = 0.58
    for x, row in enumerate(frame.itertuples(index=False)):
        open_price = float(row.open)
        close_price = float(row.close)
        high_price = float(row.high)
        low_price = float(row.low)
        color = f"#{RED}" if close_price >= open_price else f"#{TEAL}"
        ax.vlines(x, low_price, high_price, color=color, linewidth=0.8, zorder=2)
        bottom = min(open_price, close_price)
        height = max(abs(close_price - open_price), max(close_price, 1.0) * 0.001)
        ax.add_patch(
            Rectangle(
                (x - width / 2.0, bottom),
                width,
                height,
                facecolor=color,
                edgecolor=color,
                linewidth=0.6,
                zorder=3,
            )
        )


def create_event_chart(frame: pd.DataFrame, event: pd.Series, output: Path) -> None:
    """Draw qfq candles, volume, all five MAs, anchor interval, and first pullback."""
    dates = frame["trade_date"].map(_compact_date)
    event_date = _compact_date(event["pullback_date"])
    anchor_start = _compact_date(event["anchor_start_date"])
    event_positions = np.flatnonzero(dates.eq(event_date).to_numpy())
    anchor_positions = np.flatnonzero(dates.eq(anchor_start).to_numpy())
    if len(event_positions) != 1 or len(anchor_positions) != 1:
        raise ValueError(f"chart date lookup failed for {event['ts_code']}")
    event_pos = int(event_positions[0])
    anchor_pos = int(anchor_positions[0])
    lower = max(0, min(anchor_pos, event_pos - 80) - 15)
    upper = min(len(frame) - 1, event_pos + 30)
    view = frame.iloc[lower : upper + 1].reset_index(drop=True).copy()
    event_x = event_pos - lower
    anchor_x = anchor_pos - lower

    fig = plt.figure(figsize=(12.2, 6.35), dpi=150, facecolor="white")
    grid = fig.add_gridspec(5, 1, hspace=0.06)
    ax = fig.add_subplot(grid[:4, 0])
    vol_ax = fig.add_subplot(grid[4, 0], sharex=ax)
    _draw_candlesticks(ax, view)
    colors = {5: ORANGE, 10: GOLD, 30: BLUE, 60: "7B2CBF", 120: TEAL}
    for period in (5, 10, 30, 60, 120):
        ax.plot(
            np.arange(len(view)),
            pd.to_numeric(view[f"ma{period}"], errors="coerce"),
            color=f"#{colors[period]}",
            linewidth=1.15 if period < 30 else 1.45,
            label=f"MA{period}",
            zorder=4,
        )
    ax.axvspan(max(anchor_x - 0.5, -0.5), event_x + 0.5, color=f"#{CYAN}", alpha=0.08, zorder=0)
    ax.axvline(anchor_x, color=f"#{MUTED}", linewidth=1.0, linestyle=":", zorder=5)
    ax.axvline(event_x, color=f"#{RED}", linewidth=1.6, linestyle="--", zorder=5)
    event_close = float(view.iloc[event_x]["close"])
    ax.scatter([event_x], [event_close], color=f"#{RED}", s=45, zorder=6, edgecolor="white", linewidth=0.8)
    ax.annotate(
        f"首次回调 { _fmt_date(event_date) }",
        xy=(event_x, event_close),
        xytext=(8, 18),
        textcoords="offset points",
        fontsize=8,
        color=f"#{RED}",
        fontweight="bold",
        arrowprops={"arrowstyle": "->", "color": f"#{RED}", "lw": 0.8},
    )
    ax.legend(loc="upper left", ncol=5, frameon=False, fontsize=7.5)
    ax.grid(axis="y", color="#E7EDF3", linewidth=0.6)
    ax.set_ylabel("前复权价格", fontsize=8, color=f"#{MUTED}")
    ax.tick_params(axis="both", labelsize=7, colors=f"#{MUTED}")
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["left", "bottom"]].set_color("#CBD5E1")
    plt.setp(ax.get_xticklabels(), visible=False)

    volumes = pd.to_numeric(view["vol"], errors="coerce").fillna(0.0)
    volume_colors = [
        f"#{RED}" if float(close) >= float(open_) else f"#{TEAL}"
        for open_, close in zip(view["open"], view["close"])
    ]
    vol_ax.bar(np.arange(len(view)), volumes, color=volume_colors, width=0.62, alpha=0.75)
    vol_ax.axvline(event_x, color=f"#{RED}", linewidth=1.2, linestyle="--")
    vol_ax.grid(axis="y", color="#EEF2F6", linewidth=0.5)
    vol_ax.set_ylabel("成交量", fontsize=7, color=f"#{MUTED}")
    vol_ax.tick_params(axis="both", labelsize=6.5, colors=f"#{MUTED}")
    vol_ax.spines[["top", "right"]].set_visible(False)
    vol_ax.spines[["left", "bottom"]].set_color("#CBD5E1")
    step = max(1, len(view) // 8)
    ticks = sorted(set(list(range(0, len(view), step)) + [event_x, len(view) - 1]))
    vol_ax.set_xticks(ticks)
    vol_ax.set_xticklabels(
        [_fmt_date(view.iloc[index]["trade_date"])[5:] for index in ticks],
        rotation=0,
        ha="center",
        fontsize=6.5,
    )
    fig.suptitle(
        f"{event['stock_name']}  {event['ts_code']}  |  MA{int(event['anchor_period'])}锚点 "
        f"{_fmt_date(anchor_start)} → {_fmt_date(event_date)}",
        x=0.01,
        y=0.995,
        ha="left",
        va="top",
        fontsize=11,
        color=f"#{DARK_TEXT}",
        fontweight="bold",
    )
    fig.text(
        0.99,
        0.99,
        "蓝色浅区=锚点起点至首次回调；红虚线=事件日",
        ha="right",
        va="top",
        fontsize=7.5,
        color=f"#{MUTED}",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def create_distribution_chart(events: pd.DataFrame, output: Path) -> None:
    """Create a compact visual summary for the unique-stock main sample."""
    _register_fonts()
    fig, axes = plt.subplots(1, 3, figsize=(12.4, 4.2), dpi=160, facecolor="white")
    streak_data = [
        pd.to_numeric(events[f"ma{period}_streak_at_position"], errors="coerce").to_numpy()
        for period in (30, 60, 120)
    ]
    box = axes[0].boxplot(streak_data, patch_artist=True, labels=["MA30", "MA60", "MA120"])
    for patch, color in zip(box["boxes"], [ORANGE, "7B2CBF", TEAL]):
        patch.set_facecolor(f"#{color}")
        patch.set_alpha(0.78)
    for median in box["medians"]:
        median.set_color(f"#{NAVY}")
        median.set_linewidth(1.5)
    for index, values in enumerate(streak_data, start=1):
        jitter = np.linspace(-0.09, 0.09, len(values))
        axes[0].scatter(index + jitter, values, s=13, color=f"#{NAVY}", alpha=0.48, zorder=3)
    axes[0].set_title("回调日连续上涨天数（含0）", fontsize=10, fontweight="bold")
    axes[0].set_ylabel("实际交易日")

    bull = pd.to_numeric(events["ma30_ma60_ma120_bull_streak_at_position"], errors="coerce")
    bins = np.arange(0, max(10, int(bull.max()) + 10), 10)
    if len(bins) < 2:
        bins = np.array([0, 10])
    axes[1].hist(bull, bins=bins, color=f"#{BLUE}", alpha=0.85, edgecolor="white")
    axes[1].axvline(bull.median(), color=f"#{RED}", linestyle="--", linewidth=1.5, label=f"中位数 {bull.median():.0f}日")
    axes[1].set_title("MA30>MA60>MA120连续排列", fontsize=10, fontweight="bold")
    axes[1].set_xlabel("连续交易日")
    axes[1].set_ylabel("股票数")
    axes[1].legend(frameon=False, fontsize=8)

    gaps = pd.to_numeric(events["ma30_vs_ma60_pct"], errors="coerce")
    axes[2].hist(gaps, bins=10, color=f"#{ORANGE}", alpha=0.88, edgecolor="white")
    axes[2].axvline(gaps.median(), color=f"#{RED}", linestyle="--", linewidth=1.5, label=f"中位数 {gaps.median():.1f}%")
    axes[2].axvline(0, color=f"#{MUTED}", linewidth=0.8)
    axes[2].set_title("回调日MA30相对MA60间距", fontsize=10, fontweight="bold")
    axes[2].set_xlabel("(MA30 / MA60 - 1) × 100%")
    axes[2].set_ylabel("股票数")
    axes[2].legend(frameon=False, fontsize=8)
    for ax in axes:
        ax.grid(axis="y", color="#E7EDF3", linewidth=0.6)
        ax.spines[["top", "right"]].set_visible(False)
        ax.tick_params(labelsize=8, colors=f"#{MUTED}")
    fig.tight_layout()
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def _set_background(slide: Any, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(color)


def _add_text(
    slide: Any,
    text: str,
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: float = 18,
    color: str = DARK_TEXT,
    bold: bool = False,
    align: Any = PP_ALIGN.LEFT,
    valign: Any = MSO_ANCHOR.TOP,
    margin: float = 0.04,
) -> Any:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    box.text_frame.clear()
    box.text_frame.margin_left = Inches(margin)
    box.text_frame.margin_right = Inches(margin)
    box.text_frame.margin_top = Inches(margin)
    box.text_frame.margin_bottom = Inches(margin)
    box.text_frame.vertical_anchor = valign
    paragraph = box.text_frame.paragraphs[0]
    paragraph.alignment = align
    run = paragraph.add_run()
    run.text = str(text)
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = _rgb(color)
    return box


def _add_bullets(
    slide: Any,
    bullets: list[str],
    left: float,
    top: float,
    width: float,
    height: float,
    *,
    size: float = 15,
    color: str = DARK_TEXT,
) -> Any:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.05)
    frame.margin_right = Inches(0.05)
    for index, item in enumerate(bullets):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = item
        paragraph.font.name = FONT_NAME
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = _rgb(color)
        paragraph.level = 0
        paragraph.space_after = Pt(8)
        paragraph.text = f"• {item}"
    return box


def _add_header(slide: Any, title: str, subtitle: str = "") -> None:
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.12))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(BLUE)
    shape.line.fill.background()
    _add_text(slide, title, 0.42, 0.26, 9.9, 0.48, size=23, bold=True)
    if subtitle:
        _add_text(slide, subtitle, 9.75, 0.33, 3.1, 0.30, size=9.5, color=MUTED, align=PP_ALIGN.RIGHT)


def _add_footer(slide: Any, slide_no: int) -> None:
    _add_text(
        slide,
        "前复权日线｜截至2026-09-30｜相似结构观察，不构成买入信号",
        0.42,
        7.18,
        9.7,
        0.18,
        size=7.5,
        color=MUTED,
    )
    _add_text(slide, str(slide_no), 12.35, 7.18, 0.5, 0.18, size=7.5, color=MUTED, align=PP_ALIGN.RIGHT)


def _add_card(
    slide: Any,
    left: float,
    top: float,
    width: float,
    height: float,
    value: str,
    label: str,
    *,
    accent: str = BLUE,
) -> None:
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(width),
        Inches(height),
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(WHITE)
    shape.line.color.rgb = _rgb(MID)
    shape.line.width = Pt(0.8)
    bar = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE,
        Inches(left),
        Inches(top),
        Inches(0.08),
        Inches(height),
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = _rgb(accent)
    bar.line.fill.background()
    _add_text(slide, value, left + 0.18, top + 0.16, width - 0.28, 0.42, size=23, bold=True, color=accent)
    _add_text(slide, label, left + 0.18, top + 0.64, width - 0.28, height - 0.70, size=9.5, color=MUTED)


def _style_table(table: Any, *, header_size: float = 10, body_size: float = 8.5) -> None:
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.margin_left = Inches(0.04)
            cell.margin_right = Inches(0.04)
            cell.margin_top = Inches(0.03)
            cell.margin_bottom = Inches(0.03)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(NAVY if row_index == 0 else (WHITE if row_index % 2 else LIGHT))
            for paragraph in cell.text_frame.paragraphs:
                paragraph.alignment = PP_ALIGN.CENTER
                for run in paragraph.runs:
                    run.font.name = FONT_NAME
                    run.font.size = Pt(header_size if row_index == 0 else body_size)
                    run.font.bold = row_index == 0
                    run.font.color.rgb = _rgb(WHITE if row_index == 0 else DARK_TEXT)


def _add_event_table_slide(prs: Presentation, page: pd.DataFrame, page_no: int, page_count: int) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, f"29只联合去重主样本（{page_no}/{page_count}）", "按首次回调日期排序")
    headers = ["序", "股票", "回调日", "锚点", "MA30/60/120上涨", "多头排列", "MA30-MA60", "三线离差"]
    table = slide.shapes.add_table(
        len(page) + 1,
        len(headers),
        Inches(0.38),
        Inches(1.02),
        Inches(12.58),
        Inches(5.72),
    ).table
    widths = [0.45, 2.15, 1.25, 0.75, 2.05, 1.05, 1.15, 1.05]
    for index, width in enumerate(widths):
        table.columns[index].width = Inches(width)
    for index, header in enumerate(headers):
        table.cell(0, index).text = header
    for row_index, (_, row) in enumerate(page.iterrows(), start=1):
        values = [
            str(int(row["display_order"])),
            f"{row['stock_name']}\n{row['ts_code']}",
            _fmt_date(row["pullback_date"])[5:],
            f"MA{int(row['anchor_period'])}",
            f"{int(row['ma30_streak_at_position'])} / {int(row['ma60_streak_at_position'])} / {int(row['ma120_streak_at_position'])}",
            f"{int(row['ma30_ma60_ma120_bull_streak_at_position'])}日",
            f"{float(row['ma30_vs_ma60_pct']):+.1f}%",
            f"{float(row['ma5_ma10_ma30_spread_pct']):.1f}%",
        ]
        for col_index, value in enumerate(values):
            table.cell(row_index, col_index).text = value
    _style_table(table, header_size=9.5, body_size=8.2)
    _add_footer(slide, len(prs.slides))


def build_deck(
    *,
    unique_events: pd.DataFrame,
    anchor_events: pd.DataFrame,
    summary: pd.DataFrame,
    coverage: pd.DataFrame,
    chart_paths: dict[str, Path],
    distribution_chart: Path,
    output: Path,
) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, NAVY)
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.22), Inches(0.12), Inches(4.7))
    accent.fill.solid()
    accent.fill.fore_color.rgb = _rgb(CYAN)
    accent.line.fill.background()
    _add_text(slide, "MA60 / MA120上升阶段\n首次回调MA30与短均线粘合", 1.05, 1.28, 10.9, 1.65, size=30, color=WHITE, bold=True)
    _add_text(slide, "事件时点的MA30、MA60、MA120状态与个股图册", 1.08, 3.18, 10.2, 0.50, size=18, color=CYAN)
    _add_text(slide, "178只长趋势候选池｜前复权日线｜截至2026-09-30", 1.08, 4.16, 9.4, 0.36, size=12, color="BFD7EA")
    _add_text(slide, "主口径：先固定第一次MA30回调，再检查同日MA5/MA10/MA30是否粘合", 1.08, 4.72, 10.4, 0.62, size=12, color=WHITE)
    _add_text(slide, "研究用途，不构成买入建议", 1.08, 6.72, 5.0, 0.26, size=9, color="94A3B8")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, "口径：把“初期回调”固定为长均线上升段内的第一次MA30回调")
    _add_bullets(
        slide,
        [
            "股票池：MA60榜107只、MA120榜126只，并集178只；排除2个MA120左截断锚点。",
            "第一次回调：此前10个实际交易日至少有一天收盘高于自身MA30 5%；当日最低价≤MA30上方1%，收盘≥MA30下方2%。",
            "粘合：只检查已固定的第一次回调日；MA5/MA10/MA30的(最大值-最小值)/均值≤3%，失败不得后跳。",
            "连续上涨：MA_t > MA_(t-1)；持平、下降或缺失均清零。",
            "连续多头排列：每天同时满足MA30 > MA60 > MA120，统计截至回调日的连续实际交易日数。",
            "MA30—MA60间距：(MA30 / MA60 - 1) × 100%；正值表示MA30在MA60上方。",
        ],
        0.72,
        1.12,
        11.9,
        4.85,
        size=14,
    )
    note = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.72), Inches(6.04), Inches(11.9), Inches(0.72))
    note.fill.solid(); note.fill.fore_color.rgb = _rgb("E7F3FF"); note.line.color.rgb = _rgb("B9DCF7")
    _add_text(slide, "“初期”不另设10日或20日上限；由“该锚点内第一次符合价格回调定义的交易日”保证，另在结果中披露锚点已运行天数。", 0.92, 6.20, 11.4, 0.34, size=10.5, color=BLUE, bold=True)
    _add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, "覆盖：231个可搜索锚点全部出现首次MA30回调，35个当日形成粘合")
    values = [
        ("178", "候选池股票", BLUE),
        ("231", "非左截断可搜索锚点", CYAN),
        ("35", "锚点级严格事件", ORANGE),
        ("29", "联合去重主样本", RED),
    ]
    for index, (value, label, color) in enumerate(values):
        _add_card(slide, 0.58 + index * 3.12, 1.18, 2.72, 1.28, value, label, accent=color)
    ma60 = coverage.loc[coverage["anchor_period"].eq(60)].iloc[0]
    ma120 = coverage.loc[coverage["anchor_period"].eq(120)].iloc[0]
    _add_bullets(
        slide,
        [
            f"MA60锚点：{int(ma60['eligible_count'])}个可搜索，{int(ma60['first_pullback_cohesion_count'])}个首次回调当日粘合（{float(ma60['first_pullback_cohesion_pct_of_eligible']):.1f}%）。",
            f"MA120锚点：{int(ma120['eligible_count'])}个可搜索，{int(ma120['first_pullback_cohesion_count'])}个首次回调当日粘合（{float(ma120['first_pullback_cohesion_pct_of_eligible']):.1f}%）。",
            f"35个锚点事件涉及{anchor_events['ts_code'].nunique()}只代码；联合口径先按股票跨锚点取最早首次回调，再判断粘合，保留29只。",
            "因此主结论使用29只，避免MA60/MA120重叠股票或后出现锚点重复主导统计。",
        ],
        0.72,
        2.95,
        11.8,
        2.75,
        size=14,
    )
    _add_footer(slide, len(prs.slides))

    union = summary.loc[summary["scope"].eq("UNION_UNIQUE")].iloc[0]
    both_rising = int(((unique_events["ma60_streak_at_position"] > 0) & (unique_events["ma120_streak_at_position"] > 0)).sum())
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, "29只主样本：第一次回调时的均线状态")
    cards = [
        (f"{union['ma30_streak_all_median']:.0f}日", f"MA30连续上涨中位数\nIQR {union['ma30_streak_all_p25']:.0f}–{union['ma30_streak_all_p75']:.0f}", ORANGE),
        (f"{union['ma60_streak_all_median']:.0f}日", f"MA60连续上涨中位数\nIQR {union['ma60_streak_all_p25']:.0f}–{union['ma60_streak_all_p75']:.0f}", "7B2CBF"),
        (f"{union['ma120_streak_all_median']:.0f}日", f"MA120连续上涨中位数\nIQR {union['ma120_streak_all_p25']:.0f}–{union['ma120_streak_all_p75']:.0f}", TEAL),
        (f"{union['ma30_vs_ma60_pct_median']:+.1f}%", f"MA30相对MA60间距中位数\nIQR {union['ma30_vs_ma60_pct_p25']:+.1f}%–{union['ma30_vs_ma60_pct_p75']:+.1f}%", BLUE),
    ]
    for index, (value, label, color) in enumerate(cards):
        _add_card(slide, 0.56 + index * 3.13, 1.18, 2.73, 1.48, value, label, accent=color)
    _add_card(slide, 0.88, 3.12, 3.55, 1.36, f"{both_rising}/29", "回调日MA60与MA120同时仍在严格上涨", accent=BLUE)
    _add_card(slide, 4.90, 3.12, 3.55, 1.36, f"{int(union['bull_alignment_active_count'])}/29", "回调日处于MA30>MA60>MA120多头排列", accent=TEAL)
    _add_card(slide, 8.90, 3.12, 3.55, 1.36, f"{int(union['all_three_rising_bull_active_count'])}/29", "回调日三条均线上涨且多头排列", accent=ORANGE)
    _add_text(slide, f"多头排列活动样本的连续时长中位数：{union['bull_alignment_active_median']:.0f}日（IQR {union['bull_alignment_active_p25']:.0f}–{union['bull_alignment_active_p75']:.0f}日）", 0.92, 5.16, 11.5, 0.48, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    _add_text(slide, "含0口径：回调日该均线未上涨或未形成多头排列时记0；避免只看活动样本造成高估。", 1.0, 5.90, 11.3, 0.42, size=10.5, color=MUTED, align=PP_ALIGN.CENTER)
    _add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, WHITE)
    _add_header(slide, "分布：回调发生在不同长趋势阶段，不能用单一“第N天”概括")
    slide.shapes.add_picture(str(distribution_chart), Inches(0.45), Inches(1.02), width=Inches(12.42), height=Inches(5.75))
    _add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, "为何35个锚点事件最终按29只股票汇总")
    main_codes = set(unique_events["ts_code"].astype(str))
    anchor_only = anchor_events.loc[~anchor_events["ts_code"].astype(str).isin(main_codes)]
    duplicate_codes = anchor_events.loc[anchor_events["ts_code"].duplicated(False), ["stock_name", "ts_code", "anchor_period", "pullback_date"]]
    _add_text(slide, "4只仅在某一后出现锚点内通过，但股票跨锚点最早首次回调并未粘合，因此不进入联合主样本：", 0.72, 1.15, 11.8, 0.42, size=13, bold=True)
    _add_text(slide, "、".join(f"{row.stock_name}（{row.ts_code}）" for row in anchor_only.itertuples()), 0.92, 1.75, 11.2, 0.58, size=15, color=RED, bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "锚点级重复：", 0.72, 2.65, 2.0, 0.36, size=13, bold=True)
    duplicate_lines = []
    for code, group in duplicate_codes.groupby("ts_code", sort=True):
        first = group.iloc[0]
        details = "；".join(f"MA{int(row.anchor_period)} {_fmt_date(row.pullback_date)}" for row in group.itertuples())
        duplicate_lines.append(f"{first['stock_name']}（{code}）：{details}")
    _add_bullets(slide, duplicate_lines, 0.9, 3.12, 11.4, 1.45, size=14)
    _add_text(slide, "PPT后续29张个股页均采用联合去重主样本；35个锚点级明细仍完整保存在CSV。", 0.84, 5.15, 11.6, 0.62, size=13, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    _add_footer(slide, len(prs.slides))

    page_size = 10
    pages = [unique_events.iloc[start : start + page_size] for start in range(0, len(unique_events), page_size)]
    for page_no, page in enumerate(pages, start=1):
        _add_event_table_slide(prs, page, page_no, len(pages))

    for _, event in unique_events.iterrows():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_background(slide, LIGHT)
        title = f"{int(event['display_order']):02d}  {event['stock_name']}  {event['ts_code']}  ·  MA{int(event['anchor_period'])}锚点"
        _add_header(slide, title, f"首次回调 {_fmt_date(event['pullback_date'])}")
        key = str(event["event_key"])
        slide.shapes.add_picture(str(chart_paths[key]), Inches(0.38), Inches(1.02), width=Inches(9.38), height=Inches(5.95))
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.93), Inches(1.04), Inches(3.0), Inches(5.90))
        panel.fill.solid(); panel.fill.fore_color.rgb = _rgb(WHITE); panel.line.color.rgb = _rgb(MID)
        _add_text(slide, "事件指标", 10.17, 1.24, 2.5, 0.35, size=15, bold=True, color=NAVY)
        metrics = [
            ("锚点区间", f"{_fmt_date(event['anchor_start_date'])} → {_fmt_date(event['anchor_end_date'])}"),
            ("锚点运行至回调", f"{int(event['anchor_observed_days_at_pullback'])}个交易日"),
            ("三线粘合离差", f"{float(event['ma5_ma10_ma30_spread_pct']):.2f}%"),
            ("低点 / 收盘 vs MA30", f"{float(event['low_vs_ma30_pct']):+.2f}% / {float(event['close_vs_ma30_pct']):+.2f}%"),
            ("MA30连续上涨", f"{int(event['ma30_streak_at_position'])}日"),
            ("MA60连续上涨", f"{int(event['ma60_streak_at_position'])}日"),
            ("MA120连续上涨", f"{int(event['ma120_streak_at_position'])}日"),
            ("30>60>120连续排列", f"{int(event['ma30_ma60_ma120_bull_streak_at_position'])}日"),
            ("MA30—MA60间距", f"{float(event['ma30_vs_ma60_pct']):+.2f}%"),
            ("主题", str(event.get("primary_theme", ""))),
        ]
        y = 1.76
        for label, value in metrics:
            _add_text(slide, label, 10.18, y, 2.5, 0.20, size=8.4, color=MUTED)
            _add_text(slide, value, 10.18, y + 0.20, 2.5, 0.31, size=11.2, bold=True, color=DARK_TEXT)
            y += 0.50
        _add_footer(slide, len(prs.slides))

    output.parent.mkdir(parents=True, exist_ok=True)
    prs.save(output)


def write_report(
    *,
    unique_events: pd.DataFrame,
    anchor_events: pd.DataFrame,
    summary: pd.DataFrame,
    coverage: pd.DataFrame,
    output: Path,
) -> None:
    union = summary.loc[summary["scope"].eq("UNION_UNIQUE")].iloc[0]
    lines = [
        "# MA60/MA120上升阶段第一次MA30回调与短均线粘合",
        "",
        "## 核心结果",
        "",
        f"- 231个可搜索长均线锚点全部出现首次MA30回调；其中35个锚点事件在第一次回调当日即形成MA5/MA10/MA30≤3%粘合，涉及{anchor_events['ts_code'].nunique()}只代码。",
        f"- 联合股票口径先跨MA60/MA120锚点取最早首次回调，再检查粘合，最终为{len(unique_events)}只股票。",
        f"- 联合样本回调日MA30、MA60、MA120连续上涨含0中位数分别为{union['ma30_streak_all_median']:.1f}、{union['ma60_streak_all_median']:.1f}、{union['ma120_streak_all_median']:.1f}个交易日。",
        f"- MA30>MA60>MA120连续排列在{int(union['bull_alignment_active_count'])}/{len(unique_events)}只上成立；活动样本中位数{union['bull_alignment_active_median']:.1f}日。",
        f"- MA30相对MA60间距中位数{union['ma30_vs_ma60_pct_median']:+.2f}%，IQR {union['ma30_vs_ma60_pct_p25']:+.2f}%—{union['ma30_vs_ma60_pct_p75']:+.2f}%。",
        "",
        "## 口径",
        "",
        "- 前复权价格，原始成交量；只按实际交易日序列。",
        "- 第一次回调先固定：此前10日最高收盘相对自身MA30至少+5%，当日最低价不高于MA30上方1%，收盘不低于MA30下方2%。",
        "- 再检查固定日MA5/MA10/MA30离差≤3%；首次失败不得后跳。",
        "- 连续上涨为MA_t>MA_(t-1)；多头排列为每天MA30>MA60>MA120。",
        "- 股票池按截至样本日事后观察到的超长MA60/MA120上涨段筛选，存在条件选择偏差；结果不是买入信号。",
        "",
        "## 锚点组覆盖",
        "",
        "| 锚点 | 可搜索 | 首回调当日粘合 | 覆盖率 |",
        "|---|---:|---:|---:|",
    ]
    for row in coverage.itertuples():
        lines.append(
            f"| MA{int(row.anchor_period)} | {int(row.eligible_count)} | {int(row.first_pullback_cohesion_count)} | {float(row.first_pullback_cohesion_pct_of_eligible):.1f}% |"
        )
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


def convert_ppt_to_pdf(pptx_path: Path, output_dir: Path) -> Path:
    soffice = shutil.which("soffice")
    if not soffice:
        raise RuntimeError("LibreOffice soffice is not available")
    fontconfig = output_dir / "fontconfig-noto-cjk.xml"
    cache_dir = output_dir / ".font-cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    fontconfig.write_text(
        "<?xml version='1.0'?>\n"
        "<!DOCTYPE fontconfig SYSTEM 'fonts.dtd'>\n"
        "<fontconfig>\n"
        "  <dir>/System/Library/Fonts</dir>\n"
        "  <dir>/Library/Fonts</dir>\n"
        f"  <dir>{Path.home() / 'Library' / 'Fonts'}</dir>\n"
        f"  <cachedir>{cache_dir}</cachedir>\n"
        "</fontconfig>\n",
        encoding="utf-8",
    )
    pdf_path = output_dir / f"{pptx_path.stem}.pdf"
    if pdf_path.exists():
        pdf_path.unlink()
    profile_dir = output_dir / ".libreoffice-profile"
    if profile_dir.exists():
        shutil.rmtree(profile_dir)
    env = dict(**__import__("os").environ)
    env["FONTCONFIG_FILE"] = str(fontconfig.resolve())
    command = [
        soffice,
        "--headless",
        f"-env:UserInstallation={profile_dir.resolve().as_uri()}",
        "--convert-to",
        "pdf",
        "--outdir",
        str(output_dir.resolve()),
        str(pptx_path.resolve()),
    ]
    result = subprocess.run(command, check=False, capture_output=True, text=True, env=env, timeout=240)
    if result.returncode != 0 or not pdf_path.exists():
        raise RuntimeError(f"PDF conversion failed: {result.stdout}\n{result.stderr}")
    return pdf_path


def render_pdf(pdf_path: Path, output_dir: Path) -> list[Path]:
    import pymupdf

    output_dir.mkdir(parents=True, exist_ok=True)
    for stale in output_dir.glob("slide-*.png"):
        stale.unlink()
    document = pymupdf.open(pdf_path)
    paths: list[Path] = []
    matrix = pymupdf.Matrix(1.35, 1.35)
    for index, page in enumerate(document, start=1):
        path = output_dir / f"slide-{index:02d}.png"
        page.get_pixmap(matrix=matrix, alpha=False).save(path)
        paths.append(path)
    document.close()
    return paths


def create_contact_sheets(slide_images: list[Path], output_dir: Path, *, per_sheet: int = 12) -> list[Path]:
    from PIL import Image, ImageDraw

    output_dir.mkdir(parents=True, exist_ok=True)
    for stale in output_dir.glob("contact-*.jpg"):
        stale.unlink()
    sheets: list[Path] = []
    for batch_no, start in enumerate(range(0, len(slide_images), per_sheet), start=1):
        batch = slide_images[start : start + per_sheet]
        thumb_w, thumb_h = 520, 293
        cols, rows = 3, 4
        canvas = Image.new("RGB", (cols * thumb_w, rows * (thumb_h + 26)), "white")
        draw = ImageDraw.Draw(canvas)
        for index, path in enumerate(batch):
            image = Image.open(path).convert("RGB")
            image.thumbnail((thumb_w - 8, thumb_h - 8))
            col, row = index % cols, index // cols
            x = col * thumb_w + (thumb_w - image.width) // 2
            y = row * (thumb_h + 26) + 4
            canvas.paste(image, (x, y))
            draw.text((col * thumb_w + 8, row * (thumb_h + 26) + thumb_h + 3), f"Slide {start + index + 1}", fill="black")
        path = output_dir / f"contact-{batch_no:02d}.jpg"
        canvas.save(path, quality=90)
        sheets.append(path)
    return sheets


def main(argv: Sequence[str] | None = None) -> int:
    args = _parse_args(argv)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    charts_dir = output_dir / "charts"
    charts_dir.mkdir(parents=True, exist_ok=True)
    for stale in charts_dir.glob("*.png"):
        stale.unlink()
    _register_fonts()

    daily = pd.read_csv(args.daily_file, compression="infer", dtype={"ts_code": str, "trade_date": str})
    frames_by_code = prepare_stock_frames(daily)
    date_columns = ["anchor_start_date", "anchor_end_date", "pullback_date"]
    anchor_events = pd.read_csv(args.event_file, dtype={column: str for column in ["ts_code", *date_columns]})
    unique_events = pd.read_csv(args.unique_event_file, dtype={column: str for column in ["ts_code", *date_columns]})
    all_first = pd.read_csv(args.all_first_file, dtype={column: str for column in ["ts_code", *date_columns]})
    coverage = pd.read_csv(args.coverage_file)
    for frame in (anchor_events, unique_events, all_first):
        for column in date_columns:
            if column in frame:
                frame[column] = frame[column].map(_compact_date)

    anchor_events = attach_event_metrics(anchor_events, frames_by_code)
    unique_events = attach_event_metrics(unique_events, frames_by_code)
    all_first = attach_event_metrics(all_first, frames_by_code)
    anchor_events["scope"] = anchor_events["anchor_period"].map(lambda value: f"MA{int(value)}")
    unique_events["scope"] = "UNION_UNIQUE"
    combined = pd.concat([anchor_events, unique_events], ignore_index=True)
    summary = summarize_extended_metrics(combined)

    if len(anchor_events) != 35 or anchor_events["ts_code"].nunique() != 33:
        raise ValueError("anchor event count changed")
    if len(unique_events) != 29 or not unique_events["ts_code"].is_unique:
        raise ValueError("unique-stock event count changed")
    if len(all_first) != 231:
        raise ValueError("first-pullback coverage count changed")
    if set(unique_events["ts_code"]) - set(anchor_events["ts_code"]):
        raise ValueError("unique events are not a subset of anchor events")

    unique_events = unique_events.sort_values(["pullback_date", "ts_code"], kind="mergesort").reset_index(drop=True)
    unique_events.insert(0, "display_order", np.arange(1, len(unique_events) + 1))
    unique_events["event_key"] = unique_events.apply(
        lambda row: f"{row['ts_code']}|{int(row['anchor_period'])}|{row['anchor_start_date']}|{row['pullback_date']}",
        axis=1,
    )
    chart_paths: dict[str, Path] = {}
    for row in unique_events.itertuples(index=False):
        path = charts_dir / f"{int(row.display_order):02d}_{str(row.ts_code).replace('.', '_')}_{row.pullback_date}.png"
        create_event_chart(frames_by_code[str(row.ts_code)], pd.Series(row._asdict()), path)
        chart_paths[str(row.event_key)] = path
    if len(chart_paths) != len(unique_events) or len({_sha256(path) for path in chart_paths.values()}) != len(unique_events):
        raise ValueError("event chart count/hash uniqueness failed")

    distribution_chart = output_dir / "event_metric_distributions.png"
    create_distribution_chart(unique_events, distribution_chart)
    anchor_csv = output_dir / "anchor_event_metrics_35.csv"
    unique_csv = output_dir / "unique_stock_event_metrics_29.csv"
    summary_csv = output_dir / "extended_metric_summary.csv"
    anchor_events.to_csv(anchor_csv, index=False)
    unique_events.to_csv(unique_csv, index=False)
    summary.to_csv(summary_csv, index=False)
    report_path = output_dir / "report.md"
    write_report(
        unique_events=unique_events,
        anchor_events=anchor_events,
        summary=summary,
        coverage=coverage,
        output=report_path,
    )

    pptx_path = output_dir / "MA60_MA120初期首次回调MA30与三线粘合研究.pptx"
    build_deck(
        unique_events=unique_events,
        anchor_events=anchor_events,
        summary=summary,
        coverage=coverage,
        chart_paths=chart_paths,
        distribution_chart=distribution_chart,
        output=pptx_path,
    )
    pdf_path = convert_ppt_to_pdf(pptx_path, output_dir)
    rendered = render_pdf(pdf_path, output_dir / "rendered_slides")
    contacts = create_contact_sheets(rendered, output_dir / "contact_sheets")

    readback = Presentation(pptx_path)
    expected_slides = 6 + math.ceil(len(unique_events) / 10) + len(unique_events)
    if len(readback.slides) != expected_slides:
        raise ValueError(f"slide count mismatch: {len(readback.slides)} != {expected_slides}")
    if len(rendered) != expected_slides:
        raise ValueError(f"rendered page count mismatch: {len(rendered)} != {expected_slides}")
    detail_start = 6 + math.ceil(len(unique_events) / 10)
    detail_slides = list(readback.slides)[detail_start:]
    if len(detail_slides) != len(unique_events):
        raise ValueError("detail slide count mismatch")
    picture_counts = [sum(1 for shape in slide.shapes if shape.shape_type == 13) for slide in detail_slides]
    if picture_counts != [1] * len(unique_events):
        raise ValueError(f"detail picture count mismatch: {picture_counts}")

    artifacts = [
        anchor_csv,
        unique_csv,
        summary_csv,
        report_path,
        distribution_chart,
        pptx_path,
        pdf_path,
        *chart_paths.values(),
        *contacts,
    ]
    metadata = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "study": "first_ma30_pullback_cohesion_extended_ma_state_ppt",
        "data_contract": {
            "price_adjustment": "qfq=True",
            "volume": "raw volume",
            "trading_days_only": True,
            "cutoff": "2026-09-30",
            "daily_input": str(Path(args.daily_file).resolve()),
            "source_contract": "formal cache built via coreClient.data_provider with source=database_only",
        },
        "definitions": {
            "early_pullback": "earliest MA30 pullback inside each selected non-left-censored MA60/MA120 strict-rising anchor; no separate elapsed-day cap",
            "pullback": "prior 10-trading-day max(close/own-day-MA30-1)>=5%; event low/MA30-1<=1%; event close/MA30-1>=-2%",
            "cohesion": "on the already-fixed first pullback, (max(MA5,MA10,MA30)-min)/mean<=3%; no skipping",
            "ma_rising_streak": "MA_t>MA_(t-1); flat/fall/unavailable resets to zero",
            "bull_alignment_streak": "consecutive trading days through event with MA30>MA60>MA120",
            "all_three_rising_bull_streak": "consecutive trading days through event with MA30>MA60>MA120 and every MA strictly rising",
            "ma30_ma60_gap_pct": "100*(MA30/MA60-1)",
            "unique_stock_rule": "select each stock's earliest first-pullback across MA60/MA120 anchors before applying cohesion",
        },
        "counts": {
            "candidate_stocks": int(daily["ts_code"].nunique()),
            "eligible_anchors": int(coverage["eligible_count"].sum()),
            "all_first_pullbacks": int(len(all_first)),
            "anchor_events": int(len(anchor_events)),
            "anchor_event_distinct_stocks": int(anchor_events["ts_code"].nunique()),
            "unique_stock_main_events": int(len(unique_events)),
            "event_charts": int(len(chart_paths)),
            "slides": int(len(readback.slides)),
            "pdf_pages": int(len(rendered)),
        },
        "selection_bias_warning": "The 178-stock pool was selected by future-observed long MA duration; results are descriptive and are not a live-strategy backtest or buy signal.",
        "artifacts": {
            path.name: {"path": str(path.resolve()), "sha256": _sha256(path), "bytes": path.stat().st_size}
            for path in artifacts
        },
    }
    metadata_path = output_dir / "metadata.json"
    metadata_path.write_text(json.dumps(_clean_json(metadata), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(
        json.dumps(
            {
                "output_dir": str(output_dir.resolve()),
                "anchor_events": len(anchor_events),
                "anchor_event_distinct_stocks": int(anchor_events["ts_code"].nunique()),
                "unique_stock_main_events": len(unique_events),
                "slides": len(readback.slides),
                "pdf_pages": len(rendered),
                "pptx": str(pptx_path.resolve()),
                "pdf": str(pdf_path.resolve()),
                "contact_sheets": [str(path.resolve()) for path in contacts],
            },
            ensure_ascii=False,
        )
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())


"""Build a combined report and slide deck for three volume questions."""

from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
from typing import Any, Sequence

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_ma60_top1pct_kline_ppt import (
    _add_rect,
    _add_text,
    _set_slide_background,
)


ROOT = Path(__file__).resolve().parents[5]
DEFAULT_VOLUME_DIR = (
    ROOT
    / "outputs"
    / "volume_expansion_in_long_ma_episodes"
    / "2026_through_20260930_qfq"
)
DEFAULT_DIVERGENCE_DIR = (
    ROOT
    / "outputs"
    / "volume_price_divergence_top_signal"
    / "2026_through_20260930_qfq"
)
DEFAULT_DAILY = (
    ROOT
    / "outputs"
    / "ma60_ma120_top100_ppt"
    / "2026_through_20260930_qfq"
    / "daily_input_qfq.csv.gz"
)
DEFAULT_OUTPUT_DIR = (
    ROOT
    / "outputs"
    / "volume_three_questions_report"
    / "2026_through_20260930_qfq"
)

NAVY = "#0B1220"
INK = "#132238"
MUTED = "#5C6B7A"
PAPER = "#F5F7FA"
WHITE = "#FFFFFF"
BLUE = "#2563EB"
PURPLE = "#7C3AED"
TEAL = "#0F9D8B"
AMBER = "#F59E0B"
RED = "#DC4C4C"
GREEN = "#159D72"
GRID = "#DCE3EA"
LIGHT_BLUE = "#EAF2FF"
LIGHT_PURPLE = "#F1ECFF"
LIGHT_TEAL = "#E8F7F4"
LIGHT_AMBER = "#FFF4D9"
LIGHT_RED = "#FDECEC"
FONT = "Noto Sans CJK SC"


def _one(frame: pd.DataFrame, **filters: Any) -> pd.Series:
    mask = pd.Series(True, index=frame.index)
    for column, value in filters.items():
        mask &= frame[column].eq(value)
    selected = frame.loc[mask]
    if len(selected) != 1:
        raise ValueError(f"expected one row for {filters}, found {len(selected)}")
    return selected.iloc[0]


def derive_key_findings(
    *,
    long_summary: pd.DataFrame,
    segment_summary: pd.DataFrame,
    anchor_summary: pd.DataFrame,
    effectiveness: pd.DataFrame,
    lifts: pd.DataFrame,
    coverage: pd.DataFrame,
    within: pd.DataFrame,
) -> dict[str, Any]:
    """Extract the exact metrics used in the report without mixing questions."""
    q1: dict[int, dict[str, Any]] = {}
    for period in (60, 120):
        equal = _one(
            long_summary,
            anchor_period=period,
            cohort="all_observed",
            baseline="equal_length",
        )
        pre20 = _one(
            long_summary,
            anchor_period=period,
            cohort="all_observed",
            baseline="pre20",
        )
        q1[period] = {
            "comparable_count": int(equal["comparable_count"]),
            "equal_length_mean_ratio": float(equal["ratio_mean"]),
            "equal_length_median_ratio": float(equal["ratio_median"]),
            "equal_length_p25": float(equal["ratio_p25"]),
            "equal_length_p75": float(equal["ratio_p75"]),
            "amplified_gt_1_count": int(equal["amplified_gt_1_count"]),
            "amplified_gt_1_pct": float(equal["amplified_gt_1_pct"]),
            "amplified_ge_1_2_pct": float(equal["amplified_ge_1_2_pct"]),
            "amplified_ge_1_5_pct": float(equal["amplified_ge_1_5_pct"]),
            "amplified_ge_2_pct": float(equal["amplified_ge_2_pct"]),
            "pre20_comparable_count": int(pre20["comparable_count"]),
            "pre20_median_ratio": float(pre20["ratio_median"]),
            "pre20_amplified_gt_1_pct": float(pre20["amplified_gt_1_pct"]),
        }

    q2: dict[int, dict[int, dict[str, Any]]] = {}
    for period in (60, 120):
        q2[period] = {}
        for days in (1, 5, 10):
            segment = _one(
                segment_summary,
                anchor_period=period,
                comparison="full_episode_equal_prior",
                min_segment_days=days,
            )
            versus_anchor = _one(
                segment_summary,
                anchor_period=period,
                comparison="clipped_vs_anchor_mean",
                min_segment_days=days,
            )
            anchor = _one(
                anchor_summary,
                anchor_period=period,
                comparison="full_episode_equal_prior",
                min_episode_days=days,
            )
            q2[period][days] = {
                "segment_count": int(segment["comparable_count"]),
                "segment_median_ratio": float(segment["ratio_median"]),
                "segment_amplified_count": int(segment["amplified_gt_1_count"]),
                "segment_amplified_pct": float(segment["amplified_gt_1_pct"]),
                "versus_anchor_mean_amplified_pct": float(
                    versus_anchor["amplified_gt_1_pct"]
                ),
                "evaluable_anchor_count": int(anchor["evaluable_anchor_count"]),
                "all_segments_amplified_count": int(
                    anchor["all_segments_amplified_gt_1_count"]
                ),
                "all_segments_amplified_pct": float(
                    anchor["all_segments_amplified_gt_1_pct"]
                ),
                "median_anchor_amplified_segment_pct": float(
                    anchor["median_anchor_amplified_segment_pct"]
                ),
            }

    q3: dict[int, dict[str, dict[str, Any]]] = {60: {}, 120: {}}
    for period in (60, 120):
        for scope in ("all_candidate_days", "cooldown_10d"):
            divergence = _one(
                effectiveness,
                anchor_period=period,
                scope=scope,
                event_class="divergence",
            )
            control = _one(
                effectiveness,
                anchor_period=period,
                scope=scope,
                event_class="non_divergent_new_high",
            )
            lift = _one(lifts, anchor_period=period, scope=scope)
            paired = within.loc[
                within["anchor_period"].eq(period) & within["scope"].eq(scope)
            ]
            q3[period][scope] = {
                "divergence_event_count": int(divergence["event_count"]),
                "control_event_count": int(control["event_count"]),
                "divergence_drawdown_rate": float(
                    divergence["drawdown_ge_10pct_rate"]
                ),
                "control_drawdown_rate": float(control["drawdown_ge_10pct_rate"]),
                "divergence_top_like_rate": float(
                    divergence["top_like_forward_rate"]
                ),
                "control_top_like_rate": float(control["top_like_forward_rate"]),
                "divergence_near_top_rate": float(
                    divergence["near_anchor_peak_10d_rate"]
                ),
                "control_near_top_rate": float(
                    control["near_anchor_peak_10d_rate"]
                ),
                "divergence_median_forward_return": float(
                    divergence["median_forward_return_pct"]
                ),
                "control_median_forward_return": float(
                    control["median_forward_return_pct"]
                ),
                "drawdown_rate_lift": float(
                    lift["drawdown_ge_10pct_rate_lift"]
                ),
                "top_like_rate_lift": float(lift["top_like_forward_rate_lift"]),
                "near_top_rate_lift": float(
                    lift["near_anchor_peak_10d_rate_lift"]
                ),
                "paired_anchor_count": (
                    int(paired.iloc[0]["paired_anchor_count"])
                    if len(paired) == 1 else 0
                ),
                "paired_higher_count": (
                    int(paired.iloc[0]["divergence_rate_higher_count"])
                    if len(paired) == 1 else 0
                ),
                "paired_lower_count": (
                    int(paired.iloc[0]["divergence_rate_lower_count"])
                    if len(paired) == 1 else 0
                ),
                "paired_median_drawdown_rate_difference": (
                    float(paired.iloc[0]["median_drawdown_rate_difference_pct"])
                    if len(paired) == 1 else float("nan")
                ),
                "paired_median_return_difference": (
                    float(paired.iloc[0]["median_forward_return_difference_pct"])
                    if len(paired) == 1 else float("nan")
                ),
            }

    primary_coverage = _one(
        coverage, anchor_period=60, scope="all_candidate_days"
    )
    return {
        "q1": q1,
        "q2": q2,
        "q3": q3,
        "q3_top_coverage": {
            "completed_anchor_count": int(primary_coverage["completed_anchor_count"]),
            "anchors_with_any_divergence": int(
                primary_coverage["anchors_with_any_divergence"]
            ),
            "anchors_with_any_divergence_pct": float(
                primary_coverage["anchors_with_any_divergence_pct"]
            ),
            "anchors_with_divergence_within_10d_of_top": int(
                primary_coverage[
                    "anchors_with_divergence_within_10d_of_top"
                ]
            ),
            "top_coverage_rate_pct": float(
                primary_coverage["top_coverage_rate_pct"]
            ),
            "near_top_signal_lead_median_days": float(
                primary_coverage["near_top_signal_lead_median_days"]
            ),
        },
    }


def extract_equal_length_volume_ratios(
    detail: pd.DataFrame,
    anchor_period: int,
) -> pd.Series:
    """Return comparable equal-prior ratios, excluding left-censored anchors."""
    required = {
        "anchor_period", "anchor_left_censored", "equal_prior_volume_ratio"
    }
    missing = required.difference(detail.columns)
    if missing:
        raise ValueError(f"missing long-detail columns: {sorted(missing)}")
    left_censored = detail["anchor_left_censored"].map(
        lambda value: str(value).strip().lower() in {"true", "1", "yes"}
    )
    values = pd.to_numeric(
        detail.loc[
            detail["anchor_period"].eq(anchor_period) & ~left_censored,
            "equal_prior_volume_ratio",
        ],
        errors="coerce",
    )
    return values.dropna().reset_index(drop=True)


def select_cooldown_sensitivity(sensitivity: pd.DataFrame) -> pd.DataFrame:
    """Select one aggregate sensitivity row per spec and anchor period."""
    required = {"spec_name", "anchor_period", "scope", "drawdown_rate_lift"}
    missing = required.difference(sensitivity.columns)
    if missing:
        raise ValueError(f"missing sensitivity columns: {sorted(missing)}")
    return sensitivity.loc[
        sensitivity["scope"].eq("cooldown_10d")
    ].copy().reset_index(drop=True)


def humanize_divergence_spec(spec_name: str) -> str:
    """Translate internal sensitivity spec codes into slide-ready labels."""
    labels = {
        "close_l60_v5_t0.7": "缩量≥30%",
        "close_l60_v5_t0.8": "缩量≥20%",
        "close_l60_v5_t0.9": "缩量≥10%",
        "close_l40_v5_t0.8": "回看40日",
        "close_l120_v5_t0.8": "回看120日",
        "close_l60_v10_t0.8": "10日均量",
        "high_l60_v5_t0.8": "最高价新高",
    }
    return labels.get(str(spec_name), str(spec_name))


def _fmt(value: float, digits: int = 1) -> str:
    if value is None or not math.isfinite(float(value)):
        return "不可判定"
    return f"{float(value):.{digits}f}"


def _add_header(slide, section: str, title: str, number: int) -> None:
    _add_text(slide, section, 0.55, 0.28, 3.6, 0.28, size=10, color=TEAL, bold=True)
    _add_text(slide, title, 0.55, 0.64, 12.15, 0.58, size=26, color=INK, bold=True)
    _add_rect(slide, 0.55, 1.29, 0.78, 0.055, fill=TEAL)
    _add_text(
        slide,
        f"2026 YTD · 截至 2026-09-30 · 前复权均线 / 原始成交量                              {number:02d}",
        0.55,
        7.18,
        12.2,
        0.18,
        size=7.8,
        color="#7A8795",
        align=PP_ALIGN.RIGHT,
        margin=0,
    )


def _metric_card(
    slide,
    *,
    left: float,
    top: float,
    width: float,
    value: str,
    label: str,
    note: str,
    color: str,
    fill: str,
) -> None:
    _add_rect(slide, left, top, width, 1.42, fill=fill, radius=True, line=color)
    _add_text(slide, value, left + 0.18, top + 0.15, width - 0.36, 0.48, size=25, color=color, bold=True)
    _add_text(slide, label, left + 0.18, top + 0.67, width - 0.36, 0.26, size=11.5, color=INK, bold=True)
    _add_text(slide, note, left + 0.18, top + 0.99, width - 0.36, 0.28, size=9.2, color=MUTED)


def _bullet_block(
    slide,
    bullets: Sequence[str],
    *,
    left: float,
    top: float,
    width: float,
    height: float,
    size: float = 14.0,
    color: str = INK,
) -> None:
    line_height = height / max(len(bullets), 1)
    for index, bullet in enumerate(bullets):
        _add_rect(slide, left, top + index * line_height + 0.13, 0.08, 0.08, fill=TEAL, radius=True)
        _add_text(
            slide,
            bullet,
            left + 0.19,
            top + index * line_height,
            width - 0.19,
            line_height - 0.02,
            size=size,
            color=color,
            valign=MSO_ANCHOR.MIDDLE,
        )


def _add_chart(slide, path: Path, left: float, top: float, width: float, height: float) -> None:
    if not path.exists():
        raise FileNotFoundError(path)
    slide.shapes.add_picture(str(path), Inches(left), Inches(top), Inches(width), Inches(height))


def build_presentation(
    findings: dict[str, Any],
    chart_paths: dict[str, Path],
    output: Path,
) -> dict[str, int]:
    """Build the fixed 15-slide research deck."""
    expected_charts = {
        "q1_thresholds", "q1_distribution", "q2_segments", "q2_all_segments",
        "q2_benchmarks", "q3_forward", "q3_paired", "q3_sensitivity",
        "q3_examples",
    }
    missing = expected_charts.difference(chart_paths)
    if missing:
        raise ValueError(f"missing chart paths: {sorted(missing)}")

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]

    # 1 — title
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, NAVY)
    _add_rect(slide, 0.7, 0.72, 1.05, 0.08, fill=TEAL)
    _add_text(slide, "MA60 / MA120 长趋势成交量研究", 0.7, 1.25, 11.7, 0.62, size=30, color=WHITE, bold=True)
    _add_text(slide, "三个成交量问题", 0.7, 2.05, 11.7, 0.75, size=42, color="#8BE1D2", bold=True)
    _add_text(
        slide,
        "阶段整体是否放量 · MA20各段是否都放量 · 量价背离能否提示见顶",
        0.7,
        3.04,
        11.8,
        0.55,
        size=19,
        color="#CFD8E3",
    )
    _add_text(slide, "样本：MA60榜107只 / MA120榜126只 · 截至2026-09-30", 0.7, 5.72, 11.4, 0.34, size=12, color="#A7B4C4")
    _add_text(slide, "正式研究报告 · 2026 YTD", 0.7, 6.25, 11.4, 0.3, size=11, color="#708095")

    # 2 — executive summary
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, PAPER)
    _add_header(slide, "EXECUTIVE SUMMARY", "结论先行：一强、一半、一否", 2)
    q1_60, q1_120 = findings["q1"][60], findings["q1"][120]
    _metric_card(slide, left=0.62, top=1.68, width=3.85, value="是，显著放量", label="问题一 · 长均线阶段整体", note=f"中位倍率 MA60 {q1_60['equal_length_median_ratio']:.2f}× / MA120 {q1_120['equal_length_median_ratio']:.2f}×", color=BLUE, fill=LIGHT_BLUE)
    _metric_card(slide, left=4.74, top=1.68, width=3.85, value="多数段放量，但绝非每段", label="问题二 · 每个MA20上升段", note=f"逐段放量 {findings['q2'][60][1]['segment_amplified_pct']:.1f}% / {findings['q2'][120][1]['segment_amplified_pct']:.1f}%", color=PURPLE, fill=LIGHT_PURPLE)
    _metric_card(slide, left=8.86, top=1.68, width=3.85, value="否，单独使用无效", label="问题三 · 量价背离见顶", note="未来20日深回撤率未优于非背离新高日", color=RED, fill=LIGHT_RED)
    _bullet_block(
        slide,
        [
            "长趋势形成后，成交量的抬升是普遍现象；结论对等长、20/60/120日前置窗口均稳定。",
            "MA20上升段更常放量，但“同一只股票每段都放量”的严格命题大多不成立，尤其MA120阶段。",
            "量价背离出现得不算少，但精确性与后续回撤概率均未优于同类非背离新高，不能作为独立卖点。",
        ],
        left=0.8,
        top=3.55,
        width=11.9,
        height=2.55,
        size=14.1,
    )

    # 3 — scope and definitions
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, WHITE)
    _add_header(slide, "RESEARCH DESIGN", "样本、口径与可判定边界", 3)
    _metric_card(slide, left=0.62, top=1.63, width=2.8, value="107", label="MA60正式样本", note="101已结束 / 6右截断", color=BLUE, fill=LIGHT_BLUE)
    _metric_card(slide, left=3.62, top=1.63, width=2.8, value="126", label="MA120正式样本", note="全部截至样本日仍延续", color=PURPLE, fill=LIGHT_PURPLE)
    _metric_card(slide, left=6.62, top=1.63, width=2.8, value="QFQ", label="价格与均线", note="前复权日线；成交量不复权", color=TEAL, fill=LIGHT_TEAL)
    _metric_card(slide, left=9.62, top=1.63, width=2.8, value="实际交易日", label="窗口与停牌", note="不按自然日补零", color=AMBER, fill=LIGHT_AMBER)
    _bullet_block(
        slide,
        [
            "范围：带入锚点按完整已观察区间分析（可含2025）；问题一主基准为episode均量 ÷ 开始前等长交易日均量，并做20/60/120日敏感性。",
            "问题二主基准：完整MA20上升episode均量 ÷ 其开始前等长交易日均量；≥1/≥5/≥10日重新筛选。",
            "问题三主信号：收盘价创此前60日新高，同时当前5日均量 ≤ 前一价格峰值日5日均量的80%；不使用未来数据。",
            "MA120最终顶部不可观察，因此只评价未来20日表现；最终顶部覆盖率仅对已结束MA60阶段统计。",
        ],
        left=0.82,
        top=3.45,
        width=11.7,
        height=2.85,
        size=13.2,
    )

    # 4 — Q1 headline
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, PAPER)
    _add_header(slide, "QUESTION 1", "长均线连续上升阶段，整体成交量显著高于阶段前", 4)
    _metric_card(slide, left=0.62, top=1.62, width=2.7, value=f"{q1_60['equal_length_median_ratio']:.2f}×", label="MA60 中位倍率", note=f"{q1_60['amplified_gt_1_pct']:.1f}% 高于阶段前", color=BLUE, fill=LIGHT_BLUE)
    _metric_card(slide, left=3.52, top=1.62, width=2.7, value=f"{q1_120['equal_length_median_ratio']:.2f}×", label="MA120 中位倍率", note=f"{q1_120['amplified_gt_1_pct']:.1f}% 高于阶段前", color=PURPLE, fill=LIGHT_PURPLE)
    _add_chart(slide, chart_paths["q1_thresholds"], 6.45, 1.5, 6.25, 4.8)
    _bullet_block(
        slide,
        [
            f"MA60：{q1_60['amplified_ge_1_5_pct']:.1f}% 至少放大1.5倍，{q1_60['amplified_ge_2_pct']:.1f}% 至少2倍。",
            f"MA120：{q1_120['amplified_ge_1_5_pct']:.1f}% 至少放大1.5倍，{q1_120['amplified_ge_2_pct']:.1f}% 至少2倍。",
            "所以“整体平均成交量远大于阶段前”在该长趋势强样本中得到强支持。",
        ],
        left=0.72,
        top=3.47,
        width=5.5,
        height=2.45,
        size=13.4,
    )

    # 5 — Q1 distribution/sensitivity
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, WHITE)
    _add_header(slide, "QUESTION 1", "分布不是少数极端值驱动，窗口敏感性一致", 5)
    _add_chart(slide, chart_paths["q1_distribution"], 0.62, 1.5, 7.75, 4.98)
    _metric_card(slide, left=8.66, top=1.65, width=3.95, value=f"{q1_60['equal_length_p25']:.2f}–{q1_60['equal_length_p75']:.2f}×", label="MA60 四分位区间", note="股票自身前后倍率分布", color=BLUE, fill=LIGHT_BLUE)
    _metric_card(slide, left=8.66, top=3.3, width=3.95, value=f"{q1_120['equal_length_p25']:.2f}–{q1_120['equal_length_p75']:.2f}×", label="MA120 四分位区间", note="排除左截断后统计", color=PURPLE, fill=LIGHT_PURPLE)
    _add_text(slide, f"固定20日前置窗口：MA60中位 {q1_60['pre20_median_ratio']:.2f}×；MA120中位 {q1_120['pre20_median_ratio']:.2f}×。", 8.7, 5.15, 3.8, 0.72, size=12.2, color=INK, bold=True)
    _add_text(slide, "MA120等长前置窗口仅120只完整；固定窗口124只均可比较。", 8.7, 6.05, 3.8, 0.55, size=10.5, color=MUTED)

    # 6 — Q2 segment level
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, PAPER)
    _add_header(slide, "QUESTION 2", "逐个MA20上升段看：多数放量，且长段更明显", 6)
    _add_chart(slide, chart_paths["q2_segments"], 0.62, 1.46, 7.85, 5.1)
    q2_60_1, q2_120_1 = findings["q2"][60][1], findings["q2"][120][1]
    _metric_card(slide, left=8.75, top=1.62, width=3.8, value=f"{q2_60_1['segment_amplified_pct']:.1f}%", label="MA60内MA20段放量", note=f"{q2_60_1['segment_amplified_count']}/{q2_60_1['segment_count']}段", color=BLUE, fill=LIGHT_BLUE)
    _metric_card(slide, left=8.75, top=3.28, width=3.8, value=f"{q2_120_1['segment_amplified_pct']:.1f}%", label="MA120内MA20段放量", note=f"{q2_120_1['segment_amplified_count']}/{q2_120_1['segment_count']}段", color=PURPLE, fill=LIGHT_PURPLE)
    _add_text(slide, "≥10日段的放量率升至 MA60 78.4%、MA120 75.9%，短噪声段是主要拖累。", 8.78, 5.22, 3.72, 0.95, size=12.3, color=INK, bold=True)

    # 7 — Q2 every segment
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, WHITE)
    _add_header(slide, "QUESTION 2", "但“每一个MA20段都放量”是少数，不是规律", 7)
    _add_chart(slide, chart_paths["q2_all_segments"], 0.62, 1.5, 7.9, 5.0)
    _metric_card(slide, left=8.8, top=1.64, width=3.65, value=f"{q2_60_1['all_segments_amplified_pct']:.1f}%", label="MA60：所有段均放量", note=f"{q2_60_1['all_segments_amplified_count']}/{q2_60_1['evaluable_anchor_count']}只", color=BLUE, fill=LIGHT_BLUE)
    _metric_card(slide, left=8.8, top=3.3, width=3.65, value=f"{q2_120_1['all_segments_amplified_pct']:.1f}%", label="MA120：所有段均放量", note=f"{q2_120_1['all_segments_amplified_count']}/{q2_120_1['evaluable_anchor_count']}只", color=PURPLE, fill=LIGHT_PURPLE)
    _add_text(slide, f"典型股票中，放量段占比中位数为 MA60 {q2_60_1['median_anchor_amplified_segment_pct']:.1f}% / MA120 {q2_120_1['median_anchor_amplified_segment_pct']:.1f}%。", 8.82, 5.18, 3.58, 0.9, size=12.0, color=INK, bold=True)
    _add_text(slide, "正确表述：大多数段放量；错误表述：每段必然放量。", 8.82, 6.07, 3.58, 0.42, size=10.5, color=RED)

    # 8 — Q2 benchmark sensitivity
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, PAPER)
    _add_header(slide, "QUESTION 2", "基准不同，回答的是不同问题", 8)
    _add_chart(slide, chart_paths["q2_benchmarks"], 0.62, 1.5, 7.8, 5.0)
    _bullet_block(
        slide,
        [
            "相对该段开始前等长窗口：回答“进入MA20上升段时是否比之前更活跃”——约65%–69%为是。",
            "相对整个长均线阶段均量：回答“该段是否高于长趋势常态”——仅MA60 46.7%、MA120 43.2%。",
            "过滤≥5日、≥10日后放量率提高，但仍不能推出每一段都放量。",
            "这更像间歇式放量结构，而不是成交量沿MA20各段持续、单调扩张。",
        ],
        left=8.72,
        top=1.7,
        width=3.8,
        height=4.4,
        size=12.2,
    )

    # 9 — Q3 definition
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, WHITE)
    _add_header(slide, "QUESTION 3", "量价背离见顶：用无前视定义做事件检验", 9)
    _add_rect(slide, 0.7, 1.62, 11.95, 1.1, fill=LIGHT_RED, radius=True, line=RED)
    _add_text(slide, "背离事件 = 收盘价创此前60个交易日新高，且当前5日均量 ≤ 前一价格峰值日5日均量的80%", 0.95, 1.91, 11.45, 0.52, size=18, color=RED, bold=True, align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    _bullet_block(
        slide,
        [
            "对照组：同样创60日新高、但不满足缩量20%的交易日；避免拿普通日作不公平对照。",
            "主要结果：未来20日最大回撤是否达到10%；辅助结果：上涨空间≤5%且回撤≥10%的“顶部型”路径。",
            "相关事件按同类10交易日冷却去重做稳健性；另在同一个长趋势episode内做背离/非背离配对。",
            "最终顶部精确性只对101个已结束MA60 episode计算；MA120全部右截断，不作事后顶部判断。",
        ],
        left=0.95,
        top=3.2,
        width=11.35,
        height=2.95,
        size=14.0,
    )

    # 10 — Q3 forward outcomes
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, PAPER)
    _add_header(slide, "QUESTION 3", "背离后的20日深回撤率没有高于非背离新高日", 10)
    _add_chart(slide, chart_paths["q3_forward"], 0.62, 1.5, 8.0, 5.0)
    q3_60 = findings["q3"][60]["all_candidate_days"]
    q3_120 = findings["q3"][120]["all_candidate_days"]
    _metric_card(slide, left=8.9, top=1.63, width=3.5, value=f"{q3_60['drawdown_rate_lift']:.2f}×", label="MA60 深回撤率倍率", note="背离 ÷ 非背离；低于1", color=RED, fill=LIGHT_RED)
    _metric_card(slide, left=8.9, top=3.28, width=3.5, value=f"{q3_120['drawdown_rate_lift']:.2f}×", label="MA120 深回撤率倍率", note="背离 ÷ 非背离；低于1", color=RED, fill=LIGHT_RED)
    _add_text(slide, "若是有效见顶信号，倍率应稳定高于1；实际主口径两组均低于1。", 8.95, 5.28, 3.35, 0.95, size=12.0, color=INK, bold=True)

    # 11 — Q3 cooldown/paired
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, WHITE)
    _add_header(slide, "QUESTION 3", "去重与同episode配对后，仍无稳定优势", 11)
    _add_chart(slide, chart_paths["q3_paired"], 0.62, 1.5, 7.9, 5.0)
    q3_60_cool = findings["q3"][60]["cooldown_10d"]
    q3_120_cool = findings["q3"][120]["cooldown_10d"]
    _metric_card(slide, left=8.8, top=1.62, width=3.62, value=f"{q3_60_cool['drawdown_rate_lift']:.2f}×", label="MA60 去重后", note="接近1，顶部型路径仍未改善", color=AMBER, fill=LIGHT_AMBER)
    _metric_card(slide, left=8.8, top=3.28, width=3.62, value=f"{q3_120_cool['drawdown_rate_lift']:.2f}×", label="MA120 去重后", note="仍低于非背离对照", color=RED, fill=LIGHT_RED)
    _add_text(slide, f"同episode配对：MA60背离回撤率更低47只/更高41只；MA120更低81只/更高39只。", 8.84, 5.22, 3.52, 1.1, size=11.8, color=INK, bold=True)

    # 12 — Q3 peak coverage and sensitivity
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, PAPER)
    _add_header(slide, "QUESTION 3", "它会出现于顶部附近，但精确性和覆盖率都不高", 12)
    _add_chart(slide, chart_paths["q3_sensitivity"], 0.62, 1.5, 7.8, 5.0)
    coverage = findings["q3_top_coverage"]
    _metric_card(slide, left=8.72, top=1.62, width=3.78, value=f"{coverage['top_coverage_rate_pct']:.1f}%", label="MA60最终顶部覆盖率", note=f"{coverage['anchors_with_divergence_within_10d_of_top']}/{coverage['completed_anchor_count']}只顶前10日出现背离", color=PURPLE, fill=LIGHT_PURPLE)
    _metric_card(slide, left=8.72, top=3.28, width=3.78, value=f"{q3_60['divergence_near_top_rate']:.1f}%", label="背离事件顶部精确率", note=f"非背离新高日为 {q3_60['control_near_top_rate']:.1f}%", color=RED, fill=LIGHT_RED)
    _add_text(slide, "阈值、回看长度、成交量窗口和最高价口径敏感性均未形成稳定>1的优势。", 8.78, 5.26, 3.62, 0.9, size=11.8, color=INK, bold=True)

    # 13 — examples
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, WHITE)
    _add_header(slide, "QUESTION 3", "背离既能命中顶部，也会在趋势中途反复误报", 13)
    _add_chart(slide, chart_paths["q3_examples"], 0.62, 1.45, 12.05, 5.3)
    _add_text(slide, "案例用于说明机制，不参与统计结论；红虚线为背离信号日。", 0.68, 6.78, 11.8, 0.25, size=9.5, color=MUTED)

    # 14 — practical use
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, PAPER)
    _add_header(slide, "IMPLICATIONS", "如何把三项结论用于长趋势观察", 14)
    _metric_card(slide, left=0.65, top=1.62, width=3.75, value="确认", label="阶段整体放量", note="可作为长趋势已获资金参与的状态特征", color=BLUE, fill=LIGHT_BLUE)
    _metric_card(slide, left=4.78, top=1.62, width=3.75, value="结构", label="MA20段间歇放量", note="看段级分布，不要求每段都增加", color=PURPLE, fill=LIGHT_PURPLE)
    _metric_card(slide, left=8.91, top=1.62, width=3.75, value="组合", label="量价背离", note="只能作风险提示，不能独立触发卖出", color=RED, fill=LIGHT_RED)
    _bullet_block(
        slide,
        [
            "更可用的规则：量价背离 + MA20速度急降/拐头 + 跌破短期结构 + 随后放量下跌，组合确认。",
            "单次缩量创新高可能只是筹码稳定、抛压减少；趋势越强，越容易在真正见顶前多次出现。",
            "对右截断MA120样本，任何“已经见顶”的结论都属于前视误判；只能持续更新未来窗口。",
        ],
        left=0.85,
        top=3.72,
        width=11.65,
        height=2.35,
        size=14.0,
    )

    # 15 — appendix
    slide = presentation.slides.add_slide(blank)
    _set_slide_background(slide, NAVY)
    _add_text(slide, "附录｜解释边界与复算材料", 0.7, 0.72, 11.7, 0.55, size=28, color=WHITE, bold=True)
    _add_rect(slide, 0.7, 1.45, 11.9, 0.03, fill="#304158")
    _bullet_block(
        slide,
        [
            "样本是“MA60/MA120连续上升榜”条件样本，不能外推到全部A股，也不能反推放量一定产生长趋势。",
            "成交量分布右偏，因此主结论优先使用个股自身倍率的中位数、分位数与阈值比例，而不是绝对成交量。",
            "MA120全部右截断；阶段放量与MA20分段结论均表述为“截至2026-09-30的已观察阶段”。",
            "量价背离信号特征无前视；最终顶部标签仅用于已结束MA60 episode，未来20日样本不足者从结果分母排除。",
            "明细、汇总和元数据保存在对应研究目录；PPT/PDF、报告与关键结论保存在本报告目录，可逐项审计。",
        ],
        left=0.9,
        top=1.95,
        width=11.4,
        height=3.75,
        size=14.0,
        color="#D9E3EE",
    )
    _add_text(slide, "核心结论：整体放量成立；每段都放量不成立；量价背离单独见顶无效。", 0.82, 6.25, 11.7, 0.46, size=18, color="#8BE1D2", bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "15", 11.9, 7.12, 0.7, 0.2, size=8, color="#708095", align=PP_ALIGN.RIGHT, margin=0)

    output.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(output)
    return {"slide_count": len(presentation.slides)}


def write_markdown_report(findings: dict[str, Any], output: Path) -> None:
    """Write the complete three-question report in concise Chinese."""
    q1_60, q1_120 = findings["q1"][60], findings["q1"][120]
    q2_60, q2_120 = findings["q2"][60], findings["q2"][120]
    q3_60 = findings["q3"][60]["all_candidate_days"]
    q3_120 = findings["q3"][120]["all_candidate_days"]
    q3_60_c = findings["q3"][60]["cooldown_10d"]
    q3_120_c = findings["q3"][120]["cooldown_10d"]
    coverage = findings["q3_top_coverage"]
    lines = [
        "# MA60/MA120长趋势的三个成交量问题",
        "",
        "样本截至2026-09-30；均线与价格为前复权，成交量为原始成交量；按实际交易日计算。",
        "带入锚点按完整已观察区间分析，可能包含2025年事件；“2026 YTD”指样本筛选窗口，并非把事件强行截到2026年。",
        "",
        "## 结论",
        "",
        "1. **问题一：是。** MA60/MA120连续上升阶段的整体日均成交量，普遍且明显高于阶段前。",
        "2. **问题二：不是每一段。** MA20上升段中多数放量，但同一股票所有MA20段都放量的比例很低。",
        "3. **问题三：否。** 本研究定义下的缩量创新高没有形成稳定的见顶优势，不能单独作为卖出信号。",
        "",
        "## 问题一：长均线阶段整体是否显著放量",
        "",
        "主倍率 = 长均线episode内日均成交量 ÷ episode开始前等长交易日的日均成交量。",
        "",
        "| 锚定阶段 | 可比样本 | 中位倍率 | 平均倍率 | 高于阶段前 | ≥1.2倍 | ≥1.5倍 | ≥2倍 |",
        "|---|---:|---:|---:|---:|---:|---:|---:|",
        f"| MA60 | {q1_60['comparable_count']} | {q1_60['equal_length_median_ratio']:.3f}× | {q1_60['equal_length_mean_ratio']:.3f}× | {q1_60['amplified_gt_1_pct']:.1f}% | {q1_60['amplified_ge_1_2_pct']:.1f}% | {q1_60['amplified_ge_1_5_pct']:.1f}% | {q1_60['amplified_ge_2_pct']:.1f}% |",
        f"| MA120 | {q1_120['comparable_count']} | {q1_120['equal_length_median_ratio']:.3f}× | {q1_120['equal_length_mean_ratio']:.3f}× | {q1_120['amplified_gt_1_pct']:.1f}% | {q1_120['amplified_ge_1_2_pct']:.1f}% | {q1_120['amplified_ge_1_5_pct']:.1f}% | {q1_120['amplified_ge_2_pct']:.1f}% |",
        "",
        f"固定20日前置窗口的中位倍率为MA60 {q1_60['pre20_median_ratio']:.3f}×、MA120 {q1_120['pre20_median_ratio']:.3f}×，结论一致。MA120等长窗口有4只前置历史不足，因此主分母为120；固定窗口分母为124。",
        "",
        "## 问题二：每一个MA20上升段是否都放量",
        "",
        "段级主倍率 = 完整MA20上升episode日均成交量 ÷ 该episode开始前等长交易日均量。",
        "",
        "| 锚定阶段 | 最短段长 | 可比段 | 放量段比例 | 段倍率中位数 | 所有段均放量的股票比例 | 每股放量段占比中位数 |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for period, metrics in ((60, q2_60), (120, q2_120)):
        for days in (1, 5, 10):
            row = metrics[days]
            lines.append(
                f"| MA{period} | ≥{days}日 | {row['segment_count']} | {row['segment_amplified_pct']:.1f}% | {row['segment_median_ratio']:.3f}× | {row['all_segments_amplified_pct']:.1f}% | {row['median_anchor_amplified_segment_pct']:.1f}% |"
            )
    lines.extend(
        [
            "",
            f"严格看全部段，MA60仅{q2_60[1]['all_segments_amplified_count']}/{q2_60[1]['evaluable_anchor_count']}只（{q2_60[1]['all_segments_amplified_pct']:.1f}%）每段都放量；MA120为{q2_120[1]['all_segments_amplified_count']}/{q2_120[1]['evaluable_anchor_count']}只（{q2_120[1]['all_segments_amplified_pct']:.1f}%）。所以应表述为“多数段放量”，不能表述为“每段都放量”。",
            "",
            f"若与整个长均线阶段平均成交量相比，全部MA20段中只有MA60 {q2_60[1]['versus_anchor_mean_amplified_pct']:.1f}%、MA120 {q2_120[1]['versus_anchor_mean_amplified_pct']:.1f}%高于阶段均量，说明MA20上涨段不是持续高于长趋势常态，而是间歇式放量。",
            "",
            "## 问题三：量价背离是否是有效见顶信号",
            "",
            "主信号：收盘价创此前60日新高，且当前5日均量不超过此前价格峰值日5日均量的80%。对照为同样创60日新高但未背离的交易日。",
            "",
            "| 锚定阶段 | 样本口径 | 背离后20日跌幅≥10% | 非背离新高后 | 倍率 | 背离顶部型路径 | 非背离顶部型路径 |",
            "|---|---|---:|---:|---:|---:|---:|",
            f"| MA60 | 所有候选日 | {q3_60['divergence_drawdown_rate']:.1f}% | {q3_60['control_drawdown_rate']:.1f}% | {q3_60['drawdown_rate_lift']:.3f}× | {q3_60['divergence_top_like_rate']:.1f}% | {q3_60['control_top_like_rate']:.1f}% |",
            f"| MA60 | 同类10日去重 | {q3_60_c['divergence_drawdown_rate']:.1f}% | {q3_60_c['control_drawdown_rate']:.1f}% | {q3_60_c['drawdown_rate_lift']:.3f}× | {q3_60_c['divergence_top_like_rate']:.1f}% | {q3_60_c['control_top_like_rate']:.1f}% |",
            f"| MA120 | 所有候选日 | {q3_120['divergence_drawdown_rate']:.1f}% | {q3_120['control_drawdown_rate']:.1f}% | {q3_120['drawdown_rate_lift']:.3f}× | {q3_120['divergence_top_like_rate']:.1f}% | {q3_120['control_top_like_rate']:.1f}% |",
            f"| MA120 | 同类10日去重 | {q3_120_c['divergence_drawdown_rate']:.1f}% | {q3_120_c['control_drawdown_rate']:.1f}% | {q3_120_c['drawdown_rate_lift']:.3f}× | {q3_120_c['divergence_top_like_rate']:.1f}% | {q3_120_c['control_top_like_rate']:.1f}% |",
            "",
            f"在已结束MA60阶段中，{coverage['anchors_with_divergence_within_10d_of_top']}/{coverage['completed_anchor_count']}只（{coverage['top_coverage_rate_pct']:.1f}%）在最终顶部前10日出现背离；但背离事件本身位于顶前10日的比例只有{q3_60['divergence_near_top_rate']:.1f}%，低于非背离新高日的{q3_60['control_near_top_rate']:.1f}%。",
            "",
            f"同episode配对也未支持背离：MA60中背离回撤率更低{q3_60['paired_lower_count']}只、较高{q3_60['paired_higher_count']}只；MA120为更低{q3_120['paired_lower_count']}只、较高{q3_120['paired_higher_count']}只。",
            "",
            "因此，量价背离更适合作为需要组合确认的风险提示，而不是独立见顶信号。可与MA20速度急降、跌破短期结构和随后放量下跌联合使用。",
            "",
            "## 边界",
            "",
            "- 本样本是已进入强长趋势的条件样本，不能外推到全部A股。",
            "- MA120全部截至样本日仍右截断，不能判断最终顶部，只能评价已有未来20日的事件。",
            "- 量价背离结论是关联/预测表现评估，不构成交易建议。",
        ]
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--volume-dir", type=Path, default=DEFAULT_VOLUME_DIR)
    parser.add_argument("--divergence-dir", type=Path, default=DEFAULT_DIVERGENCE_DIR)
    parser.add_argument("--daily-file", type=Path, default=DEFAULT_DAILY)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    return parser.parse_args(argv)


def main() -> None:
    args = _parse_args()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    chart_dir = args.output_dir / "charts"
    chart_dir.mkdir(parents=True, exist_ok=True)

    long_summary = pd.read_csv(args.volume_dir / "long_episode_volume_summary.csv")
    segment_summary = pd.read_csv(args.volume_dir / "ma20_segment_volume_summary.csv")
    anchor_summary = pd.read_csv(args.volume_dir / "ma20_anchor_all_segments_summary.csv")
    long_detail = pd.read_csv(args.volume_dir / "long_episode_volume_detail.csv")
    divergence_events = pd.read_csv(args.divergence_dir / "divergence_event_detail.csv")
    effectiveness = pd.read_csv(args.divergence_dir / "effectiveness_summary.csv")
    lifts = pd.read_csv(args.divergence_dir / "effect_lift_summary.csv")
    coverage = pd.read_csv(args.divergence_dir / "top_coverage_summary.csv")
    within = pd.read_csv(args.divergence_dir / "within_anchor_effect_summary.csv")
    sensitivity = pd.read_csv(args.divergence_dir / "sensitivity_summary.csv")

    findings = derive_key_findings(
        long_summary=long_summary,
        segment_summary=segment_summary,
        anchor_summary=anchor_summary,
        effectiveness=effectiveness,
        lifts=lifts,
        coverage=coverage,
        within=within,
    )

    plt.rcParams["font.sans-serif"] = [
        FONT, "PingFang SC", "Heiti SC", "Arial Unicode MS", "DejaVu Sans"
    ]
    plt.rcParams["axes.unicode_minus"] = False
    chart_paths: dict[str, Path] = {}

    # Chart 1: question-one threshold shares.
    path = chart_dir / "q1_thresholds.png"
    categories = [">1.0×", "≥1.2×", "≥1.5×", "≥2.0×"]
    values60 = [
        findings["q1"][60]["amplified_gt_1_pct"],
        findings["q1"][60]["amplified_ge_1_2_pct"],
        findings["q1"][60]["amplified_ge_1_5_pct"],
        findings["q1"][60]["amplified_ge_2_pct"],
    ]
    values120 = [
        findings["q1"][120]["amplified_gt_1_pct"],
        findings["q1"][120]["amplified_ge_1_2_pct"],
        findings["q1"][120]["amplified_ge_1_5_pct"],
        findings["q1"][120]["amplified_ge_2_pct"],
    ]
    fig, ax = plt.subplots(figsize=(8.6, 5.4), facecolor="white")
    x = np.arange(len(categories)); width = 0.34
    bars1 = ax.bar(x - width / 2, values60, width, label="MA60", color=BLUE)
    bars2 = ax.bar(x + width / 2, values120, width, label="MA120", color=PURPLE)
    ax.set_ylim(0, 105); ax.set_ylabel("股票比例（%）"); ax.set_xticks(x, categories)
    ax.set_title("长均线阶段成交量倍率阈值分布", loc="left", fontweight="bold")
    ax.grid(axis="y", color=GRID, linewidth=0.7); ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, ncol=2, loc="upper right")
    ax.bar_label(bars1, fmt="%.1f", padding=3, fontsize=9); ax.bar_label(bars2, fmt="%.1f", padding=3, fontsize=9)
    fig.tight_layout(); fig.savefig(path, dpi=190, bbox_inches="tight"); plt.close(fig)
    chart_paths["q1_thresholds"] = path

    # Chart 2: equal-length distribution plus fixed-window medians.
    path = chart_dir / "q1_distribution.png"
    data = []
    labels = []
    colors = []
    for period, color in ((60, BLUE), (120, PURPLE)):
        values = extract_equal_length_volume_ratios(long_detail, period)
        data.append(values.clip(upper=4.5).to_numpy()); labels.append(f"MA{period}\n等长前置"); colors.append(color)
    fig, ax = plt.subplots(figsize=(9.2, 5.5), facecolor="white")
    box = ax.boxplot(data, labels=labels, patch_artist=True, widths=0.5, showfliers=False)
    for patch, color in zip(box["boxes"], colors): patch.set(facecolor=color, alpha=0.75)
    for median in box["medians"]: median.set(color=NAVY, linewidth=2)
    ax.axhline(1.0, color=RED, linestyle="--", linewidth=1.2, label="阶段前 = 1.0×")
    ax.set_ylim(0, 4.6); ax.set_ylabel("成交量倍率（极端值截到4.5×仅用于显示）")
    ax.set_title("个股自身前后倍率分布", loc="left", fontweight="bold")
    ax.grid(axis="y", color=GRID, linewidth=0.7); ax.spines[["top", "right"]].set_visible(False)
    ax.legend(frameon=False, loc="upper left")
    notes = [
        f"20日前置中位\nMA60 {findings['q1'][60]['pre20_median_ratio']:.2f}×",
        f"20日前置中位\nMA120 {findings['q1'][120]['pre20_median_ratio']:.2f}×",
    ]
    for i, note in enumerate(notes, start=1): ax.text(i, 4.1, note, ha="center", va="top", fontsize=10, color=colors[i-1])
    fig.tight_layout(); fig.savefig(path, dpi=190, bbox_inches="tight"); plt.close(fig)
    chart_paths["q1_distribution"] = path

    # Chart 3: segment amplified shares by minimum duration.
    path = chart_dir / "q2_segments.png"
    fig, ax = plt.subplots(figsize=(9.2, 5.5), facecolor="white")
    days = np.array([1, 5, 10]); x = np.arange(3); width = 0.34
    values60 = [findings["q2"][60][int(d)]["segment_amplified_pct"] for d in days]
    values120 = [findings["q2"][120][int(d)]["segment_amplified_pct"] for d in days]
    bars1 = ax.bar(x - width / 2, values60, width, color=BLUE, label="MA60")
    bars2 = ax.bar(x + width / 2, values120, width, color=PURPLE, label="MA120")
    ax.set_xticks(x, ["所有段", "≥5日段", "≥10日段"]); ax.set_ylim(0, 100); ax.set_ylabel("放量段比例（%）")
    ax.set_title("MA20上升段相对自身前置等长窗口", loc="left", fontweight="bold")
    ax.grid(axis="y", color=GRID, linewidth=0.7); ax.spines[["top", "right"]].set_visible(False); ax.legend(frameon=False)
    ax.bar_label(bars1, fmt="%.1f", padding=3); ax.bar_label(bars2, fmt="%.1f", padding=3)
    fig.tight_layout(); fig.savefig(path, dpi=190, bbox_inches="tight"); plt.close(fig)
    chart_paths["q2_segments"] = path

    # Chart 4: all-segments amplified shares by stock.
    path = chart_dir / "q2_all_segments.png"
    fig, ax = plt.subplots(figsize=(9.2, 5.5), facecolor="white")
    values60 = [findings["q2"][60][int(d)]["all_segments_amplified_pct"] for d in days]
    values120 = [findings["q2"][120][int(d)]["all_segments_amplified_pct"] for d in days]
    bars1 = ax.bar(x - width / 2, values60, width, color=BLUE, label="MA60")
    bars2 = ax.bar(x + width / 2, values120, width, color=PURPLE, label="MA120")
    ax.set_xticks(x, ["所有段", "≥5日段", "≥10日段"]); ax.set_ylim(0, 55); ax.set_ylabel("所有合格段均放量的股票比例（%）")
    ax.set_title("严格的“每一个MA20段都放量”", loc="left", fontweight="bold")
    ax.grid(axis="y", color=GRID, linewidth=0.7); ax.spines[["top", "right"]].set_visible(False); ax.legend(frameon=False)
    ax.bar_label(bars1, fmt="%.1f", padding=3); ax.bar_label(bars2, fmt="%.1f", padding=3)
    fig.tight_layout(); fig.savefig(path, dpi=190, bbox_inches="tight"); plt.close(fig)
    chart_paths["q2_all_segments"] = path

    # Chart 5: benchmark comparison.
    path = chart_dir / "q2_benchmarks.png"
    fig, ax = plt.subplots(figsize=(9.2, 5.5), facecolor="white")
    labels = ["MA60\n对段前", "MA60\n对阶段均量", "MA120\n对段前", "MA120\n对阶段均量"]
    values = [
        findings["q2"][60][1]["segment_amplified_pct"],
        findings["q2"][60][1]["versus_anchor_mean_amplified_pct"],
        findings["q2"][120][1]["segment_amplified_pct"],
        findings["q2"][120][1]["versus_anchor_mean_amplified_pct"],
    ]
    bars = ax.bar(np.arange(4), values, color=[BLUE, "#8BB6FF", PURPLE, "#B9A0F5"])
    ax.set_ylim(0, 85); ax.set_xticks(np.arange(4), labels); ax.set_ylabel("放量段比例（%）")
    ax.set_title("同一MA20段，在不同基准下的含义", loc="left", fontweight="bold")
    ax.grid(axis="y", color=GRID, linewidth=0.7); ax.spines[["top", "right"]].set_visible(False); ax.bar_label(bars, fmt="%.1f", padding=3)
    fig.tight_layout(); fig.savefig(path, dpi=190, bbox_inches="tight"); plt.close(fig)
    chart_paths["q2_benchmarks"] = path

    # Chart 6: divergence vs control forward drawdown.
    path = chart_dir / "q3_forward.png"
    fig, ax = plt.subplots(figsize=(9.4, 5.5), facecolor="white")
    labels = ["MA60\n全部", "MA60\n10日去重", "MA120\n全部", "MA120\n10日去重"]
    specs = [(60, "all_candidate_days"), (60, "cooldown_10d"), (120, "all_candidate_days"), (120, "cooldown_10d")]
    divergence_values = [findings["q3"][p][s]["divergence_drawdown_rate"] for p, s in specs]
    control_values = [findings["q3"][p][s]["control_drawdown_rate"] for p, s in specs]
    x4 = np.arange(4); width4 = 0.34
    bars1 = ax.bar(x4 - width4 / 2, divergence_values, width4, color=RED, label="量价背离")
    bars2 = ax.bar(x4 + width4 / 2, control_values, width4, color="#8796A8", label="非背离新高")
    ax.set_ylim(0, 52); ax.set_xticks(x4, labels); ax.set_ylabel("未来20日跌幅≥10%比例（%）")
    ax.set_title("若背离有效，红柱应稳定高于灰柱；实际并非如此", loc="left", fontweight="bold")
    ax.grid(axis="y", color=GRID, linewidth=0.7); ax.spines[["top", "right"]].set_visible(False); ax.legend(frameon=False)
    ax.bar_label(bars1, fmt="%.1f", padding=3, fontsize=9); ax.bar_label(bars2, fmt="%.1f", padding=3, fontsize=9)
    fig.tight_layout(); fig.savefig(path, dpi=190, bbox_inches="tight"); plt.close(fig)
    chart_paths["q3_forward"] = path

    # Chart 7: within-anchor pairing.
    path = chart_dir / "q3_paired.png"
    paired_all = within.loc[within["scope"].eq("all_candidate_days")].sort_values("anchor_period")
    fig, ax = plt.subplots(figsize=(9.2, 5.5), facecolor="white")
    xx = np.arange(len(paired_all)); widthp = 0.3
    higher = paired_all["divergence_rate_higher_count"].to_numpy()
    lower = paired_all["divergence_rate_lower_count"].to_numpy()
    equal = paired_all["divergence_rate_equal_count"].to_numpy()
    b1 = ax.bar(xx - widthp, higher, widthp, color=RED, label="背离回撤率更高")
    b2 = ax.bar(xx, lower, widthp, color=GREEN, label="背离回撤率更低")
    b3 = ax.bar(xx + widthp, equal, widthp, color="#A7B2BE", label="相等")
    ax.set_xticks(xx, [f"MA{int(p)}" for p in paired_all["anchor_period"]]); ax.set_ylabel("可配对episode数量")
    ax.set_title("同一长趋势内部：背离并未更常带来深回撤", loc="left", fontweight="bold")
    ax.grid(axis="y", color=GRID, linewidth=0.7); ax.spines[["top", "right"]].set_visible(False); ax.legend(frameon=False, ncol=3, fontsize=9)
    ax.bar_label(b1, padding=3); ax.bar_label(b2, padding=3); ax.bar_label(b3, padding=3)
    fig.tight_layout(); fig.savefig(path, dpi=190, bbox_inches="tight"); plt.close(fig)
    chart_paths["q3_paired"] = path

    # Chart 8: sensitivity of drawdown-rate lift.
    path = chart_dir / "q3_sensitivity.png"
    sens = select_cooldown_sensitivity(sensitivity)
    labels = []
    values60 = []
    values120 = []
    for spec in sens["spec_name"].drop_duplicates():
        labels.append(humanize_divergence_spec(str(spec)))
        row60 = sens.loc[sens["spec_name"].eq(spec) & sens["anchor_period"].eq(60)]
        row120 = sens.loc[sens["spec_name"].eq(spec) & sens["anchor_period"].eq(120)]
        values60.append(float(row60.iloc[0]["drawdown_rate_lift"]) if len(row60) else np.nan)
        values120.append(float(row120.iloc[0]["drawdown_rate_lift"]) if len(row120) else np.nan)
    fig, ax = plt.subplots(figsize=(9.5, 5.5), facecolor="white")
    sx = np.arange(len(labels)); ws = 0.34
    ax.bar(sx - ws / 2, values60, ws, color=BLUE, label="MA60")
    ax.bar(sx + ws / 2, values120, ws, color=PURPLE, label="MA120")
    ax.axhline(1.0, color=RED, linestyle="--", linewidth=1.2, label="无优势 = 1.0×")
    ax.set_xticks(sx, labels, rotation=24, ha="right"); ax.set_ylim(0.55, 1.35); ax.set_ylabel("未来20日深回撤率倍率")
    ax.set_title("不同参数下，背离优势仍不稳定", loc="left", fontweight="bold")
    ax.grid(axis="y", color=GRID, linewidth=0.7); ax.spines[["top", "right"]].set_visible(False); ax.legend(frameon=False, ncol=3, fontsize=9)
    fig.tight_layout(); fig.savefig(path, dpi=190, bbox_inches="tight"); plt.close(fig)
    chart_paths["q3_sensitivity"] = path

    # Chart 9: one hit and one false positive.
    path = chart_dir / "q3_examples.png"
    divergence_only = divergence_events.loc[
        divergence_events["is_volume_price_divergence"].astype(str).str.lower().eq("true")
        & divergence_events["anchor_period"].eq(60)
        & divergence_events["anchor_peak_observable"].astype(str).str.lower().eq("true")
    ].copy()
    hit_pool = divergence_only.loc[
        divergence_only["within_10d_before_anchor_peak"].astype(str).str.lower().eq("true")
        & divergence_only["forward_drawdown_ge_10pct"].astype(str).str.lower().eq("true")
    ]
    false_pool = divergence_only.loc[
        pd.to_numeric(divergence_only["trading_days_to_anchor_peak"], errors="coerce").ge(40)
        & pd.to_numeric(divergence_only["forward_return_pct"], errors="coerce").gt(5)
    ]
    if hit_pool.empty or false_pool.empty:
        raise ValueError("cannot select divergence illustration examples")
    hit = hit_pool.sort_values("forward_max_drawdown_pct").iloc[0]
    false = false_pool.sort_values("forward_max_runup_pct", ascending=False).iloc[0]
    example_daily = pd.read_csv(
        args.daily_file,
        usecols=["ts_code", "trade_date", "close", "vol"],
        dtype={"ts_code": str, "trade_date": str},
    )
    fig, axes = plt.subplots(2, 2, figsize=(13.2, 5.8), gridspec_kw={"height_ratios": [3.0, 1.15]}, facecolor="white")
    for column, (row, title, color) in enumerate(((hit, "命中：背离后进入最终顶部与回撤", RED), (false, "误报：背离后趋势继续显著上行", AMBER))):
        frame = example_daily.loc[example_daily["ts_code"].eq(str(row["ts_code"]))].sort_values("trade_date").reset_index(drop=True)
        signal_positions = frame.index[frame["trade_date"].eq(str(row["signal_date"]))].tolist()
        if not signal_positions:
            raise ValueError(f"signal date missing for {row['ts_code']}")
        pos = signal_positions[0]; start = max(0, pos - 70); end = min(len(frame), pos + 61)
        view = frame.iloc[start:end].copy().reset_index(drop=True); local_signal = pos - start
        xview = np.arange(len(view)); dates = pd.to_datetime(view["trade_date"], format="%Y%m%d")
        axes[0, column].plot(xview, view["close"], color=color, linewidth=1.8)
        axes[0, column].axvline(local_signal, color=RED, linestyle="--", linewidth=1.3)
        axes[0, column].set_title(f"{title}\n{row['stock_name']} {row['ts_code']} · 信号{row['signal_date']}", loc="left", fontsize=10.5, fontweight="bold")
        axes[0, column].grid(color=GRID, linewidth=0.6); axes[0, column].spines[["top", "right"]].set_visible(False); axes[0, column].set_ylabel("前复权收盘价")
        axes[1, column].bar(xview, view["vol"] / 10000.0, color="#9AA7B5", width=0.8)
        axes[1, column].plot(xview, view["vol"].rolling(5).mean() / 10000.0, color=TEAL, linewidth=1.2)
        axes[1, column].axvline(local_signal, color=RED, linestyle="--", linewidth=1.3)
        tick_count = min(5, len(view)); ticks = np.linspace(0, len(view) - 1, tick_count, dtype=int)
        axes[1, column].set_xticks(ticks, [dates.iloc[t].strftime("%Y-%m-%d") for t in ticks], rotation=20, ha="right", fontsize=8)
        axes[1, column].set_ylabel("成交量（万手）"); axes[1, column].grid(axis="y", color=GRID, linewidth=0.6); axes[1, column].spines[["top", "right"]].set_visible(False)
    fig.tight_layout(); fig.savefig(path, dpi=190, bbox_inches="tight"); plt.close(fig)
    chart_paths["q3_examples"] = path

    pptx_path = args.output_dir / "MA60_MA120_Three_Volume_Questions_2026YTD_QFQ.pptx"
    report_path = args.output_dir / "report.md"
    findings_path = args.output_dir / "key_findings.json"
    metadata_path = args.output_dir / "metadata.json"
    stats = build_presentation(findings, chart_paths, pptx_path)
    write_markdown_report(findings, report_path)
    findings_path.write_text(
        json.dumps(findings, ensure_ascii=False, indent=2, allow_nan=True),
        encoding="utf-8",
    )
    metadata = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "data_end": "20260930",
        "price_basis": "qfq",
        "volume_basis": "raw_vol",
        "slide_count": stats["slide_count"],
        "questions": [
            "长均线阶段整体是否较阶段前放量",
            "长均线阶段内每个MA20上升段是否都放量",
            "量价背离是否是有效见顶信号",
        ],
        "source_directories": {
            "volume": str(args.volume_dir.resolve()),
            "divergence": str(args.divergence_dir.resolve()),
        },
        "outputs": {},
    }
    for artifact in [pptx_path, report_path, findings_path, *chart_paths.values()]:
        metadata["outputs"][artifact.name] = {
            "path": str(artifact.resolve()),
            "sha256": _sha256(artifact),
        }
    metadata_path.write_text(
        json.dumps(metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    print(json.dumps({"pptx": str(pptx_path), "report": str(report_path), **stats}, ensure_ascii=False))


if __name__ == "__main__":
    main()

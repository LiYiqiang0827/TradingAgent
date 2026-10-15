from __future__ import annotations

from collections import Counter
import sys
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_ma60_top1pct_kline_ppt import (
    add_moving_averages,
    build_presentation,
    DAILY_QFQ,
    display_start_for_episode,
    MANUAL_CLASSIFICATIONS,
    materialize_classifications,
    render_kline_chart,
    validate_classifications,
)


def test_deck_uses_unadjusted_daily_prices() -> None:
    assert DAILY_QFQ is False


def test_add_moving_averages_sorts_each_stock_and_keeps_volume() -> None:
    daily = pd.DataFrame(
        {
            "ts_code": ["A", "A", "A", "B", "B"],
            "trade_date": ["20260103", "20260101", "20260102", "20260102", "20260101"],
            "close": [3.0, 1.0, 2.0, 20.0, 10.0],
            "vol": [30, 10, 20, 200, 100],
        }
    )

    result = add_moving_averages(daily, periods=(2,))

    assert result.loc[result.ts_code.eq("A"), "trade_date"].tolist() == [
        "20260101", "20260102", "20260103"
    ]
    assert result.loc[result.ts_code.eq("A"), "ma2"].tolist() == [pytest.approx(float("nan"), nan_ok=True), 1.5, 2.5]
    assert result.loc[result.ts_code.eq("B"), "vol"].tolist() == [100, 200]


def test_display_start_includes_long_episode_or_minimum_context() -> None:
    assert display_start_for_episode("20250604") == pd.Timestamp("2025-03-06")
    assert display_start_for_episode("20260105") == pd.Timestamp("2025-10-01")


def test_validate_classifications_requires_exact_stock_coverage() -> None:
    top = pd.DataFrame({"rank": [1, 2], "ts_code": ["A", "B"]})
    good = [
        {"ts_code": "A", "primary_theme": "芯片", "broad_category": "半导体/芯片", "secondary_themes": []},
        {"ts_code": "B", "primary_theme": "算力", "broad_category": "AI算力/数据中心", "secondary_themes": ["数据中心"]},
    ]

    merged = validate_classifications(top, good)
    assert merged["ts_code"].tolist() == ["A", "B"]

    with pytest.raises(ValueError, match="coverage"):
        validate_classifications(top, good[:1])


def test_materialize_classifications_keeps_verbatim_evidence() -> None:
    evidence = [
        {
            "rank": 1,
            "ts_code": "A",
            "stock_name": "甲公司",
            "industry": "半导体",
            "kpl_list_themes": [{"theme": "芯片", "count": 3}],
            "limit_performance_themes": [],
            "concepts": [{"theme": "先进封装", "observed_days": 120}],
        }
    ]
    manual = {
        "A": {
            "primary_theme": "半导体设备",
            "broad_category": "半导体/芯片",
            "secondary_themes": ["先进封装"],
        }
    }

    result = materialize_classifications(evidence, manual)

    assert result[0]["evidence_labels"] == ["芯片", "先进封装"]
    assert result[0]["confidence"] == "high"
    assert "芯片×3" in result[0]["evidence_note"]


def test_reviewed_classification_mapping_matches_evidence_review() -> None:
    counts = Counter(
        item["broad_category"] for item in MANUAL_CLASSIFICATIONS.values()
    )
    assert counts == {
        "PCB/通信硬件": 16,
        "半导体/芯片": 14,
        "其他科技电子": 7,
        "AI算力/数据中心": 5,
        "机器人/高端装备": 5,
        "地产/基建": 4,
        "有色金属/稀土/资源": 3,
        "化工/新材料": 1,
        "消费": 1,
        "医药医疗": 1,
        "军工/航空航天": 1,
    }
    assert MANUAL_CLASSIFICATIONS["688630.SH"]["primary_theme"] == "PCB设备"
    assert MANUAL_CLASSIFICATIONS["301115.SZ"]["primary_theme"] == "机器人概念"
    assert MANUAL_CLASSIFICATIONS["301588.SZ"]["broad_category"] == "地产/基建"


def test_render_kline_chart_contains_price_and_volume_panels(tmp_path: Path) -> None:
    dates = pd.bdate_range("2025-01-01", periods=150)
    close = pd.Series(range(100, 250), dtype=float)
    daily = pd.DataFrame(
        {
            "ts_code": "A",
            "trade_date": dates.strftime("%Y%m%d"),
            "open": close - 0.5,
            "high": close + 1.0,
            "low": close - 1.0,
            "close": close,
            "vol": 1_000 + close * 10,
        }
    )
    daily = add_moving_averages(daily)
    output = tmp_path / "chart.png"

    receipt = render_kline_chart(
        daily,
        output,
        stock_name="测试股份",
        ts_code="A",
        episode_start="20250301",
        episode_end="20250630",
        display_start=pd.Timestamp("2025-02-01"),
        display_end=pd.Timestamp("2025-07-31"),
    )

    assert output.exists()
    assert output.stat().st_size > 10_000
    assert receipt["bars"] > 0
    assert receipt["first_date"] <= receipt["last_date"]


def test_build_presentation_groups_stocks_and_adds_one_chart_slide_each(tmp_path: Path) -> None:
    image_a = tmp_path / "A.png"
    image_b = tmp_path / "B.png"
    Image.new("RGB", (1200, 650), "white").save(image_a)
    Image.new("RGB", (1200, 650), "white").save(image_b)
    stocks = pd.DataFrame(
        [
            {
                "rank": 1, "ts_code": "A", "stock_name": "甲公司", "industry": "半导体",
                "start_date": "20260101", "end_date": "20260901", "target_days": 160,
                "length": 160, "carried_in": False, "right_censored": False,
                "primary_theme": "芯片", "broad_category": "半导体/芯片",
                "secondary_themes": ["先进封装"], "evidence_labels": ["芯片"],
                "confidence": "high", "evidence_note": "KPL：芯片×3",
            },
            {
                "rank": 2, "ts_code": "B", "stock_name": "乙公司", "industry": "软件服务",
                "start_date": "20260201", "end_date": "20260930", "target_days": 150,
                "length": 150, "carried_in": False, "right_censored": True,
                "primary_theme": "算力", "broad_category": "AI算力/数据中心",
                "secondary_themes": ["数据中心"], "evidence_labels": ["算力"],
                "confidence": "high", "evidence_note": "KPL：算力×4",
            },
        ]
    )
    output = tmp_path / "atlas.pptx"

    receipt = build_presentation(
        stocks,
        {"A": image_a, "B": image_b},
        output,
        data_end="20260930",
        valid_universe=200,
        nominal_top_n=2,
        cutoff_days=150,
    )

    deck = Presentation(output)
    assert output.exists()
    assert receipt["stock_slides"] == 2
    assert receipt["category_slides"] == 2
    assert len(deck.slides) == 8
    all_text = "\n".join(
        shape.text for slide in deck.slides for shape in slide.shapes if hasattr(shape, "text")
    )
    assert "甲公司" in all_text
    assert "乙公司" in all_text
    fonts = {
        run.font.name
        for slide in deck.slides
        for shape in slide.shapes
        if hasattr(shape, "text_frame") and shape.has_text_frame
        for paragraph in shape.text_frame.paragraphs
        for run in paragraph.runs
        if run.font.name
    }
    assert fonts == {"Noto Sans CJK SC"}

"""First joint MA120 pullback + six-MA cohesion study and PPT."""

from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
from pathlib import Path
import sys
from typing import Any, Sequence

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

import study_first_ma120_pullback_six_ma_box_ppt as base


ROOT = Path(__file__).resolve().parents[5]
DEFAULT_OUTPUT_DIR = ROOT / "outputs/first_ma120_pullback_six_ma_box_first_qualifying/2026_through_20260930_qfq"
FONT_NAME = "Noto Sans CJK SC"


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def quality_tier(narrow_box_pass: bool, box_position_pass: bool) -> str:
    if narrow_box_pass and box_position_pass:
        return "A_窄箱体且均线位置保持"
    if narrow_box_pass:
        return "B_窄箱体但均线位置转弱"
    if box_position_pass:
        return "C_均线位置保持但箱体不足"
    return "D_箱体与位置均不足"


def attach_stock_first_touch_comparison(
    events: pd.DataFrame,
    raw_first_touches: pd.DataFrame,
) -> pd.DataFrame:
    """Compare selected events with each stock's earliest raw touch."""

    result = events.copy()
    if result.empty:
        return result
    result = result.rename(
        columns={
            "true_first_touch_date": "anchor_first_touch_date",
            "true_first_touch_position": "anchor_first_touch_position",
        }
    )
    earliest = (
        raw_first_touches.sort_values(
            ["ts_code", "touch_date", "anchor_period", "anchor_start_date"],
            kind="mergesort",
        )
        .drop_duplicates("ts_code", keep="first")
        .set_index("ts_code")
    )
    result["true_first_touch_date"] = result["ts_code"].map(earliest["touch_date"])
    result["true_first_touch_position"] = result["ts_code"].map(earliest["touch_position"]).astype(int)
    result["delay_from_first_touch_trading_days"] = (
        result["touch_position"].astype(int) - result["true_first_touch_position"]
    )
    if bool((result["delay_from_first_touch_trading_days"] < 0).any()):
        raise AssertionError("Selected event precedes the stock's earliest raw MA120 touch")
    result["skipped_earlier_first_touch"] = result["delay_from_first_touch_trading_days"].gt(0)
    return result


def summarize_joint_events(events: pd.DataFrame) -> pd.DataFrame:
    narrow = events.loc[events["narrow_box_pass"]]
    return pd.DataFrame(
        [
            {
                "scope": "UNION_UNIQUE_FIRST_QUALIFYING",
                "unique_event_stocks": int(len(events)),
                "skipped_earlier_first_touch_stocks": int(events["skipped_earlier_first_touch"].sum()),
                "narrow_box_stocks": int(events["narrow_box_pass"].sum()),
                "box_position_pass_stocks": int(events["box_position_descriptive_pass"].sum()),
                "best_structure_stocks": int((events["narrow_box_pass"] & events["box_position_descriptive_pass"]).sum()),
                "delay_median": float(events["delay_from_first_touch_trading_days"].median()),
                "delay_p25": float(events["delay_from_first_touch_trading_days"].quantile(0.25)),
                "delay_p75": float(events["delay_from_first_touch_trading_days"].quantile(0.75)),
                "ma120_streak_median": float(events["ma120_rising_streak_at_touch"].median()),
                "ma120_streak_p25": float(events["ma120_rising_streak_at_touch"].quantile(0.25)),
                "ma120_streak_p75": float(events["ma120_rising_streak_at_touch"].quantile(0.75)),
                "box_days_median": float(narrow["box_trading_days"].median()) if len(narrow) else float("nan"),
                "box_days_p25": float(narrow["box_trading_days"].quantile(0.25)) if len(narrow) else float("nan"),
                "box_days_p75": float(narrow["box_trading_days"].quantile(0.75)) if len(narrow) else float("nan"),
                "minimum_six_ma_spread_median": float(events["minimum_six_ma_spread_pct"].median()),
                "minimum_six_ma_spread_min": float(events["minimum_six_ma_spread_pct"].min()),
            }
        ]
    )


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    values = list(sys.argv[1:] if argv is None else argv)
    if "--output-dir" not in values:
        values.extend(["--output-dir", str(DEFAULT_OUTPUT_DIR)])
    return base._parse_args(values)


def prepare_ma_frames(daily: pd.DataFrame) -> dict[str, pd.DataFrame]:
    frames: dict[str, pd.DataFrame] = {}
    for code, group in daily.groupby("ts_code", sort=False):
        frame = group.copy()
        frame["trade_date"] = frame["trade_date"].map(base._compact_date)
        frame = (
            frame.sort_values("trade_date", kind="mergesort")
            .drop_duplicates("trade_date", keep="last")
            .reset_index(drop=True)
        )
        close = pd.to_numeric(frame["close"], errors="coerce")
        for period in (5, 10, 20, 30, 60, 120):
            frame[f"ma{period}"] = close.rolling(period, min_periods=period).mean()
        frame["ma120_rising_streak"] = base.compute_rising_streak(frame["ma120"])
        frames[str(code)] = frame
    return frames


def decorate_frames(
    ma_frames: dict[str, pd.DataFrame],
    args: argparse.Namespace,
) -> dict[str, pd.DataFrame]:
    return {
        code: base.add_ma120_pullback_features(
            frame,
            extension_lookback=args.extension_lookback,
            prior_extension_pct=args.prior_extension_pct,
            touch_upper_pct=args.touch_upper_pct,
            close_floor_pct=args.close_floor_pct,
            cohesion_threshold_pct=args.cohesion_threshold_pct,
            max_short_below_ma120_pct=args.max_short_below_ma120_pct,
            min_short_above_count=args.min_short_above_count,
        )
        for code, frame in ma_frames.items()
    }


def scan_joint_events(
    *,
    anchors: pd.DataFrame,
    ma_frames: dict[str, pd.DataFrame],
    args: argparse.Namespace,
) -> dict[str, Any]:
    features = decorate_frames(ma_frames, args)
    records: list[dict[str, Any]] = []
    raw_first_records: list[dict[str, Any]] = []
    eligibility: list[dict[str, Any]] = []
    for anchor in anchors.to_dict("records"):
        position: int | None = None
        first_position: int | None = None
        if bool(anchor["search_eligible"]):
            frame = features[str(anchor["ts_code"])]
            first_position = base.select_first_touch(
                frame,
                anchor_start_date=anchor["start_date"],
                anchor_end_date=anchor["end_date"],
            )
            if first_position is not None:
                first_row = frame.iloc[first_position]
                raw_first_records.append(
                    {
                        "ts_code": str(anchor["ts_code"]),
                        "stock_name": str(anchor.get("stock_name", "")),
                        "anchor_period": int(anchor["anchor_period"]),
                        "anchor_start_date": base._compact_date(anchor["start_date"]),
                        "anchor_end_date": base._compact_date(anchor["end_date"]),
                        "touch_date": base._compact_date(first_row["trade_date"]),
                        "touch_position": int(first_position),
                        "strict_first_touch_qualified": bool(first_row["six_ma_cohesion_position_pass"]),
                        "six_ma_spread_at_touch_pct": float(first_row["six_ma_spread_pct"]),
                        "short_ma_above_ma120_count": int(first_row["short_ma_above_ma120_count"]),
                        "min_short_ma_vs_ma120_pct": float(first_row["min_short_ma_vs_ma120_pct"]),
                        "ma120_rising_streak_at_touch": int(first_row["ma120_rising_streak"]),
                    }
                )
            position = base.select_first_qualifying_touch(
                frame,
                anchor_start_date=anchor["start_date"],
                anchor_end_date=anchor["end_date"],
            )
            if position is not None:
                record = base.build_first_touch_record(
                    frame,
                    position=position,
                    anchor=anchor,
                    args=args,
                )
                if first_position is None:
                    raise ValueError("qualifying touch exists without a prior MA120 touch")
                record["true_first_touch_position"] = int(first_position)
                record["true_first_touch_date"] = base._compact_date(frame.iloc[first_position]["trade_date"])
                record["delay_from_first_touch_trading_days"] = int(position - first_position)
                record["skipped_earlier_first_touch"] = bool(position != first_position)
                record["quality_tier"] = quality_tier(
                    bool(record["narrow_box_pass"]),
                    bool(record["box_position_descriptive_pass"]),
                )
                records.append(record)
        eligibility.append(
            {
                "ts_code": str(anchor["ts_code"]),
                "stock_name": str(anchor.get("stock_name", "")),
                "anchor_period": int(anchor["anchor_period"]),
                "anchor_start_date": base._compact_date(anchor["start_date"]),
                "anchor_end_date": base._compact_date(anchor["end_date"]),
                "left_censored": bool(anchor["left_censored"]),
                "search_eligible": bool(anchor["search_eligible"]),
                "first_ma120_touch_date": (
                    base._compact_date(features[str(anchor["ts_code"])].iloc[first_position]["trade_date"])
                    if first_position is not None else ""
                ),
                "first_qualifying_date": (
                    base._compact_date(features[str(anchor["ts_code"])].iloc[position]["trade_date"])
                    if position is not None else ""
                ),
            }
        )
    anchor_events = pd.DataFrame(records)
    if anchor_events.empty:
        raise ValueError("no first jointly qualifying events found")
    raw_first_touches = pd.DataFrame(raw_first_records)
    unique_raw_first_touches = base.select_unique_first_touches(raw_first_touches)
    anchor_events = attach_stock_first_touch_comparison(anchor_events, raw_first_touches)
    unique_events = base.select_unique_first_touches(anchor_events)
    tier_order = {
        "A_窄箱体且均线位置保持": 1,
        "B_窄箱体但均线位置转弱": 2,
        "C_均线位置保持但箱体不足": 3,
        "D_箱体与位置均不足": 4,
    }
    unique_events["quality_rank"] = unique_events["quality_tier"].map(tier_order).astype(int)
    unique_events = unique_events.sort_values(
        ["quality_rank", "minimum_six_ma_spread_pct", "touch_date", "ts_code"],
        kind="mergesort",
    ).reset_index(drop=True)
    unique_events["display_order"] = np.arange(1, len(unique_events) + 1)
    unique_events["event_key"] = unique_events.apply(
        lambda row: f"{row['ts_code']}|{int(row['anchor_period'])}|{row['anchor_start_date']}|{row['touch_date']}",
        axis=1,
    )
    coverage_rows: list[dict[str, Any]] = []
    for period in (60, 120):
        anchor_group = anchors.loc[anchors["anchor_period"].eq(period)]
        event_group = anchor_events.loc[anchor_events["anchor_period"].eq(period)]
        coverage_rows.append(
            {
                "anchor_period": period,
                "anchor_count": int(len(anchor_group)),
                "left_censored_count": int(anchor_group["left_censored"].sum()),
                "eligible_count": int(anchor_group["search_eligible"].sum()),
                "first_qualifying_event_count": int(len(event_group)),
                "narrow_box_count": int(event_group["narrow_box_pass"].sum()),
                "best_structure_count": int((event_group["narrow_box_pass"] & event_group["box_position_descriptive_pass"]).sum()),
            }
        )
    summary = summarize_joint_events(unique_events)
    summary["raw_first_touch_stocks"] = int(len(unique_raw_first_touches))
    summary["strict_first_touch_qualified_stocks"] = int(unique_raw_first_touches["strict_first_touch_qualified"].sum())
    return {
        "features_by_code": features,
        "eligibility": pd.DataFrame(eligibility),
        "raw_first_touches": raw_first_touches,
        "unique_raw_first_touches": unique_raw_first_touches,
        "anchor_events": anchor_events,
        "unique_events": unique_events,
        "best_events": unique_events.loc[
            unique_events["narrow_box_pass"] & unique_events["box_position_descriptive_pass"]
        ].copy(),
        "coverage": pd.DataFrame(coverage_rows),
        "summary": summary,
    }


def _namespace_with(args: argparse.Namespace, **changes: Any) -> argparse.Namespace:
    values = vars(args).copy()
    values.update(changes)
    return argparse.Namespace(**values)


def build_joint_sensitivity(
    *,
    anchors: pd.DataFrame,
    ma_frames: dict[str, pd.DataFrame],
    args: argparse.Namespace,
    main_result: dict[str, Any],
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    grids = (
        ("six_ma_spread_threshold", "cohesion_threshold_pct", (4.0, 5.0, 6.0, 7.0, 8.0)),
        ("box_range_threshold", "box_max_range_pct", (8.0, 10.0, 12.0)),
        ("prior_extension_threshold", "prior_extension_pct", (5.0, 8.0, 10.0)),
        ("touch_upper_threshold", "touch_upper_pct", (1.0, 2.0, 3.0)),
        ("close_floor_threshold", "close_floor_pct", (-1.0, -3.0, -5.0)),
        ("position_floor_threshold", "max_short_below_ma120_pct", (-1.0, -2.0, -3.0)),
        ("min_short_above_count", "min_short_above_count", (3, 4, 5)),
    )
    cache: dict[tuple[Any, ...], dict[str, Any]] = {}
    parameter_names = (
        "cohesion_threshold_pct", "box_max_range_pct", "prior_extension_pct",
        "touch_upper_pct", "close_floor_pct", "max_short_below_ma120_pct",
        "min_short_above_count",
    )
    main_key = tuple(getattr(args, name) for name in parameter_names)
    cache[main_key] = main_result
    for sensitivity_type, attribute, values in grids:
        for value in values:
            variant = _namespace_with(args, **{attribute: value})
            key = tuple(getattr(variant, name) for name in parameter_names)
            if key not in cache:
                cache[key] = scan_joint_events(anchors=anchors, ma_frames=ma_frames, args=variant)
            result = cache[key]
            events = result["unique_events"]
            rows.append(
                {
                    "sensitivity_type": sensitivity_type,
                    "threshold": value,
                    "event_count": int(len(events)),
                    "narrow_box_count": int(events["narrow_box_pass"].sum()),
                    "best_structure_count": int((events["narrow_box_pass"] & events["box_position_descriptive_pass"]).sum()),
                    "skipped_earlier_first_touch_count": int(events["skipped_earlier_first_touch"].sum()),
                    "box_range_threshold_pct": float(variant.box_max_range_pct),
                }
            )
    return pd.DataFrame(rows)


def create_joint_event_chart(frame: pd.DataFrame, event: pd.Series, output: Path) -> None:
    base._register_fonts()
    dates = frame["trade_date"].map(base._compact_date)
    touch = int(event["touch_position"])
    first_touch = int(event["true_first_touch_position"])
    has_box = pd.notna(event["box_start_position"]) and pd.notna(event["box_end_position"])
    box_start = int(event["box_start_position"]) if has_box else touch
    box_end = int(event["box_end_position"]) if has_box else touch
    if touch - first_touch <= 160:
        lower = max(0, first_touch - 12)
    else:
        lower = max(0, touch - 160)
    lower = min(lower, max(0, box_start - 15))
    upper = min(len(frame) - 1, max(touch + 28, box_end + 18))
    view = frame.iloc[lower : upper + 1].reset_index(drop=True).copy()
    touch_x = touch - lower
    first_x = first_touch - lower
    box_start_x = box_start - lower
    box_end_x = box_end - lower

    fig = plt.figure(figsize=(12.2, 6.35), dpi=155, facecolor="white")
    grid = fig.add_gridspec(5, 1, hspace=0.06)
    ax = fig.add_subplot(grid[:4, 0])
    vol_ax = fig.add_subplot(grid[4, 0], sharex=ax)
    base._draw_candlesticks(ax, view)
    colors = {5: base.ORANGE, 10: base.GOLD, 20: "E76F9A", 30: base.BLUE, 60: "7B2CBF", 120: base.TEAL}
    for period in (5, 10, 20, 30, 60, 120):
        ax.plot(
            np.arange(len(view)),
            pd.to_numeric(view[f"ma{period}"], errors="coerce"),
            color=f"#{colors[period]}",
            linewidth=1.05 if period < 30 else 1.4,
            label=f"MA{period}",
            zorder=4,
        )
    if 0 <= first_x < len(view) and first_touch != touch:
        ax.axvline(first_x, color=f"#{base.MUTED}", linewidth=1.0, linestyle=":", zorder=5)
        first_close = float(view.iloc[first_x]["close"])
        ax.annotate(
            f"较早首次回调\n{base._fmt_date(event['true_first_touch_date'])}",
            xy=(first_x, first_close), xytext=(5, -38), textcoords="offset points",
            fontsize=7.2, color=f"#{base.MUTED}",
            arrowprops={"arrowstyle": "->", "color": f"#{base.MUTED}", "lw": 0.7},
        )
    ax.axvline(touch_x, color=f"#{base.RED}", linewidth=1.6, linestyle="--", zorder=6)
    if has_box:
        box_slice = view.iloc[max(0, box_start_x) : box_end_x + 1]
        box_low = float(pd.to_numeric(box_slice["low"], errors="coerce").min())
        box_high = float(pd.to_numeric(box_slice["high"], errors="coerce").max())
        ax.add_patch(
            Rectangle(
                (box_start_x - 0.5, box_low), box_end_x - box_start_x + 1, box_high - box_low,
                facecolor=f"#{base.CYAN}", edgecolor=f"#{base.BLUE}", linewidth=1.2,
                alpha=0.16, zorder=1,
            )
        )
    touch_close = float(view.iloc[touch_x]["close"])
    ax.scatter([touch_x], [touch_close], s=52, color=f"#{base.RED}", edgecolor="white", linewidth=0.8, zorder=7)
    ax.annotate(
        f"首个同时满足\n{base._fmt_date(event['touch_date'])}",
        xy=(touch_x, touch_close), xytext=(8, 24), textcoords="offset points",
        fontsize=8, color=f"#{base.RED}", fontweight="bold",
        arrowprops={"arrowstyle": "->", "color": f"#{base.RED}", "lw": 0.8},
    )
    min_date = base._compact_date(event["minimum_six_ma_spread_date"])
    if min_date:
        min_matches = np.flatnonzero(view["trade_date"].map(base._compact_date).eq(min_date).to_numpy())
        if len(min_matches) == 1:
            ax.axvline(int(min_matches[0]), color=f"#{base.TEAL}", linewidth=1.0, linestyle="-.", zorder=5)
    ax.legend(loc="upper left", ncol=6, frameon=False, fontsize=7.2)
    ax.grid(axis="y", color="#E7EDF3", linewidth=0.6)
    ax.set_ylabel("前复权价格", fontsize=8, color=f"#{base.MUTED}")
    ax.tick_params(axis="both", labelsize=7, colors=f"#{base.MUTED}")
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["left", "bottom"]].set_color("#CBD5E1")
    plt.setp(ax.get_xticklabels(), visible=False)

    volumes = pd.to_numeric(view["vol"], errors="coerce").fillna(0.0)
    volume_colors = [
        f"#{base.RED}" if float(close) >= float(open_) else f"#{base.TEAL}"
        for open_, close in zip(view["open"], view["close"])
    ]
    vol_ax.bar(np.arange(len(view)), volumes, color=volume_colors, width=0.62, alpha=0.75)
    vol_ax.axvline(touch_x, color=f"#{base.RED}", linewidth=1.2, linestyle="--")
    if 0 <= first_x < len(view) and first_touch != touch:
        vol_ax.axvline(first_x, color=f"#{base.MUTED}", linewidth=0.8, linestyle=":")
    vol_ax.grid(axis="y", color="#EEF2F6", linewidth=0.5)
    vol_ax.set_ylabel("成交量", fontsize=7, color=f"#{base.MUTED}")
    vol_ax.tick_params(axis="both", labelsize=6.5, colors=f"#{base.MUTED}")
    vol_ax.spines[["top", "right"]].set_visible(False)
    vol_ax.spines[["left", "bottom"]].set_color("#CBD5E1")
    step = max(1, len(view) // 8)
    ticks = sorted(set(list(range(0, len(view), step)) + [touch_x, len(view) - 1]))
    vol_ax.set_xticks(ticks)
    vol_ax.set_xticklabels([base._fmt_date(view.iloc[index]["trade_date"])[5:] for index in ticks], fontsize=6.5)
    fig.suptitle(
        f"{event['stock_name']}  {event['ts_code']}  |  首个同时满足回调MA120＋六均线粘合",
        x=0.01, y=0.995, ha="left", va="top", fontsize=11,
        color=f"#{base.DARK_TEXT}", fontweight="bold",
    )
    box_note = (
        f"蓝框={int(event['box_trading_days'])}日 / {float(event['box_range_pct']):.1f}%振幅"
        if has_box else "未形成≤10%振幅箱体"
    )
    fig.text(
        0.99, 0.99,
        f"{box_note}；绿线=六均线最窄日",
        ha="right", va="top", fontsize=7.5, color=f"#{base.MUTED}",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def create_joint_distribution_chart(events: pd.DataFrame, output: Path) -> None:
    base._register_fonts()
    fig, axes = plt.subplots(2, 2, figsize=(11.8, 6.0), dpi=160, facecolor="white")
    plots = [
        (events["delay_from_first_touch_trading_days"], "从真正首次回调到入选点的间隔", "交易日", base.BLUE),
        (events["ma120_rising_streak_at_touch"], "入选时MA120连续上涨", "交易日", base.TEAL),
        (events["box_trading_days"], "识别出的箱体持续时间", "交易日", base.ORANGE),
        (events["minimum_six_ma_spread_pct"], "箱体内六均线最窄离差", "%", base.RED),
    ]
    for ax, (values, title, xlabel, color) in zip(axes.ravel(), plots):
        clean = pd.to_numeric(values, errors="coerce").dropna()
        bins = min(12, max(5, int(np.sqrt(len(clean))) + 1))
        ax.hist(clean, bins=bins, color=f"#{color}", alpha=0.84, edgecolor="white")
        ax.axvline(float(clean.median()), color=f"#{base.NAVY}", linewidth=1.3, linestyle="--", label=f"中位数 {clean.median():.1f}")
        ax.set_title(title, fontsize=10.5, fontweight="bold")
        ax.set_xlabel(xlabel, fontsize=8)
        ax.legend(frameon=False, fontsize=8)
        ax.grid(axis="y", color="#E7EDF3", linewidth=0.6)
        ax.spines[["top", "right"]].set_visible(False)
        ax.tick_params(labelsize=7.5, colors=f"#{base.MUTED}")
    fig.tight_layout()
    fig.savefig(output, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def _add_bullets(slide: Any, bullets: list[str], *, top: float = 1.05, height: float = 5.82) -> None:
    box = slide.shapes.add_textbox(Inches(0.72), Inches(top), Inches(11.9), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for index, item in enumerate(bullets):
        paragraph = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        paragraph.text = f"• {item}"
        paragraph.font.name = FONT_NAME
        paragraph.font.size = Pt(12.6)
        paragraph.font.color.rgb = _rgb(base.DARK_TEXT)
        paragraph.space_after = Pt(8)


def build_joint_deck(
    *,
    events: pd.DataFrame,
    anchor_events: pd.DataFrame,
    summary: pd.DataFrame,
    coverage: pd.DataFrame,
    sensitivity: pd.DataFrame,
    chart_paths: dict[str, Path],
    distribution_chart: Path,
    output: Path,
) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    row = summary.iloc[0]

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base._set_background(slide, base.NAVY)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.12), Inches(0.12), Inches(4.92))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = _rgb(base.CYAN); stripe.line.fill.background()
    base._add_text(slide, "首个同时满足回调MA120\n＋六均线粘合的位置", 1.05, 1.18, 11.2, 1.55, size=29, color=base.WHITE, bold=True)
    base._add_text(slide, "允许跳过较早但不粘合的第一次回调", 1.08, 3.02, 10.4, 0.48, size=18, color=base.CYAN)
    base._add_text(slide, "178只长趋势候选池｜前复权日线｜截至2026-09-30", 1.08, 4.00, 9.8, 0.35, size=12, color="BFD7EA")
    base._add_text(slide, f"联合去重 {len(events)} 只｜窄箱体 {int(row['narrow_box_stocks'])} 只｜最佳结构 {int(row['best_structure_stocks'])} 只", 1.08, 4.56, 10.8, 0.55, size=15, color=base.WHITE, bold=True)
    base._add_text(slide, "这是事后形态研究，不构成买入建议", 1.08, 6.72, 6.0, 0.26, size=9, color="94A3B8")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base._set_background(slide, base.LIGHT)
    base._add_header(slide, "新口径：寻找第一个同时满足“回调＋六均线粘合”的日期")
    _add_bullets(slide, [
        "在每个非左截断MA60/MA120锚点内逐日扫描；只有当回调MA120条件与六均线条件同日成立时，才锁定为该锚点事件。",
        "回调条件保持不变：此前20日最高收盘至少高于自身MA120 8%；当日最低价≤MA120上方2%，收盘≥MA120下方3%，且MA120严格上涨。",
        "六均线粘合保持不变：MA5/10/20/30/60/120最大—最小离差除以均值≤6%；MA5—60至少4/5条不低于MA120，最低不超过MA120下方2%。",
        "与旧口径的关键区别：允许跳过较早但不粘合的真正第一次回调，继续寻找后续首个合格位置；因此是形态筛选，不再是严格的“第一次回调”统计。",
        "跨锚点仍先取每只股票最早的合格事件；同日按锚点周期、锚点起点排序。六条MA均先在完整前复权历史上滚动计算。",
        "窄箱体为包含事件日、前10/后30个交易日范围内≤10%振幅的最长区间；同长度取振幅更小、再取起点更早；≥5日为窄箱体。",
        "箱体持续时间、最窄粘合和箱体内均线位置可能使用事件后最多30日数据，只作事后描述；股票池本身也由事后长MA持续时间筛选。",
    ], height=5.95)
    base._add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base._set_background(slide, base.LIGHT)
    base._add_header(slide, "为什么样本从6只增加到60只")
    cards = [
        (str(int(row["raw_first_touch_stocks"])), "找到真正首次MA120回调", base.BLUE),
        (str(int(row["strict_first_touch_qualified_stocks"])), "旧口径：首次回调即合格", base.TEAL),
        (str(len(events)), "新口径：首个同时满足", base.ORANGE),
        (str(int(row["skipped_earlier_first_touch_stocks"])), "跳过了较早不合格回调", base.RED),
    ]
    for index, (value, label, color) in enumerate(cards):
        base._add_card(slide, 0.56 + index * 3.13, 1.18, 2.73, 1.36, value, label, accent=color)
    base._add_text(slide, "旧口径", 0.90, 3.20, 2.0, 0.32, size=15, bold=True, color=base.NAVY)
    base._add_text(slide, "先固定真正第一次回调，再检查粘合；第一次失败即淘汰。", 0.90, 3.65, 5.2, 0.72, size=14, color=base.DARK_TEXT)
    base._add_text(slide, "新口径", 7.15, 3.20, 2.0, 0.32, size=15, bold=True, color=base.NAVY)
    base._add_text(slide, "在所有回调候选中寻找第一个同时满足粘合与位置约束的日期。", 7.15, 3.65, 5.2, 0.72, size=14, color=base.DARK_TEXT)
    arrow = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Inches(5.65), Inches(3.50), Inches(1.45), Inches(0.75))
    arrow.fill.solid(); arrow.fill.fore_color.rgb = _rgb(base.CYAN); arrow.line.fill.background()
    base._add_text(slide, f"54/60只使用了后续事件；从真正首次回调到入选点的间隔中位数 {row['delay_median']:.1f} 个交易日（IQR {row['delay_p25']:.1f}–{row['delay_p75']:.1f}）。", 1.10, 5.25, 11.1, 0.78, size=15, bold=True, color=base.RED, align=PP_ALIGN.CENTER)
    base._add_text(slide, "因此样本扩张来自口径改变，而不是放宽6%粘合阈值。", 1.10, 6.23, 11.1, 0.42, size=11.5, color=base.BLUE, align=PP_ALIGN.CENTER)
    base._add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base._set_background(slide, base.LIGHT)
    base._add_header(slide, "60只样本的结构分层")
    cards = [
        (str(len(events)), "联合去重事件", base.BLUE),
        (str(int(row["narrow_box_stocks"])), "形成≥5日窄箱体", base.CYAN),
        (str(int(row["box_position_pass_stocks"])), "箱体内均线位置保持", base.ORANGE),
        (str(int(row["best_structure_stocks"])), "两项同时满足", base.RED),
    ]
    for index, (value, label, color) in enumerate(cards):
        base._add_card(slide, 0.56 + index * 3.13, 1.18, 2.73, 1.36, value, label, accent=color)
    tier_names = ["A_窄箱体且均线位置保持", "B_窄箱体但均线位置转弱", "C_均线位置保持但箱体不足", "D_箱体与位置均不足"]
    counts = events["quality_tier"].value_counts()
    table = slide.shapes.add_table(5, 3, Inches(1.05), Inches(3.05), Inches(11.2), Inches(2.55)).table
    for col, text in enumerate(["层级", "股票数", "含义"]): table.cell(0, col).text = text
    explanations = ["箱体与均线位置均保持", "有箱体，但短中均线在箱体内转弱", "均线位置保持，但箱体不足5日", "两项描述条件均不足"]
    for index, (name, explanation) in enumerate(zip(tier_names, explanations), start=1):
        for col, text in enumerate([name.split("_", 1)[0], str(int(counts.get(name, 0))), explanation]): table.cell(index, col).text = text
    base._style_table(table, header_size=10.5, body_size=10.5)
    base._add_text(slide, f"入选时MA120连续上涨中位数 {row['ma120_streak_median']:.1f}日（IQR {row['ma120_streak_p25']:.1f}–{row['ma120_streak_p75']:.1f}），明显晚于严格首次回调口径。", 0.95, 6.05, 11.45, 0.55, size=12.5, color=base.NAVY, bold=True, align=PP_ALIGN.CENTER)
    base._add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base._set_background(slide, base.WHITE)
    base._add_header(slide, "分布：多数入选点已不是MA120上升初期")
    slide.shapes.add_picture(str(distribution_chart), Inches(0.55), Inches(1.02), width=Inches(12.15), height=Inches(5.72))
    base._add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base._set_background(slide, base.LIGHT)
    base._add_header(slide, "核心阈值敏感性")
    spread = sensitivity.loc[sensitivity["sensitivity_type"].eq("six_ma_spread_threshold")]
    box = sensitivity.loc[sensitivity["sensitivity_type"].eq("box_range_threshold")]
    table1 = slide.shapes.add_table(len(spread) + 1, 4, Inches(0.55), Inches(1.42), Inches(6.0), Inches(3.85)).table
    for col, text in enumerate(["六线离差", "事件数", "窄箱体", "最佳结构"]): table1.cell(0, col).text = text
    for idx, item in enumerate(spread.itertuples(), start=1):
        for col, text in enumerate([f"≤{item.threshold:.0f}%", str(int(item.event_count)), str(int(item.narrow_box_count)), str(int(item.best_structure_count))]): table1.cell(idx, col).text = text
    base._style_table(table1, header_size=9.5, body_size=10)
    table2 = slide.shapes.add_table(len(box) + 1, 4, Inches(6.80), Inches(1.42), Inches(5.95), Inches(2.75)).table
    for col, text in enumerate(["箱体振幅", "事件数", "窄箱体", "最佳结构"]): table2.cell(0, col).text = text
    for idx, item in enumerate(box.itertuples(), start=1):
        for col, text in enumerate([f"≤{item.threshold:.0f}%", str(int(item.event_count)), str(int(item.narrow_box_count)), str(int(item.best_structure_count))]): table2.cell(idx, col).text = text
    base._style_table(table2, header_size=9.5, body_size=10)
    base._add_text(slide, "主口径仍使用六均线离差≤6%、箱体振幅≤10%；本次只改变“先固定第一次”这一选择顺序。", 0.85, 5.72, 11.7, 0.60, size=13, color=base.BLUE, bold=True, align=PP_ALIGN.CENTER)
    base._add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base._set_background(slide, base.LIGHT)
    base._add_header(slide, "回调与均线位置阈值敏感性：每组重新寻找首个同时满足位置")
    specs = [
        ("prior_extension_threshold", "此前离开MA120", lambda x: f"≥{x:.0f}%"),
        ("touch_upper_threshold", "最低价触及容忍", lambda x: f"≤+{x:.0f}%"),
        ("close_floor_threshold", "收盘破位容忍", lambda x: f"≥{x:.0f}%"),
        ("position_floor_threshold", "短中均线最低位置", lambda x: f"≥{x:.0f}%"),
        ("min_short_above_count", "位于MA120上方条数", lambda x: f"≥{int(x)}/5"),
    ]
    positions = [(0.36, 1.20), (4.58, 1.20), (8.80, 1.20), (2.48, 4.10), (6.70, 4.10)]
    for (kind, title, formatter), (left, top) in zip(specs, positions):
        part = sensitivity.loc[sensitivity["sensitivity_type"].eq(kind)]
        base._add_text(slide, title, left, top, 3.82, 0.30, size=11.5, bold=True, color=base.NAVY, align=PP_ALIGN.CENTER)
        table = slide.shapes.add_table(len(part) + 1, 4, Inches(left), Inches(top + 0.45), Inches(3.82), Inches(1.88)).table
        for col, text in enumerate(["阈值", "事件", "窄箱体", "最佳"]): table.cell(0, col).text = text
        for row_idx, item in enumerate(part.itertuples(), start=1):
            values = [formatter(float(item.threshold)), str(int(item.event_count)), str(int(item.narrow_box_count)), str(int(item.best_structure_count))]
            for col, text in enumerate(values): table.cell(row_idx, col).text = text
        base._style_table(table, header_size=7.8, body_size=8.2)
    base._add_footer(slide, len(prs.slides))

    page_size = 20
    total_list_pages = int(np.ceil(len(events) / page_size))
    for page_index, start in enumerate(range(0, len(events), page_size), start=1):
        subset = events.iloc[start : start + page_size]
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        base._set_background(slide, base.LIGHT)
        base._add_header(slide, f"{len(events)}只事件清单（{page_index}/{total_list_pages}）")
        headers = ["#", "股票", "入选日", "真正首次", "间隔", "MA120", "箱体", "层级"]
        table = slide.shapes.add_table(len(subset) + 1, len(headers), Inches(0.28), Inches(1.0), Inches(12.78), Inches(6.12)).table
        widths = [0.50, 5.05, 1.55, 1.55, 1.05, 1.00, 1.25, 0.83]
        for idx, width in enumerate(widths): table.columns[idx].width = Inches(width)
        for col, text in enumerate(headers): table.cell(0, col).text = text
        for row_index, event in enumerate(subset.itertuples(), start=1):
            values = [
                str(int(event.display_order)), f"{event.stock_name} {event.ts_code}", base._fmt_date(event.touch_date),
                base._fmt_date(event.true_first_touch_date), f"{int(event.delay_from_first_touch_trading_days)}日",
                f"{int(event.ma120_rising_streak_at_touch)}日",
                f"{int(event.box_trading_days)}日" if pd.notna(event.box_trading_days) else "—",
                str(event.quality_tier).split("_", 1)[0],
            ]
            for col, text in enumerate(values): table.cell(row_index, col).text = text
        base._style_table(table, header_size=7.5, body_size=6.8)
        base._add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    base._set_background(slide, base.LIGHT)
    base._add_header(slide, "个股页阅读说明")
    _add_bullets(slide, [
        "红色虚线与红点：该股票第一个同时满足回调MA120、六均线粘合和事件日位置约束的日期。",
        "灰色点线：更早的真正第一次MA120回调；若不在当前窗口，说明两者间隔较长。",
        "蓝色框：包含入选日、振幅不超过10%的最长价格箱体；没有蓝框表示单日高低振幅已超过阈值或没有合格区间。",
        "绿色点划线：识别箱体内六均线最大—最小离差最小的日期。",
        "A层为窄箱体和箱体内均线位置均保持；B/C/D依次表示其中一项或两项描述条件不足。",
        "右侧同时列真正首次回调日期、入选日期和间隔，避免把后续形态误称为真正第一次回调。",
    ])
    base._add_footer(slide, len(prs.slides))

    for _, event in events.iterrows():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        base._set_background(slide, base.LIGHT)
        tier_letter = str(event["quality_tier"]).split("_", 1)[0]
        base._add_header(slide, f"{int(event['display_order']):02d}  {event['stock_name']}  {event['ts_code']}  ·  {tier_letter}层", f"入选日 {base._fmt_date(event['touch_date'])}")
        slide.shapes.add_picture(str(chart_paths[str(event["event_key"])]), Inches(0.38), Inches(1.02), width=Inches(9.38), height=Inches(5.95))
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.93), Inches(1.04), Inches(3.0), Inches(5.90))
        panel.fill.solid(); panel.fill.fore_color.rgb = _rgb(base.WHITE); panel.line.color.rgb = _rgb(base.MID)
        base._add_text(slide, "事件指标", 10.17, 1.20, 2.5, 0.34, size=14.5, bold=True, color=base.NAVY)
        has_box = pd.notna(event["box_start_position"])
        box_value = f"{int(event['box_trading_days'])}日 / {float(event['box_range_pct']):.2f}%" if has_box else "无≤10%箱体"
        metrics = [
            ("锚点", f"MA{int(event['anchor_period'])}  {base._fmt_date(event['anchor_start_date'])}"),
            ("真正首次回调", base._fmt_date(event["true_first_touch_date"])),
            ("首个同时满足", base._fmt_date(event["touch_date"])),
            ("与首次回调间隔", f"{int(event['delay_from_first_touch_trading_days'])}个交易日"),
            ("MA120连续上涨", f"{int(event['ma120_rising_streak_at_touch'])}日"),
            ("事件日六线离差", f"{float(event['six_ma_spread_at_touch_pct']):.2f}%"),
            ("MA5—60位置", f"{int(event['short_ma_above_ma120_count'])}/5条上方；最低{float(event['min_short_ma_vs_ma120_pct']):+.2f}%"),
            ("箱体天数 / 振幅", box_value),
            ("六均线最窄", f"{float(event['minimum_six_ma_spread_pct']):.2f}%" if pd.notna(event["minimum_six_ma_spread_pct"]) else "—"),
            ("箱体内均线位置", f"上方{100.0 * float(event['box_short_ma_observations_above_ma120_ratio']):.0f}% / ≥-2% {100.0 * float(event['box_short_ma_observations_within_floor_ratio']):.0f}%" if pd.notna(event["box_short_ma_observations_above_ma120_ratio"]) else "—"),
            ("结构层级", str(event["quality_tier"])),
        ]
        y = 1.62
        for label, value in metrics:
            base._add_text(slide, label, 10.18, y, 2.5, 0.18, size=7.5, color=base.MUTED)
            base._add_text(slide, value, 10.18, y + 0.17, 2.5, 0.27, size=9.7, bold=True, color=base.DARK_TEXT)
            y += 0.47
        base._add_footer(slide, len(prs.slides))

    prs.save(output)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_joint_report(
    *,
    output: Path,
    args: argparse.Namespace,
    result: dict[str, Any],
    sensitivity: pd.DataFrame,
    pptx_name: str | None,
    pdf_name: str | None,
) -> None:
    summary = result["summary"].iloc[0]
    events = result["unique_events"]
    coverage = result["coverage"]
    tier_counts = events["quality_tier"].value_counts()
    lines = [
        "# 首个同时满足回调MA120＋六均线粘合研究",
        "",
        "## 结论",
        "",
        f"- 178只长趋势候选股、231个非左截断可搜索锚点中，找到 **{len(result['anchor_events'])} 个锚点事件、{len(events)} 只唯一股票**。",
        f"- 严格旧口径有 {int(summary['raw_first_touch_stocks'])} 只股票找到真正首次MA120回调，其中仅 {int(summary['strict_first_touch_qualified_stocks'])} 只在该日已经满足六均线条件。",
        f"- 新口径的 {int(summary['skipped_earlier_first_touch_stocks'])}/{len(events)} 只跳过了更早不合格回调；入选点相对真正首次回调间隔中位数 {summary['delay_median']:.1f} 个交易日，IQR {summary['delay_p25']:.1f}–{summary['delay_p75']:.1f}。",
        f"- {int(summary['narrow_box_stocks'])} 只形成至少{args.box_min_days}日窄箱体；{int(summary['box_position_pass_stocks'])} 只在箱体内保持短中均线位置；两项同时满足的A层为 {int(summary['best_structure_stocks'])} 只。",
        f"- 入选时MA120严格连续上涨中位数 {summary['ma120_streak_median']:.1f}日；窄箱体持续中位数 {summary['box_days_median']:.1f}日；识别箱体内六均线最窄离差中位数 {summary['minimum_six_ma_spread_median']:.2f}%。",
        "",
        "## 口径变化",
        "",
        "旧口径先固定真正第一次MA120回调，再检查同日六均线条件，失败不得后跳。新口径在每个锚点内按交易日顺序扫描，锁定第一个同时满足回调MA120、六均线粘合和事件日均线位置条件的日期，因此允许跳过更早不合格回调。新事件不能再称为真正第一次回调，属于事后形态筛选。",
        "",
        "## 硬筛选条件",
        "",
        f"1. 仅搜索非左截断MA60/MA120锚点；在完整前复权历史上先计算MA5/10/20/30/60/120，再裁锚点。",
        f"2. 此前{args.extension_lookback}个实际交易日最高收盘相对各日自身MA120至少离开 {args.prior_extension_pct:.0f}%。",
        f"3. 事件日最低价不高于MA120上方 {args.touch_upper_pct:.0f}%，收盘不低于MA120 {args.close_floor_pct:.0f}%。",
        f"4. MA120当日严格上涨；`MA_t > MA_(t-1)` 才续段，持平、下降或不可比较均清零。",
        f"5. 六均线离差 `(max-min)/mean` 不超过 {args.cohesion_threshold_pct:.0f}%。",
        f"6. MA5/10/20/30/60至少 {args.min_short_above_count}/5 条不低于MA120，最低一条不低于MA120 {args.max_short_below_ma120_pct:.0f}%。",
        "7. 每个锚点取首个联合合格日；跨锚点取股票最早合格日，同日按锚点周期、锚点起点稳定排序。",
        "",
        "## 箱体与层级",
        "",
        f"- 在事件日前最多{args.box_lookback}个、事件后最多{args.box_lookforward}个实际交易日内，寻找包含事件日且高低振幅≤{args.box_max_range_pct:.0f}%的最长区间；同长度取振幅更小，再取起点更早。",
        f"- ≥{args.box_min_days}日判为窄箱体。箱体内位置保持定义为：MA5—MA60全部有效观测至少80%不低于MA120，全部不低于MA120下方2%。",
        "- A：窄箱体＋位置保持；B：窄箱体但位置转弱；C：位置保持但箱体不足；D：两项均不足。箱体和最窄粘合可能使用事件后数据，只作事后描述。",
        "",
        "## 锚点覆盖",
        "",
        "| 锚点 | 原始锚点 | 左截断 | 可搜索 | 首个联合事件 | 窄箱体 | A层 |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ]
    for item in coverage.itertuples():
        lines.append(
            f"| MA{int(item.anchor_period)} | {int(item.anchor_count)} | {int(item.left_censored_count)} | {int(item.eligible_count)} | {int(item.first_qualifying_event_count)} | {int(item.narrow_box_count)} | {int(item.best_structure_count)} |"
        )
    lines.extend([
        "",
        "## 股票级结构分层",
        "",
        "| 层级 | 股票数 |",
        "|---|---:|",
    ])
    for tier in ("A_窄箱体且均线位置保持", "B_窄箱体但均线位置转弱", "C_均线位置保持但箱体不足", "D_箱体与位置均不足"):
        lines.append(f"| {tier} | {int(tier_counts.get(tier, 0))} |")
    lines.extend([
        "",
        "## 阈值敏感性",
        "",
        "每一行都重新按日期搜索首个同时满足位置，不能在主样本上简单过滤。",
        "",
        "| 类型 | 阈值 | 事件数 | 窄箱体数 | A层数 | 跳过较早回调数 | 实际箱体阈值 |",
        "|---|---:|---:|---:|---:|---:|---:|",
    ])
    for item in sensitivity.itertuples():
        lines.append(
            f"| {item.sensitivity_type} | {float(item.threshold):.2f} | {int(item.event_count)} | {int(item.narrow_box_count)} | {int(item.best_structure_count)} | {int(item.skipped_earlier_first_touch_count)} | {float(item.box_range_threshold_pct):.2f}% |"
        )
    lines.extend([
        "",
        "## 数据与边界",
        "",
        f"- 日线输入：`{Path(args.daily_file)}`；前复权价格与均线，原始成交量。",
        "- 行情截止：2026-09-30，因此仅表述为截至该日或2026 YTD。停牌和缺失只按实际存在的交易日序列处理。",
        "- 候选池由截至样本日事后观察到的MA60/MA120超长连续上涨段筛选，存在条件样本选择偏差；本研究不是实时可执行策略或买入建议。",
        "",
        "## 产物",
        "",
        "- `unique_stock_first_qualifying_events.csv`：60只股票级主样本与全部指标。",
        "- `anchor_first_qualifying_events.csv`：锚点级首个联合事件。",
        "- `unique_stock_raw_first_ma120_touches.csv`：真正首次回调对照。",
        "- `threshold_sensitivity.csv`、`metric_summary.csv`、`coverage_summary.csv`。",
    ])
    if pptx_name:
        lines.append(f"- `{pptx_name}`：完整PPT。")
    if pdf_name:
        lines.append(f"- `{pdf_name}`：对应PDF。")
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


def verify_joint_outputs(
    *,
    result: dict[str, Any],
    sensitivity: pd.DataFrame,
    chart_paths: dict[str, Path],
    pptx_path: Path | None,
    pdf_path: Path | None,
    rendered: list[Path],
) -> dict[str, Any]:
    events = result["unique_events"]
    raw_unique = result["unique_raw_first_touches"]
    assert len(events) == events["ts_code"].nunique()
    assert len(raw_unique) == raw_unique["ts_code"].nunique()
    assert int(raw_unique["strict_first_touch_qualified"].sum()) == int(result["summary"].iloc[0]["strict_first_touch_qualified_stocks"])
    assert bool(events["skipped_earlier_first_touch"].eq(events["delay_from_first_touch_trading_days"].gt(0)).all())
    assert bool(events["delay_from_first_touch_trading_days"].ge(0).all())
    assert bool(events["six_ma_cohesion_position_pass"].all())
    assert int(events["narrow_box_pass"].sum()) == int(result["summary"].iloc[0]["narrow_box_stocks"])
    assert int((events["narrow_box_pass"] & events["box_position_descriptive_pass"]).sum()) == int(result["summary"].iloc[0]["best_structure_stocks"])
    assert set(chart_paths) == set(events["event_key"].astype(str))
    assert len({_sha256(path) for path in chart_paths.values()}) == len(chart_paths)
    main_sensitivity = sensitivity.loc[
        sensitivity["sensitivity_type"].eq("six_ma_spread_threshold")
        & sensitivity["threshold"].eq(6.0)
    ]
    assert len(main_sensitivity) == 1
    assert int(main_sensitivity.iloc[0]["event_count"]) == len(events)
    for event in events.itertuples():
        frame = result["features_by_code"][str(event.ts_code)]
        selected = base.select_first_qualifying_touch(
            frame,
            anchor_start_date=event.anchor_start_date,
            anchor_end_date=event.anchor_end_date,
        )
        assert selected == int(event.touch_position)
        assert base._compact_date(frame.iloc[selected]["trade_date"]) == str(event.touch_date)
    verification: dict[str, Any] = {
        "unique_event_count": int(len(events)),
        "raw_first_touch_stock_count": int(len(raw_unique)),
        "strict_first_touch_qualified_count": int(raw_unique["strict_first_touch_qualified"].sum()),
        "chart_count": int(len(chart_paths)),
        "chart_hashes_unique": True,
    }
    if pptx_path is not None:
        deck = Presentation(pptx_path)
        expected_slides = 11 + len(events)
        assert len(deck.slides) == expected_slides
        detail_slides = list(deck.slides)[-len(events):]
        for slide, code in zip(detail_slides, events["ts_code"].astype(str)):
            text = "\n".join(shape.text for shape in slide.shapes if hasattr(shape, "text"))
            assert code in text
        picture_count = sum(
            1 for slide in deck.slides for shape in slide.shapes
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
        )
        assert picture_count == len(events) + 1
        verification.update({
            "pptx_slide_count": int(len(deck.slides)),
            "pptx_picture_count": int(picture_count),
            "detail_slide_stock_coverage": True,
        })
    if pdf_path is not None:
        import pymupdf
        document = pymupdf.open(pdf_path)
        pdf_pages = len(document)
        document.close()
        assert pdf_pages == 11 + len(events)
        assert len(rendered) == pdf_pages
        verification.update({"pdf_page_count": int(pdf_pages), "rendered_slide_count": int(len(rendered))})
    return verification


def main(argv: Sequence[str] | None = None) -> int:
    args = _parse_args(argv)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    charts_dir = output_dir / "charts"
    charts_dir.mkdir(parents=True, exist_ok=True)
    for stale in charts_dir.glob("*.png"):
        stale.unlink()
    base._register_fonts()

    daily = pd.read_csv(args.daily_file, compression="infer", dtype={"ts_code": str, "trade_date": str})
    daily["trade_date"] = daily["trade_date"].map(base._compact_date)
    anchors = base.load_anchors(args)
    ma_frames = prepare_ma_frames(daily)
    result = scan_joint_events(anchors=anchors, ma_frames=ma_frames, args=args)
    sensitivity = build_joint_sensitivity(
        anchors=anchors,
        ma_frames=ma_frames,
        args=args,
        main_result=result,
    )

    csv_outputs = {
        "anchor_eligibility.csv": result["eligibility"],
        "anchor_raw_first_ma120_touches.csv": result["raw_first_touches"],
        "unique_stock_raw_first_ma120_touches.csv": result["unique_raw_first_touches"],
        "anchor_first_qualifying_events.csv": result["anchor_events"],
        "unique_stock_first_qualifying_events.csv": result["unique_events"],
        "unique_stock_best_structure_events.csv": result["best_events"],
        "coverage_summary.csv": result["coverage"],
        "metric_summary.csv": result["summary"],
        "threshold_sensitivity.csv": sensitivity,
    }
    for filename, frame in csv_outputs.items():
        frame.to_csv(output_dir / filename, index=False, encoding="utf-8-sig")

    pptx_path: Path | None = None
    pdf_path: Path | None = None
    rendered: list[Path] = []
    contacts: list[Path] = []
    chart_paths: dict[str, Path] = {}
    distribution_chart = output_dir / "joint_event_distributions.png"
    if not args.metrics_only:
        for _, event in result["unique_events"].iterrows():
            path = charts_dir / f"{int(event['display_order']):02d}_{str(event['ts_code']).replace('.', '_')}_{event['touch_date']}.png"
            create_joint_event_chart(result["features_by_code"][str(event["ts_code"])], event, path)
            chart_paths[str(event["event_key"])] = path
        create_joint_distribution_chart(result["unique_events"], distribution_chart)
        pptx_path = output_dir / "首个同时满足回调MA120与六均线粘合_窄箱体研究.pptx"
        build_joint_deck(
            events=result["unique_events"],
            anchor_events=result["anchor_events"],
            summary=result["summary"],
            coverage=result["coverage"],
            sensitivity=sensitivity,
            chart_paths=chart_paths,
            distribution_chart=distribution_chart,
            output=pptx_path,
        )
        pdf_path = base.convert_ppt_to_pdf(pptx_path, output_dir)
        rendered = base.render_pdf(pdf_path, output_dir / "rendered_slides")
        contacts = base.create_contact_sheets(rendered, output_dir / "contact_sheets")

    report_path = output_dir / "report.md"
    write_joint_report(
        output=report_path,
        args=args,
        result=result,
        sensitivity=sensitivity,
        pptx_name=pptx_path.name if pptx_path else None,
        pdf_name=pdf_path.name if pdf_path else None,
    )
    verification = verify_joint_outputs(
        result=result,
        sensitivity=sensitivity,
        chart_paths=chart_paths if not args.metrics_only else {
            str(event["event_key"]): charts_dir / f"{int(event['display_order']):02d}_{str(event['ts_code']).replace('.', '_')}_{event['touch_date']}.png"
            for _, event in result["unique_events"].iterrows()
            if (charts_dir / f"{int(event['display_order']):02d}_{str(event['ts_code']).replace('.', '_')}_{event['touch_date']}.png").exists()
        },
        pptx_path=pptx_path,
        pdf_path=pdf_path,
        rendered=rendered,
    ) if not args.metrics_only else {
        "unique_event_count": int(len(result["unique_events"])),
        "raw_first_touch_stock_count": int(len(result["unique_raw_first_touches"])),
        "strict_first_touch_qualified_count": int(result["unique_raw_first_touches"]["strict_first_touch_qualified"].sum()),
    }

    artifact_paths = [output_dir / name for name in csv_outputs] + [report_path]
    artifact_paths.extend(chart_paths.values())
    if distribution_chart.exists(): artifact_paths.append(distribution_chart)
    if pptx_path is not None: artifact_paths.append(pptx_path)
    if pdf_path is not None: artifact_paths.append(pdf_path)
    artifact_paths.extend(rendered)
    artifact_paths.extend(contacts)
    metadata = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "study": "first_joint_ma120_pullback_and_six_ma_cohesion",
        "selection_semantics": "Within each eligible anchor, choose the first date simultaneously satisfying MA120 pullback, six-MA cohesion, and event-day MA-position constraints; earlier nonqualifying touches may be skipped.",
        "warning": "This is not the true first MA120 pullback. Box metrics may use up to 30 future trading days, and the long-trend candidate pool is selected ex post.",
        "data": {
            "daily_file": str(Path(args.daily_file).resolve()),
            "ma60_anchor_file": str(Path(args.ma60_file).resolve()),
            "ma120_anchor_file": str(Path(args.ma120_file).resolve()),
            "adjustment": "qfq=True for prices and moving averages; raw volume",
            "observed_end_date": "20260930",
            "actual_daily_rows": int(len(daily)),
            "candidate_stock_count": int(daily["ts_code"].nunique()),
        },
        "definitions": {
            "anchor_eligibility": "exclude every left-censored MA60/MA120 anchor",
            "moving_averages": "compute MA5/10/20/30/60/120 on each stock's complete available qfq close history before slicing anchors",
            "ma120_rising_streak": "continue only when MA120_t > MA120_(t-1); reset to zero on flat, down, or non-comparable observations",
            "first_qualifying_selection": "within each eligible anchor, scan chronologically and select the first date jointly satisfying pullback, cohesion, and event-day MA-position rules",
            "raw_first_touch_comparison": "independently select the first pullback candidate without cohesion; the joint event may skip it",
            "box_position_pass": "at least 80% of valid MA5-MA60 observations are at or above MA120 and 100% are at or above MA120 minus 2%",
            "future_data_boundary": "box duration, minimum six-MA spread, and box-period MA-position statistics may use up to 30 trading days after the event",
            "candidate_pool_bias": "the MA60/MA120 long-trend pool was selected using ex-post observed streak duration",
        },
        "thresholds": {
            "extension_lookback_trading_days": int(args.extension_lookback),
            "prior_extension_pct": float(args.prior_extension_pct),
            "touch_upper_pct": float(args.touch_upper_pct),
            "close_floor_pct": float(args.close_floor_pct),
            "six_ma_cohesion_threshold_pct": float(args.cohesion_threshold_pct),
            "min_short_above_count": int(args.min_short_above_count),
            "max_short_below_ma120_pct": float(args.max_short_below_ma120_pct),
            "box_lookback_trading_days": int(args.box_lookback),
            "box_lookforward_trading_days": int(args.box_lookforward),
            "box_max_range_pct": float(args.box_max_range_pct),
            "box_min_days": int(args.box_min_days),
        },
        "tie_breaking": {
            "box": "longest interval, then smaller range, then earlier start",
            "stock": "earliest selected date, then anchor_period, then anchor_start_date",
        },
        "counts": result["summary"].iloc[0].to_dict(),
        "verification": verification,
        "artifacts": [
            {"path": str(path.resolve()), "sha256": _sha256(path), "bytes": int(path.stat().st_size)}
            for path in artifact_paths if path.exists()
        ],
    }
    metadata_path = output_dir / "metadata.json"
    metadata_path.write_text(json.dumps(base._clean_json(metadata), ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps({
        "output_dir": str(output_dir.resolve()),
        "anchor_events": int(len(result["anchor_events"])),
        "unique_events": int(len(result["unique_events"])),
        "narrow_box_events": int(result["unique_events"]["narrow_box_pass"].sum()),
        "best_structure_events": int(len(result["best_events"])),
        "skipped_earlier": int(result["unique_events"]["skipped_earlier_first_touch"].sum()),
        "pptx": str(pptx_path.resolve()) if pptx_path else None,
        "pdf": str(pdf_path.resolve()) if pdf_path else None,
        "slides": verification.get("pptx_slide_count"),
    }, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

"""Build a combined MA60/MA120 nominal Top-100 themed K-line atlas."""

from __future__ import annotations

import argparse
from datetime import datetime, timedelta
import hashlib
import json
import math
from pathlib import Path
import sys
from typing import Any, Iterable, Sequence

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches

from build_ma60_top1pct_kline_ppt import (
    ALLOWED_BROAD_CATEGORIES,
    CATEGORY_COLORS,
    CATEGORY_ORDER,
    _add_pill,
    _add_rect,
    _add_text,
    _set_slide_background,
    add_moving_averages,
)


SPAN_COLORS = {60: "#2563EB", 120: "#7C3AED"}


def price_basis_label(qfq: bool) -> str:
    return "前复权" if qfq else "未复权"


def _compact_date(value: Any) -> str:
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if len(text) == 8 and text.isdigit():
        return text
    parsed = pd.Timestamp(text)
    return parsed.strftime("%Y%m%d")


def display_start_for_spans(spans: Sequence[dict[str, Any]]) -> pd.Timestamp:
    """Keep episode context without showing dates before the 2023 analysis history."""
    if not spans:
        raise ValueError("at least one episode span is required")
    earliest = min(
        pd.to_datetime(_compact_date(span["start_date"]), format="%Y%m%d")
        for span in spans
    )
    contextual = min(earliest - timedelta(days=90), pd.Timestamp("2025-10-01"))
    return max(contextual, pd.Timestamp("2023-01-01"))


def episode_spans_from_row(row: pd.Series) -> list[dict[str, Any]]:
    """Extract period-specific episode spans from a union row."""
    spans: list[dict[str, Any]] = []
    for period in (60, 120):
        if not bool(row.get(f"in_ma{period}_top100", False)):
            continue
        spans.append(
            {
                "period": period,
                "start_date": _compact_date(row[f"ma{period}_start_date"]),
                "end_date": _compact_date(row[f"ma{period}_end_date"]),
                "target_days": int(row[f"ma{period}_target_days"]),
            }
        )
    return spans


def materialize_reviewed_classifications(
    evidence: Sequence[dict[str, Any]],
    reviewed: Sequence[dict[str, Any]],
) -> list[dict[str, Any]]:
    """Validate analyst review and attach verbatim local evidence labels."""
    evidence_by_code = {str(item["ts_code"]): item for item in evidence}
    reviewed_by_code = {str(item["ts_code"]): item for item in reviewed}
    if len(reviewed_by_code) != len(reviewed):
        raise ValueError("classification coverage contains duplicate ts_code")
    if set(evidence_by_code) != set(reviewed_by_code):
        raise ValueError(
            "classification coverage mismatch: "
            f"missing={sorted(set(evidence_by_code) - set(reviewed_by_code))}, "
            f"extra={sorted(set(reviewed_by_code) - set(evidence_by_code))}"
        )
    output: list[dict[str, Any]] = []
    ordered_evidence = sorted(
        evidence,
        key=lambda item: int(item.get("observation_rank", item.get("rank", 0))),
    )
    for item in ordered_evidence:
        code = str(item["ts_code"])
        classification = reviewed_by_code[code]
        category = str(classification["broad_category"])
        if category not in ALLOWED_BROAD_CATEGORIES:
            raise ValueError(f"invalid broad category for {code}: {category}")
        confidence = str(classification.get("confidence", "low"))
        if confidence not in {"high", "medium", "low"}:
            raise ValueError(f"invalid confidence for {code}: {confidence}")
        labels: list[str] = []
        for source in ("kpl_list_themes", "concepts", "limit_performance_themes"):
            for record in item.get(source, []):
                label = str(record.get("theme", "")).strip()
                if label and label not in labels:
                    labels.append(label)
        output.append(
            {
                "ts_code": code,
                "primary_theme": str(classification["primary_theme"]),
                "broad_category": category,
                "secondary_themes": list(classification.get("secondary_themes", []))[:3],
                "evidence_labels": labels[:8],
                "confidence": confidence,
                "evidence_note": str(
                    classification.get("rationale", classification.get("evidence_note", ""))
                ),
            }
        )
    return output


def render_dual_period_chart(
    daily: pd.DataFrame,
    output: Path,
    *,
    stock_name: str,
    ts_code: str,
    episode_spans: Sequence[dict[str, Any]],
    display_start: pd.Timestamp,
    display_end: pd.Timestamp,
    price_basis: str = "未复权",
) -> dict[str, Any]:
    """Render daily candles, four moving averages and MA60/120 spans."""
    if price_basis not in {"前复权", "未复权"}:
        raise ValueError(f"unsupported price basis: {price_basis}")
    required = {
        "trade_date", "open", "high", "low", "close", "vol",
        "ma20", "ma30", "ma60", "ma120",
    }
    if missing := required.difference(daily.columns):
        raise ValueError(f"chart data missing columns: {sorted(missing)}")
    frame = daily.copy()
    frame["date"] = pd.to_datetime(frame["trade_date"].astype(str), format="%Y%m%d")
    frame = frame.loc[frame["date"].between(display_start, display_end)].copy()
    frame = frame.sort_values("date").drop_duplicates("date", keep="last").reset_index(drop=True)
    if frame.empty:
        raise ValueError(f"no chart bars for {ts_code}")
    if frame[["open", "high", "low", "close", "vol"]].isna().any().any():
        raise ValueError(f"missing OHLCV values for {ts_code}")

    plt.rcParams["font.sans-serif"] = [
        "Noto Sans CJK SC", "PingFang SC", "Heiti SC", "Arial Unicode MS", "DejaVu Sans"
    ]
    plt.rcParams["axes.unicode_minus"] = False
    figure, (price_ax, volume_ax) = plt.subplots(
        2, 1, figsize=(14.2, 7.2), sharex=True,
        gridspec_kw={"height_ratios": [4.25, 1.15], "hspace": 0.03},
        facecolor="white",
    )
    for axis in (price_ax, volume_ax):
        axis.set_facecolor("#FBFCFE")
        axis.grid(True, color="#CBD5E1", alpha=0.38, linewidth=0.65)
        axis.spines[["top", "right"]].set_visible(False)
        axis.tick_params(labelsize=9, colors="#334155")

    x = np.arange(len(frame))
    up = frame["close"].gt(frame["open"])
    down = frame["close"].lt(frame["open"])
    colors = np.where(up, "#D9363E", np.where(down, "#159D72", "#64748B"))
    price_ax.vlines(x, frame["low"], frame["high"], color=colors, linewidth=0.72, zorder=2)
    body_bottom = np.minimum(frame["open"], frame["close"])
    body_height = np.abs(frame["close"] - frame["open"])
    min_body = max(float(frame["close"].median()) * 0.0008, 1e-6)
    for index, (bottom, height, color) in enumerate(zip(body_bottom, body_height, colors)):
        price_ax.add_patch(
            Rectangle(
                (index - 0.31, float(bottom)), 0.62, max(float(height), min_body),
                facecolor=color, edgecolor=color, linewidth=0.45, zorder=3,
            )
        )
    ma_colors = {20: "#F59E0B", 30: "#0D9488", 60: "#2563EB", 120: "#111827"}
    for period, color in ma_colors.items():
        price_ax.plot(
            x, frame[f"ma{period}"], label=f"MA{period}", color=color,
            linewidth=1.4 if period in {60, 120} else 1.0, alpha=0.95, zorder=4,
        )
    volume_ax.bar(x, frame["vol"] / 10000.0, width=0.68, color=colors, alpha=0.82)

    present_periods: list[int] = []
    for span in episode_spans:
        period = int(span["period"])
        present_periods.append(period)
        mask = frame["date"].between(
            pd.to_datetime(_compact_date(span["start_date"]), format="%Y%m%d"),
            pd.to_datetime(_compact_date(span["end_date"]), format="%Y%m%d"),
        )
        if not mask.any():
            continue
        indexes = np.flatnonzero(mask.to_numpy())
        left, right = int(indexes[0]) - 0.5, int(indexes[-1]) + 0.5
        color = SPAN_COLORS[period]
        for axis in (price_ax, volume_ax):
            axis.axvspan(left, right, color=color, alpha=0.085, zorder=0)
            axis.axvline(left + 0.5, color=color, linestyle="--", linewidth=0.9, alpha=0.8)
            axis.axvline(right - 0.5, color=color, linestyle="--", linewidth=0.9, alpha=0.8)

    price_ax.set_ylabel(f"{price_basis}价格（元）", fontsize=10, color="#334155")
    volume_ax.set_ylabel("成交量（万手）", fontsize=10, color="#334155")
    price_ax.legend(
        loc="upper left", ncol=4, fontsize=9, frameon=True,
        facecolor="white", edgecolor="#E2E8F0", framealpha=0.92,
    )
    price_ax.set_title(
        f"{stock_name}  {ts_code}｜日K + 成交量", loc="left", fontsize=15,
        fontweight="bold", color="#0F172A", pad=10,
    )
    tick_count = min(10, len(frame))
    ticks = np.unique(np.linspace(0, len(frame) - 1, tick_count, dtype=int))
    volume_ax.set_xticks(ticks)
    volume_ax.set_xticklabels(
        [frame.iloc[index]["date"].strftime("%Y-%m") for index in ticks], fontsize=9
    )
    price_ax.set_xlim(-1, len(frame))
    legend_bits = []
    if 60 in present_periods:
        legend_bits.append("蓝色区间/虚线：MA60入选段")
    if 120 in present_periods:
        legend_bits.append("紫色区间/虚线：MA120入选段")
    figure.text(
        0.995, 0.012,
        "红涨绿跌｜" + "｜".join(legend_bits) + f"｜价格与成交量均{price_basis}",
        ha="right", va="bottom", fontsize=8.3, color="#64748B",
    )
    figure.subplots_adjust(left=0.065, right=0.985, top=0.91, bottom=0.105)
    output.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(output, dpi=150, facecolor="white", bbox_inches="tight")
    plt.close(figure)
    return {
        "ts_code": ts_code,
        "image": output.name,
        "bars": int(len(frame)),
        "first_date": frame.iloc[0]["date"].strftime("%Y%m%d"),
        "last_date": frame.iloc[-1]["date"].strftime("%Y%m%d"),
        "periods": sorted(set(present_periods)),
        "price_basis": price_basis,
    }


def _category_sequence(stocks: pd.DataFrame) -> list[str]:
    present = set(stocks["broad_category"].astype(str))
    return [category for category in CATEGORY_ORDER if category in present] + sorted(
        present - set(CATEGORY_ORDER)
    )


def _add_slide_number(slide, number: int, *, dark: bool = False) -> None:
    _add_text(
        slide,
        f"MA60 / MA120 NOMINAL TOP100 · 2026 YTD                         {number:03d}",
        0.38,
        7.16,
        12.55,
        0.22,
        size=8,
        color="#94A3B8" if dark else "#64748B",
        align=PP_ALIGN.RIGHT,
        margin=0,
    )


def _date_text(value: Any) -> str:
    text = _compact_date(value)
    return f"{text[:4]}-{text[4:6]}-{text[6:8]}"


def _period_summary(row: pd.Series, period: int) -> str:
    if not bool(row[f"in_ma{period}_top100"]):
        return f"MA{period}：未入选"
    rank = int(row[f"ma{period}_rank"])
    days = int(row[f"ma{period}_target_days"])
    status = "仍延续" if bool(row[f"ma{period}_right_censored"]) else "已结束"
    carry = " · 跨年" if bool(row[f"ma{period}_carried_in"]) else ""
    return (
        f"MA{period} #{rank} · {days}d · "
        f"{_date_text(row[f'ma{period}_start_date'])} → "
        f"{_date_text(row[f'ma{period}_end_date'])} · {status}{carry}"
    )


def build_combined_presentation(
    stocks: pd.DataFrame,
    chart_paths: dict[str, Path],
    output: Path,
    *,
    data_end: str,
    ma60_cutoff: int,
    ma120_cutoff: int,
    price_basis: str = "未复权",
) -> dict[str, int]:
    """Build one deduplicated, theme-grouped deck for both ranked lists."""
    if price_basis not in {"前复权", "未复权"}:
        raise ValueError(f"unsupported price basis: {price_basis}")
    required = {
        "observation_rank", "ts_code", "stock_name", "industry", "membership",
        "in_ma60_top100", "in_ma120_top100", "ma60_rank", "ma60_start_date",
        "ma60_end_date", "ma60_target_days", "ma60_carried_in",
        "ma60_right_censored", "ma120_rank", "ma120_start_date",
        "ma120_end_date", "ma120_target_days", "ma120_carried_in",
        "ma120_right_censored", "primary_theme", "broad_category",
        "secondary_themes", "evidence_labels", "confidence", "evidence_note",
    }
    if missing := required.difference(stocks.columns):
        raise ValueError(f"presentation data missing columns: {sorted(missing)}")
    codes = set(stocks["ts_code"].astype(str))
    if codes != set(chart_paths):
        raise ValueError(
            f"chart coverage mismatch: missing={sorted(codes - set(chart_paths))}, "
            f"extra={sorted(set(chart_paths) - codes)}"
        )
    for path in chart_paths.values():
        if not Path(path).exists():
            raise FileNotFoundError(path)

    ordered = stocks.sort_values("observation_rank").reset_index(drop=True)
    categories = _category_sequence(ordered)
    counts = ordered["broad_category"].value_counts().to_dict()
    ma60_count = int(ordered["in_ma60_top100"].sum())
    ma120_count = int(ordered["in_ma120_top100"].sum())
    overlap = int((ordered["in_ma60_top100"] & ordered["in_ma120_top100"]).sum())
    deck = Presentation()
    deck.slide_width = Inches(13.333)
    deck.slide_height = Inches(7.5)
    blank = deck.slide_layouts[6]

    # Cover
    slide = deck.slides.add_slide(blank)
    _set_slide_background(slide, "#0B1220")
    _add_rect(slide, 0, 0, 0.16, 7.5, fill="#2563EB")
    _add_rect(slide, 0.16, 0, 0.10, 7.5, fill="#7C3AED")
    _add_text(
        slide, "MA60 / MA120 上升持续性\n观察名单图册", 0.82, 1.1, 8.5, 1.75,
        size=34, color="#F8FAFC", bold=True, valign=MSO_ANCHOR.MIDDLE,
    )
    _add_text(
        slide, f"名义前100 · 并列保留 · {len(ordered)}只去重股票 · 题材归类 · 日K与成交量",
        0.86, 3.12, 8.65, 0.62, size=17, color="#C4B5FD", bold=True,
    )
    _add_text(
        slide,
        f"数据截至 {_date_text(data_end)}  |  本地日线数据库  |  图表为{price_basis}价格",
        0.86, 3.86, 8.8, 0.38, size=12.5, color="#CBD5E1",
    )
    _add_rect(slide, 9.72, 0.96, 2.72, 5.35, fill="#111C32", radius=True, line="#263653")
    cover_stats = [
        ("MA60 入选", ma60_count, "#60A5FA"),
        ("MA120 入选", ma120_count, "#C4B5FD"),
        ("双榜重合", overlap, "#5EEAD4"),
        ("去重观察", len(ordered), "#FBBF24"),
    ]
    for index, (label, value, color) in enumerate(cover_stats):
        top = 1.35 + index * 1.18
        _add_text(slide, str(value), 10.05, top, 2.05, 0.58, size=30, color=color,
                  bold=True, align=PP_ALIGN.CENTER)
        _add_text(slide, label, 10.05, top + 0.57, 2.05, 0.28, size=11.5,
                  color="#94A3B8", align=PP_ALIGN.CENTER)
    _add_slide_number(slide, 1, dark=True)

    # Methodology
    slide = deck.slides.add_slide(blank)
    _set_slide_background(slide, "#F8FAFC")
    _add_text(slide, "筛选口径与图表口径", 0.58, 0.42, 8.2, 0.58, size=25, bold=True)
    stats = [
        ("MA60 第100名阈值", f"{ma60_cutoff}d", "#2563EB"),
        ("MA60 并列后", str(ma60_count), "#0EA5E9"),
        ("MA120 第100名阈值", f"{ma120_cutoff}d", "#7C3AED"),
        ("MA120 并列后", str(ma120_count), "#A855F7"),
    ]
    for index, (label, value, color) in enumerate(stats):
        left = 0.58 + index * 3.13
        _add_rect(slide, left, 1.27, 2.78, 1.15, fill="#FFFFFF", radius=True, line="#E2E8F0")
        _add_rect(slide, left, 1.27, 0.08, 1.15, fill=color)
        _add_text(slide, value, left + 0.24, 1.45, 2.25, 0.48, size=23, color=color, bold=True)
        _add_text(slide, label, left + 0.24, 1.96, 2.25, 0.28, size=10.5, color="#64748B")
    method = (
        "排名对象\n按股票排名；每只股票在对应均线的全部上升段中，取2026目标窗口内 target_days 最大的一段。\n\n"
        "并列规则\n名义前100；第100名 target_days 阈值同分全部纳入，因此两榜均可能超过100只。\n\n"
        f"价格口径\n排名与展示均使用{price_basis}OHLCV，并叠加MA20/30/60/120。\n\n"
        "区间标识\n蓝色为MA60入选段，紫色为MA120入选段；右截断表示截至数据库末日仍在上升。"
    )
    _add_rect(slide, 0.58, 2.78, 8.05, 3.66, fill="#FFFFFF", radius=True, line="#E2E8F0")
    _add_text(slide, method, 0.88, 3.02, 7.45, 3.18, size=13.3, color="#334155")
    caveat = (
        f"数据截止\n{_date_text(data_end)}\n\n"
        f"去重观察池\n{len(ordered)}只\n\n"
        f"双榜重合\n{overlap}只\n\n"
        "提示\n2026 YTD，非完整全年；MA120榜第100名与更多股票同为181个窗口交易日。"
    )
    _add_rect(slide, 8.92, 2.78, 3.82, 3.66, fill="#0F172A", radius=True)
    _add_text(slide, caveat, 9.25, 3.03, 3.18, 3.15, size=13.2, color="#E2E8F0")
    _add_slide_number(slide, 2)

    # Membership overview
    slide = deck.slides.add_slide(blank)
    _set_slide_background(slide, "#F8FAFC")
    _add_text(slide, "两份观察名单的重合关系", 0.58, 0.42, 8.5, 0.58, size=25, bold=True)
    ma60_only = ma60_count - overlap
    ma120_only = ma120_count - overlap
    _add_rect(slide, 1.05, 1.48, 5.4, 4.75, fill="#DBEAFE", radius=True, line="#93C5FD")
    _add_rect(slide, 6.0, 1.48, 5.4, 4.75, fill="#EDE9FE", radius=True, line="#C4B5FD")
    _add_text(slide, "MA60", 1.42, 1.86, 2.0, 0.48, size=24, color="#1D4ED8", bold=True)
    _add_text(slide, "MA120", 9.02, 1.86, 1.8, 0.48, size=24, color="#6D28D9", bold=True,
              align=PP_ALIGN.RIGHT)
    overview = [
        (2.03, "仅MA60", ma60_only, "#2563EB"),
        (5.22, "双榜重合", overlap, "#0F766E"),
        (8.42, "仅MA120", ma120_only, "#7C3AED"),
    ]
    for left, label, value, color in overview:
        _add_rect(slide, left, 2.75, 2.85, 2.1, fill="#FFFFFF", radius=True, line="#D8E1EC")
        _add_text(slide, str(value), left + 0.22, 3.0, 2.4, 0.78, size=40, color=color,
                  bold=True, align=PP_ALIGN.CENTER)
        _add_text(slide, label, left + 0.22, 3.85, 2.4, 0.4, size=13, color="#64748B",
                  bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, f"合并后共 {len(ordered)} 只，个股页只呈现一次，并同时标注两榜资格。",
              1.4, 5.5, 10.2, 0.48, size=15, color="#334155", bold=True,
              align=PP_ALIGN.CENTER)
    _add_slide_number(slide, 3)

    # Theme distribution
    slide = deck.slides.add_slide(blank)
    _set_slide_background(slide, "#F8FAFC")
    _add_text(slide, "合并观察池题材结构", 0.58, 0.42, 8.0, 0.58, size=25, bold=True)
    _add_text(slide, "按主归类去重统计；个股页保留次题材与本地历史证据标签。",
              0.6, 0.98, 8.8, 0.34, size=11.5, color="#64748B")
    max_count = max(counts.values())
    spacing = min(0.42, 5.45 / max(len(categories), 1))
    for index, category in enumerate(categories):
        count = int(counts[category])
        top = 1.38 + index * spacing
        color = CATEGORY_COLORS.get(category, "#64748B")
        _add_text(slide, category, 0.62, top, 2.42, 0.27, size=10.6, color="#334155",
                  bold=True, valign=MSO_ANCHOR.MIDDLE)
        _add_rect(slide, 3.06, top + 0.02, 7.65 * count / max_count, 0.24, fill=color, radius=True)
        _add_text(slide, str(count), 10.87, top - 0.02, 0.6, 0.3, size=11.5,
                  color=color, bold=True, align=PP_ALIGN.RIGHT)
    _add_rect(slide, 11.55, 1.38, 1.15, 5.45, fill="#FFFFFF", radius=True, line="#E2E8F0")
    _add_text(slide, str(len(categories)), 11.64, 2.15, 0.98, 0.65, size=30,
              color="#5B6CFF", bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "个主归类", 11.65, 2.82, 0.96, 0.4, size=11,
              color="#64748B", align=PP_ALIGN.CENTER)
    _add_text(slide, str(len(ordered)), 11.64, 4.0, 0.98, 0.65, size=28,
              color="#0F766E", bold=True, align=PP_ALIGN.CENTER)
    _add_text(slide, "只去重股票", 11.63, 4.68, 1.0, 0.55, size=10.5,
              color="#64748B", align=PP_ALIGN.CENTER)
    _add_slide_number(slide, 4)

    # Ranking inventories, 28 entries per page.
    ranking_slides = 0
    for period, color in ((60, "#2563EB"), (120, "#7C3AED")):
        subset = ordered.loc[ordered[f"in_ma{period}_top100"]].sort_values(f"ma{period}_rank")
        pages = math.ceil(len(subset) / 28)
        for page in range(pages):
            slide = deck.slides.add_slide(blank)
            _set_slide_background(slide, "#F8FAFC")
            _add_text(slide, f"MA{period} 名义前100（并列保留）· {page + 1}/{pages}",
                      0.58, 0.4, 9.0, 0.56, size=23, bold=True)
            _add_pill(slide, f"{len(subset)}只", 10.7, 0.46, 1.45, color)
            page_rows = subset.iloc[page * 28:(page + 1) * 28]
            for position, (_, row) in enumerate(page_rows.iterrows()):
                column = position // 14
                row_index = position % 14
                left = 0.58 + column * 6.25
                top = 1.12 + row_index * 0.41
                rank = int(row[f"ma{period}_rank"])
                _add_text(slide, f"{rank:03d}", left, top, 0.52, 0.34, size=10,
                          color=color, bold=True, valign=MSO_ANCHOR.MIDDLE)
                _add_text(slide, f"{row['stock_name']}  {row['ts_code']}", left + 0.55, top, 2.55, 0.34,
                          size=10.4, bold=True, valign=MSO_ANCHOR.MIDDLE)
                _add_text(slide, str(row["primary_theme"]), left + 3.08, top, 2.25, 0.34,
                          size=9.8, color=CATEGORY_COLORS.get(row["broad_category"], "#64748B"),
                          bold=True, valign=MSO_ANCHOR.MIDDLE)
                _add_text(slide, f"{int(row[f'ma{period}_target_days'])}d", left + 5.35, top, 0.55, 0.34,
                          size=9.8, color="#334155", align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)
            _add_slide_number(slide, len(deck.slides))
            ranking_slides += 1

    # Deduplicated classification inventory, 24 entries per page.
    inventory_pages = math.ceil(len(ordered) / 24)
    for page in range(inventory_pages):
        slide = deck.slides.add_slide(blank)
        _set_slide_background(slide, "#F8FAFC")
        _add_text(slide, f"去重观察池题材归类 · {page + 1}/{inventory_pages}",
                  0.58, 0.4, 8.5, 0.56, size=23, bold=True)
        subset = ordered.iloc[page * 24:(page + 1) * 24]
        for position, (_, row) in enumerate(subset.iterrows()):
            column = position // 12
            row_index = position % 12
            left = 0.58 + column * 6.27
            top = 1.19 + row_index * 0.47
            color = CATEGORY_COLORS.get(row["broad_category"], "#64748B")
            _add_rect(slide, left, top + 0.04, 0.09, 0.33, fill=color, radius=True)
            _add_text(slide, f"{int(row['observation_rank']):03d}", left + 0.15, top, 0.48, 0.4,
                      size=10, color="#64748B", bold=True, valign=MSO_ANCHOR.MIDDLE)
            _add_text(slide, f"{row['stock_name']}  {row['ts_code']}", left + 0.66, top, 2.5, 0.4,
                      size=10.7, bold=True, valign=MSO_ANCHOR.MIDDLE)
            _add_text(slide, str(row["primary_theme"]), left + 3.08, top, 2.18, 0.4,
                      size=10.1, color=color, bold=True, valign=MSO_ANCHOR.MIDDLE)
            badge = "双榜" if row["membership"] == "MA60+MA120" else str(row["membership"])
            _add_text(slide, badge, left + 5.27, top, 0.7, 0.4, size=9.7, color="#334155",
                      bold=True, align=PP_ALIGN.RIGHT, valign=MSO_ANCHOR.MIDDLE)
        _add_slide_number(slide, len(deck.slides))

    # Theme sections and one page per unique stock.
    category_slides = 0
    stock_slides = 0
    for category_index, category in enumerate(categories, start=1):
        group = ordered.loc[ordered["broad_category"].eq(category)].sort_values("observation_rank")
        color = CATEGORY_COLORS.get(category, "#64748B")
        divider_pages = math.ceil(len(group) / 24)
        for divider_page in range(divider_pages):
            slide = deck.slides.add_slide(blank)
            _set_slide_background(slide, "#0B1220")
            _add_rect(slide, 0, 0, 0.18, 7.5, fill=color)
            _add_text(slide, f"{category_index:02d}", 0.72, 0.72, 1.3, 0.7,
                      size=31, color=color, bold=True)
            title = category if divider_pages == 1 else f"{category} · {divider_page + 1}/{divider_pages}"
            _add_text(slide, title, 0.72, 1.44, 9.8, 0.82, size=30, color="#FFFFFF", bold=True)
            _add_text(slide, f"{len(group)}只去重股票", 0.76, 2.37, 3.2, 0.4,
                      size=15, color="#CBD5E1")
            page_group = group.iloc[divider_page * 24:(divider_page + 1) * 24]
            midpoint = math.ceil(len(page_group) / 2)
            for column, chunk in enumerate((page_group.iloc[:midpoint], page_group.iloc[midpoint:])):
                lines = []
                for _, row in chunk.iterrows():
                    membership = "双榜" if row["membership"] == "MA60+MA120" else row["membership"]
                    lines.append(
                        f"{int(row['observation_rank']):03d}  {row['stock_name']} · "
                        f"{row['primary_theme']} · {membership}"
                    )
                _add_text(slide, "\n".join(lines), 0.8 + column * 6.15, 2.98, 5.72, 3.55,
                          size=11.8, color="#E2E8F0")
            _add_slide_number(slide, len(deck.slides), dark=True)
            category_slides += 1

        for _, row in group.iterrows():
            slide = deck.slides.add_slide(blank)
            _set_slide_background(slide, "#F1F5F9")
            _add_rect(slide, 0, 0, 0.13, 7.5, fill=color)
            _add_rect(slide, 0.36, 0.26, 0.7, 0.58, fill=color, radius=True)
            _add_text(slide, f"{int(row['observation_rank']):03d}", 0.37, 0.27, 0.68, 0.56,
                      size=16, color="#FFFFFF", bold=True, align=PP_ALIGN.CENTER,
                      valign=MSO_ANCHOR.MIDDLE, margin=0)
            _add_text(slide, f"{row['stock_name']}  {row['ts_code']}", 1.2, 0.27, 5.75, 0.47,
                      size=22, bold=True, valign=MSO_ANCHOR.MIDDLE)
            _add_pill(slide, category, 8.73, 0.34, 2.3, color)
            membership_color = "#0F766E" if row["membership"] == "MA60+MA120" else (
                "#2563EB" if row["membership"] == "MA60" else "#7C3AED"
            )
            _add_pill(slide, "双榜" if row["membership"] == "MA60+MA120" else str(row["membership"]),
                      11.18, 0.34, 1.16, membership_color)
            period_lines = [_period_summary(row, 60), _period_summary(row, 120)]
            _add_text(slide, "\n".join(period_lines), 1.22, 0.75, 11.2, 0.48,
                      size=9.4, color="#64748B")
            slide.shapes.add_picture(
                str(chart_paths[str(row["ts_code"])]), Inches(0.36), Inches(1.27),
                width=Inches(10.25), height=Inches(5.37),
            )
            _add_rect(slide, 10.79, 1.27, 2.2, 5.37, fill="#FFFFFF", radius=True, line="#DDE5EF")
            _add_text(slide, "主归类", 11.06, 1.53, 1.65, 0.25,
                      size=9.5, color="#94A3B8", bold=True)
            _add_text(slide, str(row["primary_theme"]), 11.03, 1.82, 1.7, 0.72,
                      size=14.5, color=color, bold=True)
            _add_text(slide, "次题材", 11.06, 2.68, 1.65, 0.25,
                      size=9.5, color="#94A3B8", bold=True)
            secondary = " / ".join(row["secondary_themes"]) or "—"
            _add_text(slide, secondary, 11.03, 2.97, 1.7, 0.86, size=10.5, color="#334155")
            _add_text(slide, "证据标签", 11.06, 3.95, 1.65, 0.25,
                      size=9.5, color="#94A3B8", bold=True)
            labels = " / ".join(row["evidence_labels"][:5]) or "行业兜底"
            _add_text(slide, labels, 11.03, 4.24, 1.7, 0.84, size=10.1, color="#475569")
            _add_text(slide, f"行业  {row['industry'] or '未标注'}", 11.04, 5.18, 1.72, 0.3,
                      size=9.8, color="#64748B", bold=True)
            confidence_map = {"high": "高", "medium": "中", "low": "低"}
            _add_text(slide, f"归类置信度  {confidence_map.get(row['confidence'], row['confidence'])}",
                      11.04, 5.53, 1.72, 0.3, size=9.8, color="#64748B", bold=True)
            _add_text(slide, "来源：本地日线 + 开盘啦历史题材；不构成投资建议。",
                      10.98, 6.04, 1.82, 0.45, size=8.0, color="#94A3B8")
            _add_slide_number(slide, len(deck.slides))
            stock_slides += 1

    output.parent.mkdir(parents=True, exist_ok=True)
    deck.save(output)
    return {
        "slides": len(deck.slides),
        "stock_slides": stock_slides,
        "category_slides": category_slides,
        "inventory_slides": inventory_pages,
        "ranking_slides": ranking_slides,
    }


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _write_report(stocks: pd.DataFrame, output: Path, metadata: dict[str, Any]) -> None:
    lines = [
        "# 2026 YTD MA60 / MA120 名义前100观察名单",
        "",
        f"- 数据截止：`{metadata['actual_target_end']}`",
        f"- MA60：名义前100，并列后 {metadata['ma60']['selected_with_ties']} 只，阈值 {metadata['ma60']['cutoff_target_days']} 天。",
        f"- MA120：名义前100，并列后 {metadata['ma120']['selected_with_ties']} 只，阈值 {metadata['ma120']['cutoff_target_days']} 天。",
        f"- 双榜重合：{metadata['overlap']} 只；去重观察池：{len(stocks)} 只。",
        f"- 排名与图表均使用{metadata['price_basis_zh']}日K和成交量。",
        "- 题材证据：本地开盘啦历史涨停题材、概念成分与行业兜底。",
        "",
    ]
    for category in _category_sequence(stocks):
        group = stocks.loc[stocks["broad_category"].eq(category)].sort_values("observation_rank")
        lines.extend(
            [
                f"## {category}（{len(group)}只）",
                "",
                "| 观察序 | 代码 | 名称 | 榜单 | 主归类 | MA60排名 | MA120排名 |",
                "|---:|---|---|---|---|---:|---:|",
            ]
        )
        for _, row in group.iterrows():
            rank60 = "—" if pd.isna(row["ma60_rank"]) else str(int(row["ma60_rank"]))
            rank120 = "—" if pd.isna(row["ma120_rank"]) else str(int(row["ma120_rank"]))
            lines.append(
                f"| {int(row['observation_rank'])} | {row['ts_code']} | {row['stock_name']} | "
                f"{row['membership']} | {row['primary_theme']} | {rank60} | {rank120} |"
            )
        lines.append("")
    output.write_text("\n".join(lines), encoding="utf-8")


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--analysis-dir",
        type=Path,
        default=Path("outputs/ma60_ma120_top100_themes/2026_through_20260930"),
    )
    parser.add_argument(
        "--classification-file",
        type=Path,
        default=Path(
            "outputs/ma60_ma120_top100_themes/2026_through_20260930/"
            "classifications_reviewed.json"
        ),
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("outputs/ma60_ma120_top100_ppt/2026_through_20260930"),
    )
    parser.add_argument("--daily-start", default="20220101")
    parser.add_argument(
        "--qfq",
        action="store_true",
        help="use forward-adjusted daily prices for every chart and moving average",
    )
    return parser.parse_args(argv)


def main() -> None:
    root = Path(__file__).resolve().parents[5]
    if str(root) not in sys.path:
        sys.path.insert(0, str(root))
    from coreClient.data_provider import get_day

    args = _parse_args()
    daily_qfq = bool(args.qfq)
    basis_label = price_basis_label(daily_qfq)
    args.output_dir.mkdir(parents=True, exist_ok=True)
    charts_dir = args.output_dir / "charts"
    charts_dir.mkdir(parents=True, exist_ok=True)
    selection_metadata = json.loads(
        (args.analysis_dir / "metadata.json").read_text(encoding="utf-8")
    )
    expected_basis_token = f"qfq={daily_qfq}"
    ranking_basis = str(selection_metadata.get("ranking_price_basis", ""))
    if expected_basis_token not in ranking_basis:
        raise ValueError(
            "ranking/chart price-basis mismatch: "
            f"ranking={ranking_basis!r}, chart={expected_basis_token}"
        )
    union = pd.read_csv(
        args.analysis_dir / "observation_union.csv", dtype={"ts_code": str}
    )
    evidence = json.loads(
        (args.analysis_dir / "theme_evidence_summary.json").read_text(encoding="utf-8")
    )
    reviewed_payload = json.loads(args.classification_file.read_text(encoding="utf-8"))
    reviewed = (
        reviewed_payload["classifications"]
        if isinstance(reviewed_payload, dict)
        else reviewed_payload
    )
    classifications = materialize_reviewed_classifications(evidence, reviewed)
    evidence_details = pd.DataFrame(
        [
            {
                "ts_code": str(item["ts_code"]),
                "stock_name": str(item.get("stock_name", "")),
                "industry": str(item.get("industry", "")),
            }
            for item in evidence
        ]
    )
    stocks = union.merge(
        evidence_details, on="ts_code", how="left", validate="one_to_one"
    ).merge(
        pd.DataFrame(classifications), on="ts_code", how="left", validate="one_to_one"
    )
    expected_union = int(selection_metadata["union"])
    if len(stocks) != expected_union or stocks["ts_code"].nunique() != expected_union:
        raise ValueError(
            f"expected {expected_union} unique observation stocks, got "
            f"rows={len(stocks)}, unique={stocks['ts_code'].nunique()}"
        )
    if stocks[["stock_name", "industry", "primary_theme", "broad_category"]].isna().any().any():
        raise ValueError("missing stock identity or classification fields")
    stocks = stocks.sort_values("observation_rank").reset_index(drop=True)

    classification_records = json.loads(stocks.to_json(orient="records", force_ascii=False))
    (args.output_dir / "classifications.json").write_text(
        json.dumps(classification_records, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    classification_csv = stocks.copy()
    for column in ("secondary_themes", "evidence_labels"):
        classification_csv[column] = classification_csv[column].map(
            lambda values: " | ".join(values)
        )
    classification_csv.to_csv(
        args.output_dir / "classifications.csv", index=False, encoding="utf-8-sig"
    )

    codes = stocks["ts_code"].astype(str).tolist()
    data_end = str(selection_metadata["actual_target_end"])
    daily = get_day(
        ts_codes=codes,
        start_date=args.daily_start,
        end_date=data_end,
        qfq=daily_qfq,
        source="database_only",
    )
    if daily.empty:
        raise ValueError("local daily database returned no rows")
    daily["ts_code"] = daily["ts_code"].astype(str)
    actual_codes = set(daily["ts_code"])
    if set(codes) != actual_codes:
        raise ValueError(
            f"daily coverage mismatch: missing={sorted(set(codes) - actual_codes)}, "
            f"extra={sorted(actual_codes - set(codes))}"
        )
    if daily.duplicated(["ts_code", "trade_date"]).any():
        raise ValueError("duplicate daily ts_code/trade_date rows")
    daily = add_moving_averages(daily)
    daily_input_name = f"daily_input_{'qfq' if daily_qfq else 'unadjusted'}.csv.gz"
    daily.to_csv(
        args.output_dir / daily_input_name, index=False, compression="gzip"
    )

    chart_paths: dict[str, Path] = {}
    receipts: list[dict[str, Any]] = []
    for _, row in stocks.iterrows():
        code = str(row["ts_code"])
        spans = episode_spans_from_row(row)
        image_path = charts_dir / f"{int(row['observation_rank']):03d}_{code.replace('.', '_')}.png"
        receipt = render_dual_period_chart(
            daily.loc[daily["ts_code"].eq(code)].copy(),
            image_path,
            stock_name=str(row["stock_name"]),
            ts_code=code,
            episode_spans=spans,
            display_start=display_start_for_spans(spans),
            display_end=pd.Timestamp(data_end),
            price_basis=basis_label,
        )
        receipt["observation_rank"] = int(row["observation_rank"])
        receipt["sha256"] = _sha256(image_path)
        receipts.append(receipt)
        chart_paths[code] = image_path
        print(
            f"[{int(row['observation_rank']):03d}/{len(stocks)}] "
            f"{row['stock_name']} {code}",
            flush=True,
        )

    basis_file_tag = "QFQ" if daily_qfq else "Unadjusted"
    deck_path = args.output_dir / (
        f"MA60_MA120_Top100_{len(stocks)}stocks_2026YTD_{basis_file_tag}_Kline_Atlas.pptx"
    )
    deck_receipt = build_combined_presentation(
        stocks,
        chart_paths,
        deck_path,
        data_end=data_end,
        ma60_cutoff=int(selection_metadata["ma60"]["cutoff_target_days"]),
        ma120_cutoff=int(selection_metadata["ma120"]["cutoff_target_days"]),
        price_basis=basis_label,
    )
    reopened = Presentation(deck_path)
    embedded_images = sum(
        1
        for slide in reopened.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    )
    if len(reopened.slides) != deck_receipt["slides"]:
        raise ValueError("saved deck slide count changed after reopen")
    if embedded_images != len(stocks):
        raise ValueError(
            f"expected {len(stocks)} embedded chart images, got {embedded_images}"
        )
    deck_text = "\n".join(
        shape.text
        for slide in reopened.slides
        for shape in slide.shapes
        if hasattr(shape, "text_frame")
    )
    opposite_basis = "未复权" if daily_qfq else "前复权"
    if opposite_basis in deck_text:
        raise ValueError(f"deck contains contradictory price basis text: {opposite_basis}")
    if f"排名与展示均使用{basis_label}" not in deck_text:
        raise ValueError("deck is missing the ranking/chart price-basis declaration")

    output_metadata = {
        **selection_metadata,
        "generated_at": datetime.now().astimezone().isoformat(),
        "daily_access": (
            "coreClient.data_provider.get_day(source='database_only', "
            f"qfq={daily_qfq})"
        ),
        "price_basis_zh": basis_label,
        "ranking_and_chart_basis_match": True,
        "daily_input_file": daily_input_name,
        "daily_start": args.daily_start,
        "daily_end": data_end,
        "daily_rows": int(len(daily)),
        "daily_codes": int(daily["ts_code"].nunique()),
        "categories": {
            str(key): int(value)
            for key, value in stocks["broad_category"].value_counts().to_dict().items()
        },
        "charts": len(receipts),
        "deck": deck_receipt,
        "embedded_images": embedded_images,
        "pptx_file": deck_path.name,
        "pptx_sha256": _sha256(deck_path),
        "pptx_bytes": deck_path.stat().st_size,
    }
    (args.output_dir / "chart_receipts.json").write_text(
        json.dumps(receipts, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    (args.output_dir / "metadata.json").write_text(
        json.dumps(output_metadata, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    _write_report(stocks, args.output_dir / "report.md", output_metadata)
    print(json.dumps(output_metadata, ensure_ascii=False))


if __name__ == "__main__":
    main()

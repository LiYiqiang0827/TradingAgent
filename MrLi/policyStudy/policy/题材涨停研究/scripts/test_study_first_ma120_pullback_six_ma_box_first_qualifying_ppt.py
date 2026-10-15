from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
import pytest
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))

from study_first_ma120_pullback_six_ma_box_first_qualifying_ppt import (
    _add_bullets,
    _rgb,
    attach_stock_first_touch_comparison,
    quality_tier,
    summarize_joint_events,
)


def test_quality_tier_prioritizes_box_and_ma_position() -> None:
    assert quality_tier(True, True) == "A_窄箱体且均线位置保持"
    assert quality_tier(True, False) == "B_窄箱体但均线位置转弱"
    assert quality_tier(False, True) == "C_均线位置保持但箱体不足"
    assert quality_tier(False, False) == "D_箱体与位置均不足"


def test_rgb_helper_builds_a_powerpoint_color() -> None:
    assert tuple(_rgb("12ABEF")) == (0x12, 0xAB, 0xEF)


def test_add_bullets_does_not_depend_on_an_unexported_base_font() -> None:
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])

    _add_bullets(slide, ["测试项目"])

    assert "测试项目" in slide.shapes[-1].text


def test_joint_summary_reports_skipped_first_touch_and_core_metrics() -> None:
    events = pd.DataFrame(
        [
            {
                "skipped_earlier_first_touch": True,
                "delay_from_first_touch_trading_days": 10,
                "ma120_rising_streak_at_touch": 20,
                "narrow_box_pass": True,
                "box_position_descriptive_pass": True,
                "box_trading_days": 12,
                "minimum_six_ma_spread_pct": 3.0,
            },
            {
                "skipped_earlier_first_touch": False,
                "delay_from_first_touch_trading_days": 0,
                "ma120_rising_streak_at_touch": 40,
                "narrow_box_pass": False,
                "box_position_descriptive_pass": True,
                "box_trading_days": 3,
                "minimum_six_ma_spread_pct": 5.0,
            },
        ]
    )

    summary = summarize_joint_events(events).iloc[0]

    assert summary["unique_event_stocks"] == 2
    assert summary["skipped_earlier_first_touch_stocks"] == 1
    assert summary["narrow_box_stocks"] == 1
    assert summary["best_structure_stocks"] == 1
    assert summary["delay_median"] == pytest.approx(5.0)
    assert summary["ma120_streak_median"] == pytest.approx(30.0)
    assert summary["minimum_six_ma_spread_median"] == pytest.approx(4.0)


def test_stock_comparison_uses_earliest_raw_touch_across_anchors() -> None:
    events = pd.DataFrame(
        {
            "ts_code": ["000001.SZ"],
            "touch_date": ["20250108"],
            "touch_position": [5],
            "true_first_touch_date": ["20250106"],
            "true_first_touch_position": [3],
        }
    )
    raw_first = pd.DataFrame(
        {
            "ts_code": ["000001.SZ", "000001.SZ"],
            "touch_date": ["20250102", "20250106"],
            "touch_position": [1, 3],
            "anchor_period": [60, 120],
            "anchor_start_date": ["20241220", "20250101"],
        }
    )

    compared = attach_stock_first_touch_comparison(events, raw_first)

    row = compared.iloc[0]
    assert row["true_first_touch_date"] == "20250102"
    assert row["true_first_touch_position"] == 1
    assert row["delay_from_first_touch_trading_days"] == 4
    assert bool(row["skipped_earlier_first_touch"])
    assert row["anchor_first_touch_date"] == "20250106"

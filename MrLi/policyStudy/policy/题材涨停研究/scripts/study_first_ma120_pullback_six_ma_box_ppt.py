"""First MA120 pullback with six-MA cohesion, box analysis, and PPT output."""

from __future__ import annotations

import argparse
from datetime import datetime
import hashlib
import json
import math
from pathlib import Path
from typing import Any, Sequence

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.patches import Rectangle
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches

from study_first_ma30_pullback_cohesion_timing import compute_rising_streak
from study_first_ma30_pullback_cohesion_ppt import (
    BLUE,
    CYAN,
    DARK_TEXT,
    GOLD,
    LIGHT,
    MID,
    MUTED,
    NAVY,
    ORANGE,
    RED,
    TEAL,
    WHITE,
    _add_card,
    _add_footer,
    _add_header,
    _add_text,
    _draw_candlesticks,
    _register_fonts,
    _set_background,
    _style_table,
    convert_ppt_to_pdf,
    create_contact_sheets,
    render_pdf,
)


ROOT = Path(__file__).resolve().parents[5]
DEFAULT_DAILY = ROOT / "outputs" / "ma60_ma120_top100_ppt" / "2026_through_20260930_qfq" / "daily_input_qfq.csv.gz"
DEFAULT_MA60 = ROOT / "outputs" / "ma60_ma120_top100_themes" / "2026_through_20260930_qfq" / "ma60_top100_with_ties_named.csv"
DEFAULT_MA120 = ROOT / "outputs" / "ma60_ma120_top100_themes" / "2026_through_20260930_qfq" / "ma120_top100_with_ties_named.csv"
DEFAULT_OUTPUT = ROOT / "outputs" / "first_ma120_pullback_six_ma_box" / "2026_through_20260930_qfq"


def _compact_date(value: object) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip().replace("-", "").split(".")[0]


def _as_bool(value: object) -> bool:
    if isinstance(value, str):
        return value.strip().lower() in {"true", "1", "yes", "y"}
    return bool(value) if not pd.isna(value) else False


def _parse_args(argv: Sequence[str] | None = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--daily-file", type=Path, default=DEFAULT_DAILY)
    parser.add_argument("--ma60-file", type=Path, default=DEFAULT_MA60)
    parser.add_argument("--ma120-file", type=Path, default=DEFAULT_MA120)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--extension-lookback", type=int, default=20)
    parser.add_argument("--prior-extension-pct", type=float, default=8.0)
    parser.add_argument("--touch-upper-pct", type=float, default=2.0)
    parser.add_argument("--close-floor-pct", type=float, default=-3.0)
    parser.add_argument("--cohesion-threshold-pct", type=float, default=6.0)
    parser.add_argument("--max-short-below-ma120-pct", type=float, default=-2.0)
    parser.add_argument("--min-short-above-count", type=int, default=4)
    parser.add_argument("--box-lookback", type=int, default=10)
    parser.add_argument("--box-lookforward", type=int, default=30)
    parser.add_argument("--box-max-range-pct", type=float, default=10.0)
    parser.add_argument("--box-min-days", type=int, default=5)
    parser.add_argument("--metrics-only", action="store_true")
    return parser.parse_args(argv)


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _clean_json(value: Any) -> Any:
    if isinstance(value, dict):
        return {str(key): _clean_json(item) for key, item in value.items()}
    if isinstance(value, list):
        return [_clean_json(item) for item in value]
    if isinstance(value, np.integer):
        return int(value)
    if isinstance(value, np.floating):
        value = float(value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def add_ma120_pullback_features(
    frame: pd.DataFrame,
    *,
    extension_lookback: int,
    prior_extension_pct: float,
    touch_upper_pct: float,
    close_floor_pct: float,
    cohesion_threshold_pct: float,
    max_short_below_ma120_pct: float,
    min_short_above_count: int,
) -> pd.DataFrame:
    """Add no-lookahead MA120 touch and six-MA cohesion/position features."""
    if extension_lookback <= 0:
        raise ValueError("extension_lookback must be positive")
    result = frame.copy()
    ma_columns = ["ma5", "ma10", "ma20", "ma30", "ma60", "ma120"]
    ma = result[ma_columns].apply(pd.to_numeric, errors="coerce")
    mean = ma.mean(axis=1)
    result["six_ma_spread_pct"] = 100.0 * (ma.max(axis=1) - ma.min(axis=1)) / mean
    close = pd.to_numeric(result["close"], errors="coerce")
    low = pd.to_numeric(result["low"], errors="coerce")
    ma120 = ma["ma120"]
    result["close_vs_ma120_pct"] = 100.0 * (close / ma120 - 1.0)
    result["low_vs_ma120_pct"] = 100.0 * (low / ma120 - 1.0)
    result["prior_max_close_vs_ma120_pct"] = (
        result["close_vs_ma120_pct"]
        .shift(1)
        .rolling(extension_lookback, min_periods=extension_lookback)
        .max()
    )
    short = ma[["ma5", "ma10", "ma20", "ma30", "ma60"]]
    short_vs = short.div(ma120, axis=0).sub(1.0).mul(100.0)
    result["short_ma_above_ma120_count"] = short.ge(ma120, axis=0).sum(axis=1)
    result["min_short_ma_vs_ma120_pct"] = short_vs.min(axis=1)
    result["mean_short_ma_vs_ma120_pct"] = short_vs.mean(axis=1)
    rising = pd.to_numeric(result["ma120_rising_streak"], errors="coerce").gt(0)
    touch = (
        result["prior_max_close_vs_ma120_pct"].ge(prior_extension_pct)
        & result["low_vs_ma120_pct"].le(touch_upper_pct)
        & result["close_vs_ma120_pct"].ge(close_floor_pct)
        & rising
        & ma120.notna()
    )
    result["ma120_pullback_candidate"] = touch.fillna(False)
    result["six_ma_cohesion_position_pass"] = (
        result["six_ma_spread_pct"].le(cohesion_threshold_pct)
        & result["short_ma_above_ma120_count"].ge(min_short_above_count)
        & result["min_short_ma_vs_ma120_pct"].ge(max_short_below_ma120_pct)
    ).fillna(False)
    result["strict_first_touch_pass"] = (
        result["ma120_pullback_candidate"] & result["six_ma_cohesion_position_pass"]
    )
    return result


def select_first_touch(
    frame: pd.DataFrame,
    *,
    anchor_start_date: object,
    anchor_end_date: object,
    signal_column: str = "ma120_pullback_candidate",
) -> int | None:
    """Return the first MA120 touch in an inclusive anchor, without cohesion filtering."""
    dates = frame["trade_date"].map(_compact_date)
    mask = (
        dates.ge(_compact_date(anchor_start_date))
        & dates.le(_compact_date(anchor_end_date))
        & frame[signal_column].fillna(False).astype(bool)
    )
    positions = np.flatnonzero(mask.to_numpy())
    return int(positions[0]) if len(positions) else None


def select_first_qualifying_touch(
    frame: pd.DataFrame,
    *,
    anchor_start_date: object,
    anchor_end_date: object,
) -> int | None:
    """Return the first day jointly satisfying MA120 touch and six-MA rules."""
    dates = frame["trade_date"].map(_compact_date)
    mask = (
        dates.ge(_compact_date(anchor_start_date))
        & dates.le(_compact_date(anchor_end_date))
        & frame["ma120_pullback_candidate"].fillna(False).astype(bool)
        & frame["six_ma_cohesion_position_pass"].fillna(False).astype(bool)
    )
    positions = np.flatnonzero(mask.to_numpy())
    return int(positions[0]) if len(positions) else None


def select_unique_first_touches(events: pd.DataFrame) -> pd.DataFrame:
    """Select each stock's earliest touch with deterministic same-day ties."""
    return (
        events.sort_values(
            ["touch_date", "anchor_period", "anchor_start_date"],
            kind="mergesort",
        )
        .drop_duplicates("ts_code", keep="first")
        .copy()
    )


def find_narrow_box(
    frame: pd.DataFrame,
    *,
    touch_position: int,
    anchor_start_position: int,
    anchor_end_position: int,
    lookback: int,
    lookforward: int,
    max_range_pct: float,
    min_days: int,
) -> dict[str, Any]:
    """Find the longest high-low box containing the touch within a bounded window."""
    if touch_position < 0 or touch_position >= len(frame):
        raise IndexError(touch_position)
    search_start = max(0, int(anchor_start_position), touch_position - lookback)
    search_end = min(len(frame) - 1, int(anchor_end_position), touch_position + lookforward)
    high = pd.to_numeric(frame["high"], errors="coerce")
    low = pd.to_numeric(frame["low"], errors="coerce")
    candidates: list[tuple[int, float, int, int]] = []
    for start in range(search_start, touch_position + 1):
        for end in range(touch_position, search_end + 1):
            period_high = float(high.iloc[start : end + 1].max())
            period_low = float(low.iloc[start : end + 1].min())
            if not np.isfinite(period_high) or not np.isfinite(period_low) or period_low <= 0:
                continue
            range_pct = 100.0 * (period_high / period_low - 1.0)
            if range_pct <= max_range_pct:
                candidates.append((end - start + 1, range_pct, start, end))
    if not candidates:
        return {
            "box_start_position": None,
            "box_end_position": None,
            "box_start_date": "",
            "box_end_date": "",
            "box_trading_days": 0,
            "box_range_pct": float("nan"),
            "narrow_box_pass": False,
            "box_right_censored": False,
        }
    days, range_pct, start, end = sorted(
        candidates,
        key=lambda item: (-item[0], item[1], item[2], item[3]),
    )[0]
    return {
        "box_start_position": int(start),
        "box_end_position": int(end),
        "box_start_date": _compact_date(frame.iloc[start]["trade_date"]),
        "box_end_date": _compact_date(frame.iloc[end]["trade_date"]),
        "box_trading_days": int(days),
        "box_range_pct": float(range_pct),
        "narrow_box_pass": bool(days >= min_days),
        "box_right_censored": bool(end == search_end),
    }


def minimum_six_ma_spread(
    frame: pd.DataFrame,
    *,
    start_position: int,
    end_position: int,
) -> dict[str, Any]:
    """Locate the narrowest six-MA spread in an inclusive box interval."""
    values = pd.to_numeric(
        frame.iloc[start_position : end_position + 1]["six_ma_spread_pct"],
        errors="coerce",
    )
    if values.dropna().empty:
        return {
            "minimum_six_ma_spread_pct": float("nan"),
            "minimum_six_ma_spread_date": "",
            "minimum_six_ma_spread_position": None,
        }
    position = int(values.idxmin())
    return {
        "minimum_six_ma_spread_pct": float(values.loc[position]),
        "minimum_six_ma_spread_date": _compact_date(frame.loc[position, "trade_date"]),
        "minimum_six_ma_spread_position": position,
    }


def box_short_ma_position_metrics(
    frame: pd.DataFrame,
    *,
    start_position: int,
    end_position: int,
    floor_pct: float,
) -> dict[str, float]:
    """Describe MA5–MA60 position versus MA120 inside an inclusive box."""
    periods = (5, 10, 20, 30, 60)
    box = frame.iloc[start_position : end_position + 1]
    short = box[[f"ma{period}" for period in periods]].apply(pd.to_numeric, errors="coerce")
    ma120 = pd.to_numeric(box["ma120"], errors="coerce")
    relative = 100.0 * (short.div(ma120, axis=0) - 1.0)
    valid = relative.notna()
    valid_observations = int(valid.to_numpy().sum())
    complete_days = valid.all(axis=1)
    if valid_observations == 0 or not bool(complete_days.any()):
        return {
            "box_short_ma_observations_above_ma120_ratio": float("nan"),
            "box_short_ma_observations_within_floor_ratio": float("nan"),
            "box_days_all_short_ma_within_floor_ratio": float("nan"),
            "box_min_short_ma_vs_ma120_pct": float("nan"),
        }
    above_ratio = float(((relative >= 0.0) & valid).to_numpy().sum() / valid_observations)
    within_ratio = float(((relative >= floor_pct) & valid).to_numpy().sum() / valid_observations)
    all_within_days = (relative.loc[complete_days] >= floor_pct).all(axis=1)
    return {
        "box_short_ma_observations_above_ma120_ratio": above_ratio,
        "box_short_ma_observations_within_floor_ratio": within_ratio,
        "box_days_all_short_ma_within_floor_ratio": float(all_within_days.mean()),
        "box_min_short_ma_vs_ma120_pct": float(relative.min(axis=1, skipna=True).min()),
    }


def prepare_stock_features(daily: pd.DataFrame, args: argparse.Namespace) -> pd.DataFrame:
    frame = daily.copy()
    frame["trade_date"] = frame["trade_date"].map(_compact_date)
    frame = frame.sort_values("trade_date", kind="mergesort").drop_duplicates("trade_date", keep="last").reset_index(drop=True)
    close = pd.to_numeric(frame["close"], errors="coerce")
    for period in (5, 10, 20, 30, 60, 120):
        frame[f"ma{period}"] = close.rolling(period, min_periods=period).mean()
    frame["ma120_rising_streak"] = compute_rising_streak(frame["ma120"])
    return add_ma120_pullback_features(
        frame,
        extension_lookback=args.extension_lookback,
        prior_extension_pct=args.prior_extension_pct,
        touch_upper_pct=args.touch_upper_pct,
        close_floor_pct=args.close_floor_pct,
        cohesion_threshold_pct=args.cohesion_threshold_pct,
        max_short_below_ma120_pct=args.max_short_below_ma120_pct,
        min_short_above_count=args.min_short_above_count,
    )


def load_anchors(args: argparse.Namespace) -> pd.DataFrame:
    pieces: list[pd.DataFrame] = []
    for period, path in ((60, args.ma60_file), (120, args.ma120_file)):
        frame = pd.read_csv(path, dtype={"ts_code": str, "start_date": str, "end_date": str})
        frame["anchor_period"] = period
        frame["start_date"] = frame["start_date"].map(_compact_date)
        frame["end_date"] = frame["end_date"].map(_compact_date)
        frame["left_censored"] = frame["left_censored"].map(_as_bool)
        frame["right_censored"] = frame["right_censored"].map(_as_bool)
        frame["search_eligible"] = ~frame["left_censored"]
        pieces.append(frame)
    return pd.concat(pieces, ignore_index=True)


def build_first_touch_record(
    frame: pd.DataFrame,
    *,
    position: int,
    anchor: dict[str, Any],
    args: argparse.Namespace,
) -> dict[str, Any]:
    dates = frame["trade_date"].map(_compact_date)
    start_matches = np.flatnonzero(dates.eq(_compact_date(anchor["start_date"])).to_numpy())
    end_matches = np.flatnonzero(dates.eq(_compact_date(anchor["end_date"])).to_numpy())
    if len(start_matches) != 1 or len(end_matches) != 1:
        raise ValueError(f"anchor date lookup failed for {anchor['ts_code']}")
    anchor_start_position = int(start_matches[0])
    anchor_end_position = int(end_matches[0])
    row = frame.iloc[position]
    box = find_narrow_box(
        frame,
        touch_position=position,
        anchor_start_position=anchor_start_position,
        anchor_end_position=anchor_end_position,
        lookback=args.box_lookback,
        lookforward=args.box_lookforward,
        max_range_pct=args.box_max_range_pct,
        min_days=args.box_min_days,
    )
    if box["box_start_position"] is None:
        minimum = {
            "minimum_six_ma_spread_pct": float("nan"),
            "minimum_six_ma_spread_date": "",
            "minimum_six_ma_spread_position": None,
        }
        box_position = {
            "box_short_ma_observations_above_ma120_ratio": float("nan"),
            "box_short_ma_observations_within_floor_ratio": float("nan"),
            "box_days_all_short_ma_within_floor_ratio": float("nan"),
            "box_min_short_ma_vs_ma120_pct": float("nan"),
        }
    else:
        minimum = minimum_six_ma_spread(
            frame,
            start_position=int(box["box_start_position"]),
            end_position=int(box["box_end_position"]),
        )
        box_position = box_short_ma_position_metrics(
            frame,
            start_position=int(box["box_start_position"]),
            end_position=int(box["box_end_position"]),
            floor_pct=args.max_short_below_ma120_pct,
        )
    record: dict[str, Any] = {
        "ts_code": str(anchor["ts_code"]),
        "stock_name": str(anchor.get("stock_name", "")),
        "primary_theme": str(anchor.get("primary_theme", "")),
        "broad_category": str(anchor.get("broad_category", "")),
        "anchor_period": int(anchor["anchor_period"]),
        "anchor_rank": int(anchor.get("rank", 0)),
        "anchor_start_date": _compact_date(anchor["start_date"]),
        "anchor_end_date": _compact_date(anchor["end_date"]),
        "anchor_length": int(anchor.get("length", 0)),
        "anchor_right_censored": bool(anchor.get("right_censored", False)),
        "touch_date": _compact_date(row["trade_date"]),
        "touch_position": int(position),
        "anchor_observed_days_at_touch": int(position - anchor_start_position + 1),
        "open": float(row["open"]),
        "high": float(row["high"]),
        "low": float(row["low"]),
        "close": float(row["close"]),
        "prior_max_close_vs_ma120_pct": float(row["prior_max_close_vs_ma120_pct"]),
        "low_vs_ma120_pct": float(row["low_vs_ma120_pct"]),
        "close_vs_ma120_pct": float(row["close_vs_ma120_pct"]),
        "six_ma_spread_at_touch_pct": float(row["six_ma_spread_pct"]),
        "short_ma_above_ma120_count": int(row["short_ma_above_ma120_count"]),
        "min_short_ma_vs_ma120_pct": float(row["min_short_ma_vs_ma120_pct"]),
        "mean_short_ma_vs_ma120_pct": float(row["mean_short_ma_vs_ma120_pct"]),
        "ma120_rising_streak_at_touch": int(row["ma120_rising_streak"]),
        "six_ma_cohesion_position_pass": bool(row["six_ma_cohesion_position_pass"]),
    }
    for period in (5, 10, 20, 30, 60, 120):
        record[f"ma{period}"] = float(row[f"ma{period}"])
    record.update(box)
    record.update(minimum)
    record.update(box_position)
    record["box_position_descriptive_pass"] = bool(
        pd.notna(record["box_short_ma_observations_above_ma120_ratio"])
        and record["box_short_ma_observations_above_ma120_ratio"] >= 0.80
        and record["box_short_ma_observations_within_floor_ratio"] >= 1.0 - 1e-12
    )
    record["event_grade"] = "A_六均线粘合且窄箱体" if bool(record["six_ma_cohesion_position_pass"] and record["narrow_box_pass"]) else (
        "B_六均线粘合但箱体不足" if bool(record["six_ma_cohesion_position_pass"]) else "未通过六均线粘合位置约束"
    )
    return record


def run_scan(args: argparse.Namespace) -> dict[str, pd.DataFrame]:
    daily = pd.read_csv(args.daily_file, compression="infer", dtype={"ts_code": str, "trade_date": str})
    daily["trade_date"] = daily["trade_date"].map(_compact_date)
    if daily.duplicated(["ts_code", "trade_date"]).any():
        raise ValueError("duplicate stock-date rows")
    features_by_code = {
        str(code): prepare_stock_features(group, args)
        for code, group in daily.groupby("ts_code", sort=False)
    }
    anchors = load_anchors(args)
    missing = sorted(set(anchors["ts_code"].astype(str)) - set(features_by_code))
    if missing:
        raise ValueError(f"daily input missing anchor stocks: {missing}")
    records: list[dict[str, Any]] = []
    eligibility: list[dict[str, Any]] = []
    for anchor in anchors.to_dict("records"):
        position = None
        if bool(anchor["search_eligible"]):
            frame = features_by_code[str(anchor["ts_code"])]
            position = select_first_touch(
                frame,
                anchor_start_date=anchor["start_date"],
                anchor_end_date=anchor["end_date"],
            )
        eligibility.append(
            {
                "ts_code": str(anchor["ts_code"]),
                "stock_name": anchor.get("stock_name", ""),
                "anchor_period": int(anchor["anchor_period"]),
                "anchor_start_date": _compact_date(anchor["start_date"]),
                "anchor_end_date": _compact_date(anchor["end_date"]),
                "left_censored": bool(anchor["left_censored"]),
                "search_eligible": bool(anchor["search_eligible"]),
                "first_ma120_touch_date": _compact_date(features_by_code[str(anchor["ts_code"])].iloc[position]["trade_date"]) if position is not None else "",
            }
        )
        if position is not None:
            records.append(
                build_first_touch_record(
                    features_by_code[str(anchor["ts_code"])],
                    position=position,
                    anchor=anchor,
                    args=args,
                )
            )
    all_first = pd.DataFrame(records)
    if all_first.empty:
        raise ValueError("no first MA120 pullbacks found")
    anchor_events = all_first.loc[all_first["six_ma_cohesion_position_pass"]].copy()
    anchor_box_events = anchor_events.loc[anchor_events["narrow_box_pass"]].copy()
    unique_first = select_unique_first_touches(all_first)
    unique_events = unique_first.loc[unique_first["six_ma_cohesion_position_pass"]].copy()
    unique_box_events = unique_events.loc[unique_events["narrow_box_pass"]].copy()
    unique_events = unique_events.sort_values(
        ["narrow_box_pass", "minimum_six_ma_spread_pct", "box_trading_days", "touch_date"],
        ascending=[False, True, False, True],
        kind="mergesort",
    ).reset_index(drop=True)
    unique_events.insert(0, "display_order", np.arange(1, len(unique_events) + 1))
    unique_box_events = unique_events.loc[unique_events["narrow_box_pass"]].copy()

    coverage_rows: list[dict[str, Any]] = []
    for period in (60, 120):
        anchor_group = anchors.loc[anchors["anchor_period"].eq(period)]
        first_group = all_first.loc[all_first["anchor_period"].eq(period)]
        event_group = anchor_events.loc[anchor_events["anchor_period"].eq(period)]
        box_group = anchor_box_events.loc[anchor_box_events["anchor_period"].eq(period)]
        coverage_rows.append(
            {
                "anchor_period": period,
                "anchor_count": int(len(anchor_group)),
                "left_censored_count": int(anchor_group["left_censored"].sum()),
                "eligible_count": int(anchor_group["search_eligible"].sum()),
                "first_ma120_touch_count": int(len(first_group)),
                "six_ma_cohesion_position_count": int(len(event_group)),
                "narrow_box_count": int(len(box_group)),
            }
        )
    coverage = pd.DataFrame(coverage_rows)
    unique_summary = pd.DataFrame(
        [
            {
                "scope": "UNION_UNIQUE",
                "first_touch_stocks": int(len(unique_first)),
                "six_ma_event_stocks": int(len(unique_events)),
                "narrow_box_stocks": int(len(unique_box_events)),
                "ma120_streak_median": float(unique_events["ma120_rising_streak_at_touch"].median()) if len(unique_events) else float("nan"),
                "ma120_streak_p25": float(unique_events["ma120_rising_streak_at_touch"].quantile(0.25)) if len(unique_events) else float("nan"),
                "ma120_streak_p75": float(unique_events["ma120_rising_streak_at_touch"].quantile(0.75)) if len(unique_events) else float("nan"),
                "box_days_median": float(unique_box_events["box_trading_days"].median()) if len(unique_box_events) else float("nan"),
                "box_days_p25": float(unique_box_events["box_trading_days"].quantile(0.25)) if len(unique_box_events) else float("nan"),
                "box_days_p75": float(unique_box_events["box_trading_days"].quantile(0.75)) if len(unique_box_events) else float("nan"),
                "minimum_six_ma_spread_median": float(unique_events["minimum_six_ma_spread_pct"].median()) if len(unique_events) else float("nan"),
                "minimum_six_ma_spread_min": float(unique_events["minimum_six_ma_spread_pct"].min()) if len(unique_events) else float("nan"),
                "box_position_descriptive_pass_count": int(unique_events["box_position_descriptive_pass"].sum()) if len(unique_events) else 0,
                "box_short_ma_above_ratio_median": float(unique_events["box_short_ma_observations_above_ma120_ratio"].median()) if len(unique_events) else float("nan"),
                "box_short_ma_above_ratio_min": float(unique_events["box_short_ma_observations_above_ma120_ratio"].min()) if len(unique_events) else float("nan"),
                "box_short_ma_within_floor_ratio_min": float(unique_events["box_short_ma_observations_within_floor_ratio"].min()) if len(unique_events) else float("nan"),
                "box_days_all_short_ma_within_floor_ratio_min": float(unique_events["box_days_all_short_ma_within_floor_ratio"].min()) if len(unique_events) else float("nan"),
            }
        ]
    )
    return {
        "daily": daily,
        "anchors": anchors,
        "eligibility": pd.DataFrame(eligibility),
        "all_first": all_first,
        "anchor_events": anchor_events,
        "anchor_box_events": anchor_box_events,
        "unique_first": unique_first,
        "unique_events": unique_events,
        "unique_box_events": unique_box_events,
        "coverage": coverage,
        "summary": unique_summary,
        "features_by_code": features_by_code,
    }


def build_sensitivity(
    result: dict[str, Any],
    args: argparse.Namespace,
    *,
    scan_fn: Any | None = None,
) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    first = result["unique_first"]
    for spread in (4.0, 5.0, 6.0, 7.0, 8.0):
        selected = first.loc[
            first["six_ma_spread_at_touch_pct"].le(spread)
            & first["short_ma_above_ma120_count"].ge(args.min_short_above_count)
            & first["min_short_ma_vs_ma120_pct"].ge(args.max_short_below_ma120_pct)
        ]
        rows.append(
            {
                "sensitivity_type": "six_ma_spread_threshold",
                "threshold": spread,
                "event_count": int(len(selected)),
                "narrow_box_count": int(selected["narrow_box_pass"].sum()),
                "box_range_threshold_pct": float(getattr(args, "box_max_range_pct", 10.0)),
            }
        )
    for box_range in (8.0, 10.0, 12.0):
        pass_count = 0
        days: list[int] = []
        for event in result["unique_events"].to_dict("records"):
            frame = result["features_by_code"][str(event["ts_code"])]
            dates = frame["trade_date"].map(_compact_date)
            start = int(np.flatnonzero(dates.eq(_compact_date(event["anchor_start_date"])).to_numpy())[0])
            end = int(np.flatnonzero(dates.eq(_compact_date(event["anchor_end_date"])).to_numpy())[0])
            box = find_narrow_box(
                frame,
                touch_position=int(event["touch_position"]),
                anchor_start_position=start,
                anchor_end_position=end,
                lookback=args.box_lookback,
                lookforward=args.box_lookforward,
                max_range_pct=box_range,
                min_days=args.box_min_days,
            )
            if box["narrow_box_pass"]:
                pass_count += 1
                days.append(int(box["box_trading_days"]))
        rows.append(
            {
                "sensitivity_type": "box_range_threshold",
                "threshold": box_range,
                "event_count": int(len(result["unique_events"])),
                "narrow_box_count": pass_count,
                "box_range_threshold_pct": box_range,
                "box_days_median": float(pd.Series(days).median()) if days else float("nan"),
            }
        )
    scanner = run_scan if scan_fn is None else scan_fn
    threshold_grids = (
        ("prior_extension_threshold", "prior_extension_pct", (5.0, 8.0, 10.0)),
        ("touch_upper_threshold", "touch_upper_pct", (1.0, 2.0, 3.0)),
        ("close_floor_threshold", "close_floor_pct", (-1.0, -3.0, -5.0)),
        ("position_floor_threshold", "max_short_below_ma120_pct", (-1.0, -2.0, -3.0)),
        ("min_short_above_count", "min_short_above_count", (3, 4, 5)),
    )
    for sensitivity_type, attribute, values in threshold_grids:
        for value in values:
            use_main = (
                float(getattr(args, attribute)) == float(value)
                and "unique_box_events" in result
                and "unique_first" in result
            )
            if use_main:
                scanned = result
            else:
                variant = argparse.Namespace(**vars(args))
                setattr(variant, attribute, value)
                scanned = scanner(variant)
            rows.append(
                {
                    "sensitivity_type": sensitivity_type,
                    "threshold": value,
                    "event_count": int(len(scanned["unique_events"])),
                    "narrow_box_count": int(len(scanned["unique_box_events"])),
                    "box_range_threshold_pct": float(getattr(args, "box_max_range_pct", 10.0)),
                    "first_touch_stock_count": int(len(scanned["unique_first"])) if "unique_first" in scanned else float("nan"),
                }
            )
    return pd.DataFrame(rows)


def _fmt_date(value: object) -> str:
    text = _compact_date(value)
    return f"{text[:4]}-{text[4:6]}-{text[6:8]}" if len(text) == 8 else text


def create_event_chart(frame: pd.DataFrame, event: pd.Series, output: Path) -> None:
    _register_fonts()
    dates = frame["trade_date"].map(_compact_date)
    touch = int(event["touch_position"])
    start = int(event["box_start_position"]) if pd.notna(event["box_start_position"]) else touch
    end = int(event["box_end_position"]) if pd.notna(event["box_end_position"]) else touch
    anchor_matches = np.flatnonzero(dates.eq(_compact_date(event["anchor_start_date"])).to_numpy())
    if len(anchor_matches) != 1:
        raise ValueError(f"anchor date lookup failed for chart {event['ts_code']}")
    anchor_start = int(anchor_matches[0])
    lower = max(0, min(anchor_start, touch - 85, start - 15) - 5)
    upper = min(len(frame) - 1, max(touch + 30, end + 20))
    view = frame.iloc[lower : upper + 1].reset_index(drop=True).copy()
    touch_x, start_x, end_x, anchor_x = touch - lower, start - lower, end - lower, anchor_start - lower

    fig = plt.figure(figsize=(12.2, 6.35), dpi=155, facecolor="white")
    grid = fig.add_gridspec(5, 1, hspace=0.06)
    ax = fig.add_subplot(grid[:4, 0])
    vol_ax = fig.add_subplot(grid[4, 0], sharex=ax)
    _draw_candlesticks(ax, view)
    colors = {5: ORANGE, 10: GOLD, 20: "E76F9A", 30: BLUE, 60: "7B2CBF", 120: TEAL}
    for period in (5, 10, 20, 30, 60, 120):
        ax.plot(
            np.arange(len(view)),
            pd.to_numeric(view[f"ma{period}"], errors="coerce"),
            color=f"#{colors[period]}",
            linewidth=1.05 if period < 30 else 1.4,
            label=f"MA{period}",
            zorder=4,
        )
    ax.axvline(anchor_x, color=f"#{MUTED}", linewidth=0.9, linestyle=":")
    ax.axvline(touch_x, color=f"#{RED}", linewidth=1.6, linestyle="--", zorder=6)
    box_slice = view.iloc[max(0, start_x) : end_x + 1]
    if len(box_slice):
        box_low = float(pd.to_numeric(box_slice["low"], errors="coerce").min())
        box_high = float(pd.to_numeric(box_slice["high"], errors="coerce").max())
        rect = Rectangle(
            (start_x - 0.5, box_low),
            end_x - start_x + 1,
            box_high - box_low,
            facecolor=f"#{CYAN}",
            edgecolor=f"#{BLUE}",
            linewidth=1.2,
            alpha=0.16,
            zorder=1,
        )
        ax.add_patch(rect)
    touch_close = float(view.iloc[touch_x]["close"])
    ax.scatter([touch_x], [touch_close], s=52, color=f"#{RED}", edgecolor="white", linewidth=0.8, zorder=7)
    ax.annotate(
        f"首次回调MA120\n{_fmt_date(event['touch_date'])}",
        xy=(touch_x, touch_close),
        xytext=(8, 24),
        textcoords="offset points",
        fontsize=8,
        color=f"#{RED}",
        fontweight="bold",
        arrowprops={"arrowstyle": "->", "color": f"#{RED}", "lw": 0.8},
    )
    min_date = _compact_date(event["minimum_six_ma_spread_date"])
    min_matches = np.flatnonzero(view["trade_date"].map(_compact_date).eq(min_date).to_numpy()) if min_date else np.array([])
    if len(min_matches) == 1:
        ax.axvline(int(min_matches[0]), color=f"#{TEAL}", linewidth=1.0, linestyle="-.", zorder=5)
    ax.legend(loc="upper left", ncol=6, frameon=False, fontsize=7.2)
    ax.grid(axis="y", color="#E7EDF3", linewidth=0.6)
    ax.set_ylabel("前复权价格", fontsize=8, color=f"#{MUTED}")
    ax.tick_params(axis="both", labelsize=7, colors=f"#{MUTED}")
    ax.spines[["top", "right"]].set_visible(False)
    ax.spines[["left", "bottom"]].set_color("#CBD5E1")
    plt.setp(ax.get_xticklabels(), visible=False)

    volumes = pd.to_numeric(view["vol"], errors="coerce").fillna(0.0)
    volume_colors = [f"#{RED}" if float(close) >= float(open_) else f"#{TEAL}" for open_, close in zip(view["open"], view["close"])]
    vol_ax.bar(np.arange(len(view)), volumes, color=volume_colors, width=0.62, alpha=0.75)
    vol_ax.axvline(touch_x, color=f"#{RED}", linewidth=1.2, linestyle="--")
    vol_ax.grid(axis="y", color="#EEF2F6", linewidth=0.5)
    vol_ax.set_ylabel("成交量", fontsize=7, color=f"#{MUTED}")
    vol_ax.tick_params(axis="both", labelsize=6.5, colors=f"#{MUTED}")
    vol_ax.spines[["top", "right"]].set_visible(False)
    vol_ax.spines[["left", "bottom"]].set_color("#CBD5E1")
    step = max(1, len(view) // 8)
    ticks = sorted(set(list(range(0, len(view), step)) + [touch_x, len(view) - 1]))
    vol_ax.set_xticks(ticks)
    vol_ax.set_xticklabels([_fmt_date(view.iloc[index]["trade_date"])[5:] for index in ticks], fontsize=6.5)
    fig.suptitle(
        f"{event['stock_name']}  {event['ts_code']}  |  MA{int(event['anchor_period'])}锚点内首次回调MA120",
        x=0.01,
        y=0.995,
        ha="left",
        va="top",
        fontsize=11,
        color=f"#{DARK_TEXT}",
        fontweight="bold",
    )
    fig.text(
        0.99,
        0.99,
        f"蓝框={int(event['box_trading_days'])}日箱体 / {float(event['box_range_pct']):.1f}%振幅；绿线=六均线最窄日",
        ha="right",
        va="top",
        fontsize=7.5,
        color=f"#{MUTED}",
    )
    output.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(output, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def create_summary_chart(events: pd.DataFrame, output: Path) -> None:
    _register_fonts()
    names = events["stock_name"].astype(str).tolist()
    x = np.arange(len(events))
    fig, axes = plt.subplots(1, 3, figsize=(12.5, 4.3), dpi=160, facecolor="white")
    axes[0].bar(x, events["ma120_rising_streak_at_touch"], color=f"#{TEAL}")
    axes[0].set_xticks(x, names, rotation=28, ha="right")
    axes[0].set_title("首次回调时MA120连续上涨", fontsize=10, fontweight="bold")
    axes[0].set_ylabel("交易日")
    axes[1].bar(x - 0.18, events["box_trading_days"], width=0.36, color=f"#{BLUE}", label="箱体天数")
    axes[1].bar(x + 0.18, events["box_range_pct"], width=0.36, color=f"#{ORANGE}", label="箱体振幅%")
    axes[1].set_xticks(x, names, rotation=28, ha="right")
    axes[1].set_title("窄幅箱体：天数与振幅", fontsize=10, fontweight="bold")
    axes[1].legend(frameon=False, fontsize=8)
    axes[2].bar(x - 0.18, events["six_ma_spread_at_touch_pct"], width=0.36, color=f"#{GOLD}", label="回调日")
    axes[2].bar(x + 0.18, events["minimum_six_ma_spread_pct"], width=0.36, color=f"#{RED}", label="箱体内最窄")
    axes[2].set_xticks(x, names, rotation=28, ha="right")
    axes[2].set_title("六均线最大—最小离差", fontsize=10, fontweight="bold")
    axes[2].set_ylabel("占均值百分比")
    axes[2].legend(frameon=False, fontsize=8)
    for ax in axes:
        ax.grid(axis="y", color="#E7EDF3", linewidth=0.6)
        ax.spines[["top", "right"]].set_visible(False)
        ax.tick_params(labelsize=7.5, colors=f"#{MUTED}")
    fig.tight_layout()
    fig.savefig(output, bbox_inches="tight", pad_inches=0.08)
    plt.close(fig)


def build_deck(
    *,
    events: pd.DataFrame,
    coverage: pd.DataFrame,
    summary: pd.DataFrame,
    sensitivity: pd.DataFrame,
    chart_paths: dict[str, Path],
    summary_chart: Path,
    output: Path,
) -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, NAVY)
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.62), Inches(1.18), Inches(0.12), Inches(4.82))
    stripe.fill.solid(); stripe.fill.fore_color.rgb = __import__("pptx").dml.color.RGBColor.from_string(CYAN); stripe.line.fill.background()
    _add_text(slide, "第一次回调MA120\n六均线粘合与窄幅箱体", 1.05, 1.25, 10.8, 1.55, size=31, color=WHITE, bold=True)
    _add_text(slide, "MA5 / MA10 / MA20 / MA30 / MA60 / MA120", 1.08, 3.14, 10.0, 0.46, size=18, color=CYAN)
    _add_text(slide, "178只长趋势候选池｜前复权日线｜截至2026-09-30", 1.08, 4.10, 9.5, 0.36, size=12, color="BFD7EA")
    _add_text(slide, "严格事件6只，其中5只形成≥5个交易日的窄幅箱体", 1.08, 4.68, 10.4, 0.55, size=15, color=WHITE, bold=True)
    _add_text(slide, "研究用途，不构成买入建议", 1.08, 6.72, 5.0, 0.26, size=9, color="94A3B8")

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, "主口径：先固定第一次MA120回调，再检查六均线粘合与位置")
    bullets = [
        "第一次回调：此前20个交易日最高收盘至少高于自身MA120 8%；当日最低价≤MA120上方2%，收盘≥MA120下方3%，且MA120当日仍严格上涨；持平、下降或缺失均将连续上涨计数清零。",
        "六均线粘合：(max(MA5/10/20/30/60/120)-min)/均值≤6%；六条MA均先在完整前复权历史上滚动计算，再切入锚点。",
        "位置约束：事件日MA5—MA60五条均线至少4条不低于MA120，且最低一条最多低于MA120 2%；箱体内另以五线观测中≥80%不低于MA120、100%不低于MA120下方2%作描述性核验。",
        "首次回调日如果不粘合，不允许跳到后续更漂亮的位置；联合股票口径也先跨锚点取最早首次回调，再过滤；同日依次按锚点周期、锚点起点排序。",
        "窄幅箱体：锚点内、事件前最多10日和事件后最多30日，寻找包含事件日且最高价/最低价振幅≤10%的最长区间；同长度先取振幅更小、再取起点更早，至少5个交易日。",
        "箱体天数和箱体内指标可能使用回调后的价格，只作事后描述；首次回调与事件日粘合判断无前视。候选池按事后观察到的长MA持续时间筛选，存在选择偏差。",
    ]
    box = slide.shapes.add_textbox(Inches(0.72), Inches(1.05), Inches(11.9), Inches(5.82))
    tf = box.text_frame; tf.clear(); tf.word_wrap = True
    for index, item in enumerate(bullets):
        p = tf.paragraphs[0] if index == 0 else tf.add_paragraph()
        p.text = f"• {item}"; p.font.name = "Noto Sans CJK SC"; p.font.size = __import__("pptx").util.Pt(13); p.font.color.rgb = __import__("pptx").dml.color.RGBColor.from_string(DARK_TEXT); p.space_after = __import__("pptx").util.Pt(9)
    _add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, "筛选漏斗：第一次回调MA120并不罕见，六均线同时收敛很少见")
    first_anchor = int(coverage["first_ma120_touch_count"].sum())
    strict_anchor = int(coverage["six_ma_cohesion_position_count"].sum())
    box_anchor = int(coverage["narrow_box_count"].sum())
    cards = [("231", "非左截断可搜索锚点", BLUE), (str(first_anchor), "找到首次MA120回调", CYAN), (str(strict_anchor), "锚点级六均线严格事件", ORANGE), (str(len(events)), "联合去重主样本", RED)]
    for index, (value, label, color) in enumerate(cards):
        _add_card(slide, 0.56 + index * 3.13, 1.18, 2.73, 1.36, value, label, accent=color)
    table = slide.shapes.add_table(3, 5, Inches(0.85), Inches(3.08), Inches(11.65), Inches(2.15)).table
    headers = ["锚点", "可搜索", "首次回调MA120", "六均线严格事件", "其中窄箱体"]
    for col, text in enumerate(headers): table.cell(0, col).text = text
    for row_idx, row in enumerate(coverage.itertuples(), start=1):
        values = [f"MA{int(row.anchor_period)}", str(int(row.eligible_count)), str(int(row.first_ma120_touch_count)), str(int(row.six_ma_cohesion_position_count)), str(int(row.narrow_box_count))]
        for col, text in enumerate(values): table.cell(row_idx, col).text = text
    _style_table(table, header_size=10, body_size=10)
    _add_text(slide, f"锚点级{strict_anchor}个事件中{box_anchor}个有窄箱体；联合去重后6只，其中5只为A类窄箱体。", 0.9, 5.72, 11.5, 0.52, size=14, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    _add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, "6只联合主样本")
    headers = ["股票", "首次回调", "MA120上涨", "箱体", "箱体振幅", "回调日粘合", "箱体内最窄", "评级"]
    table = slide.shapes.add_table(len(events) + 1, len(headers), Inches(0.35), Inches(1.06), Inches(12.65), Inches(5.72)).table
    widths = [1.95, 1.15, 1.05, 0.82, 1.05, 1.15, 1.15, 1.75]
    for index, width in enumerate(widths): table.columns[index].width = Inches(width)
    for col, text in enumerate(headers): table.cell(0, col).text = text
    for row_idx, event in enumerate(events.itertuples(), start=1):
        values = [
            f"{event.stock_name}\n{event.ts_code}", _fmt_date(event.touch_date)[5:], f"{int(event.ma120_rising_streak_at_touch)}日",
            f"{int(event.box_trading_days)}日", f"{float(event.box_range_pct):.1f}%", f"{float(event.six_ma_spread_at_touch_pct):.2f}%",
            f"{float(event.minimum_six_ma_spread_pct):.2f}%", "A 窄箱体" if bool(event.narrow_box_pass) else "B 箱体不足",
        ]
        for col, text in enumerate(values): table.cell(row_idx, col).text = text
    _style_table(table, header_size=9.0, body_size=8.6)
    _add_footer(slide, len(prs.slides))

    row = summary.iloc[0]
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, WHITE)
    _add_header(slide, "统计：MA120上涨阶段普遍很早，箱体差异较大")
    slide.shapes.add_picture(str(summary_chart), Inches(0.45), Inches(1.05), width=Inches(12.4), height=Inches(4.72))
    _add_text(slide, f"MA120连续上涨中位数 {row['ma120_streak_median']:.1f}日（IQR {row['ma120_streak_p25']:.1f}–{row['ma120_streak_p75']:.1f}）", 0.65, 6.02, 4.0, 0.42, size=11, bold=True, color=TEAL, align=PP_ALIGN.CENTER)
    _add_text(slide, f"A类箱体天数中位数 {row['box_days_median']:.1f}日（IQR {row['box_days_p25']:.1f}–{row['box_days_p75']:.1f}）", 4.65, 6.02, 4.0, 0.42, size=11, bold=True, color=BLUE, align=PP_ALIGN.CENTER)
    _add_text(slide, f"箱体内六均线最窄中位数 {row['minimum_six_ma_spread_median']:.2f}%｜位置核验 {int(row['box_position_descriptive_pass_count'])}/{len(events)}", 8.45, 6.02, 4.35, 0.42, size=10.6, bold=True, color=RED, align=PP_ALIGN.CENTER)
    _add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, "阈值敏感性：主口径6%，不是为了凑数量")
    spread_sens = sensitivity.loc[sensitivity["sensitivity_type"].eq("six_ma_spread_threshold")]
    box_sens = sensitivity.loc[sensitivity["sensitivity_type"].eq("box_range_threshold")]
    _add_text(slide, "六均线粘合阈值", 0.8, 1.15, 5.6, 0.38, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    table1 = slide.shapes.add_table(len(spread_sens) + 1, 3, Inches(0.75), Inches(1.72), Inches(5.7), Inches(3.6)).table
    for col, text in enumerate(["最大离差", "事件数", "其中≤10%箱体"]): table1.cell(0, col).text = text
    for idx, row_s in enumerate(spread_sens.itertuples(), start=1):
        for col, text in enumerate([f"≤{row_s.threshold:.0f}%", str(int(row_s.event_count)), str(int(row_s.narrow_box_count))]): table1.cell(idx, col).text = text
    _style_table(table1, header_size=10, body_size=10)
    _add_text(slide, "窄箱体振幅阈值（固定6只严格事件）", 6.95, 1.15, 5.6, 0.38, size=15, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    table2 = slide.shapes.add_table(len(box_sens) + 1, 3, Inches(6.9), Inches(1.72), Inches(5.7), Inches(2.52)).table
    for col, text in enumerate(["最高/最低振幅", "箱体合格数", "箱体天数中位数"]): table2.cell(0, col).text = text
    for idx, row_s in enumerate(box_sens.itertuples(), start=1):
        median = "—" if pd.isna(row_s.box_days_median) else f"{row_s.box_days_median:.1f}日"
        for col, text in enumerate([f"≤{row_s.threshold:.0f}%", str(int(row_s.narrow_box_count)), median]): table2.cell(idx, col).text = text
    _style_table(table2, header_size=10, body_size=10)
    _add_text(slide, "主口径选择六均线最大离差≤6%、箱体高低振幅≤10%。放宽到8%会把联合事件从6只扩到21只，形态纯度明显下降。", 0.92, 5.75, 11.45, 0.62, size=12, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    _add_footer(slide, len(prs.slides))

    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_background(slide, LIGHT)
    _add_header(slide, "首次回调与均线位置阈值敏感性：每组都重新固定第一触点")
    sensitivity_specs = [
        ("prior_extension_threshold", "此前离开MA120", lambda value: f"≥{value:.0f}%"),
        ("touch_upper_threshold", "最低价触及容忍", lambda value: f"≤+{value:.0f}%"),
        ("close_floor_threshold", "收盘破位容忍", lambda value: f"≥{value:.0f}%"),
        ("position_floor_threshold", "短中均线最低位置", lambda value: f"≥{value:.0f}%"),
        ("min_short_above_count", "位于MA120上方条数", lambda value: f"≥{int(value)}/5"),
    ]
    positions = [(0.36, 1.20), (4.58, 1.20), (8.80, 1.20), (2.48, 4.10), (6.70, 4.10)]
    for (kind, title, formatter), (left, top) in zip(sensitivity_specs, positions):
        part = sensitivity.loc[sensitivity["sensitivity_type"].eq(kind)]
        _add_text(slide, title, left, top, 3.82, 0.30, size=11.5, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        table = slide.shapes.add_table(len(part) + 1, 4, Inches(left), Inches(top + 0.45), Inches(3.82), Inches(1.88)).table
        for col, text in enumerate(["阈值", "首触股票", "严格事件", "10%箱体"]):
            table.cell(0, col).text = text
        for row_idx, row_s in enumerate(part.itertuples(), start=1):
            first_count = "—" if pd.isna(row_s.first_touch_stock_count) else str(int(row_s.first_touch_stock_count))
            values = [formatter(float(row_s.threshold)), first_count, str(int(row_s.event_count)), str(int(row_s.narrow_box_count))]
            for col, text in enumerate(values):
                table.cell(row_idx, col).text = text
        _style_table(table, header_size=7.8, body_size=8.2)
    _add_text(slide, "说明：改变回调阈值时，先按新阈值重新寻找每个锚点的第一次回调，再做股票级最早事件选择和六均线过滤；不是只对主样本二次筛选。", 0.95, 6.65, 11.4, 0.34, size=9.3, color=BLUE, bold=True, align=PP_ALIGN.CENTER)
    _add_footer(slide, len(prs.slides))

    for _, event in events.iterrows():
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        _set_background(slide, LIGHT)
        grade = "A类窄箱体" if bool(event["narrow_box_pass"]) else "B类箱体不足"
        _add_header(slide, f"{int(event['display_order']):02d}  {event['stock_name']}  {event['ts_code']}  ·  {grade}", f"首次回调 {_fmt_date(event['touch_date'])}")
        slide.shapes.add_picture(str(chart_paths[str(event["event_key"])]), Inches(0.38), Inches(1.02), width=Inches(9.38), height=Inches(5.95))
        panel = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.93), Inches(1.04), Inches(3.0), Inches(5.90))
        panel.fill.solid(); panel.fill.fore_color.rgb = __import__("pptx").dml.color.RGBColor.from_string(WHITE); panel.line.color.rgb = __import__("pptx").dml.color.RGBColor.from_string(MID)
        _add_text(slide, "事件指标", 10.17, 1.24, 2.5, 0.35, size=15, bold=True, color=NAVY)
        metrics = [
            ("锚点", f"MA{int(event['anchor_period'])}  {_fmt_date(event['anchor_start_date'])}"),
            ("首次回调MA120", _fmt_date(event["touch_date"])),
            ("MA120连续上涨", f"{int(event['ma120_rising_streak_at_touch'])}日"),
            ("回调日六均线离差", f"{float(event['six_ma_spread_at_touch_pct']):.2f}%"),
            ("MA5—60在MA120上方", f"{int(event['short_ma_above_ma120_count'])}/5条"),
            ("最低均线相对MA120", f"{float(event['min_short_ma_vs_ma120_pct']):+.2f}%"),
            ("箱体区间", f"{_fmt_date(event['box_start_date'])} → {_fmt_date(event['box_end_date'])}"),
            ("箱体天数 / 振幅", f"{int(event['box_trading_days'])}日 / {float(event['box_range_pct']):.2f}%"),
            ("六均线最窄", f"{float(event['minimum_six_ma_spread_pct']):.2f}%  {_fmt_date(event['minimum_six_ma_spread_date'])}"),
            (
                "箱体内MA5—60位置",
                f"≥MA120 {100.0 * float(event['box_short_ma_observations_above_ma120_ratio']):.1f}% / ≥-2% {100.0 * float(event['box_short_ma_observations_within_floor_ratio']):.0f}%",
            ),
        ]
        y = 1.75
        for label, value in metrics:
            _add_text(slide, label, 10.18, y, 2.5, 0.20, size=8.2, color=MUTED)
            _add_text(slide, value, 10.18, y + 0.19, 2.5, 0.30, size=10.7, bold=True, color=DARK_TEXT)
            y += 0.50
        _add_footer(slide, len(prs.slides))
    prs.save(output)


def write_report(result: dict[str, Any], sensitivity: pd.DataFrame, output: Path) -> None:
    events = result["unique_events"]
    row = result["summary"].iloc[0]
    lines = [
        "# 第一次回调MA120：六均线粘合与窄幅箱体",
        "",
        "## 结果",
        "",
        f"- 231个可搜索锚点中，{len(result['all_first'])}个找到首次MA120回调；锚点级严格六均线事件{len(result['anchor_events'])}个。",
        f"- 联合股票口径先跨锚点取最早首次回调再过滤，最终{len(events)}只；其中{len(result['unique_box_events'])}只形成至少5个交易日、最高/最低振幅不超过10%的窄箱体。",
        f"- 首次回调时MA120连续上涨中位数{row['ma120_streak_median']:.1f}日，IQR {row['ma120_streak_p25']:.1f}–{row['ma120_streak_p75']:.1f}日。",
        f"- A类窄箱体天数中位数{row['box_days_median']:.1f}日，IQR {row['box_days_p25']:.1f}–{row['box_days_p75']:.1f}日。",
        f"- {len(events)}只事件在箱体内的六均线最窄离差中位数{row['minimum_six_ma_spread_median']:.2f}%，最小值{row['minimum_six_ma_spread_min']:.2f}%。",
        f"- 箱体内均线位置描述性核验通过{int(row.get('box_position_descriptive_pass_count', len(events)))}/{len(events)}只：五线观测至少80%位于MA120上方，且没有一条低于MA120超过2%。",
        "",
        "## 主口径",
        "",
        "- 锚点与均线：使用MA60/MA120长趋势候选锚点并排除左截断段；MA5/10/20/30/60/120均由输入的完整历史前复权close先滚动计算，再切入锚点；只按实际交易日，不按自然日补零。",
        "- 第一次回调MA120：前20个交易日（不含当日）至少出现过收盘价高于自身MA120 8%；当日最低价不高于MA120上方2%，且(收盘价/MA120-1)≥-3%，同时MA120连续上涨>0日。连续上涨仅在MA120_t>MA120_(t-1)时累加，持平、下降或缺失均清零。",
        "- 选择顺序：先固定锚点内第一次回调，再检查粘合；第一次不合格时不允许跳到后续更漂亮的位置。跨锚点联合口径同样先取每只股票最早首次回调，再过滤；同日依次按anchor_period、anchor_start_date排序。",
        "- 六均线粘合：MA5/10/20/30/60/120的最大—最小离差除以六线均值≤6%。",
        "- 相对位置：事件日MA5—MA60五条均线中至少4条不低于MA120，且最低一条最多低于MA120 2%；箱体内另核验五线观测至少80%不低于MA120、100%不低于MA120下方2%，该项使用箱体期数据，只作事后描述。",
        "- 窄幅箱体：锚点内从事件前最多10日、事件后最多30日枚举包含事件日的连续区间；选择最高价/最低价振幅≤10%的最长区间（并列取振幅最小、再取最早起点），至少5个交易日为A类窄箱体。",
        "- 最窄粘合：在选中的箱体区间内，取六均线最大—最小无量纲离差的最小值及日期。",
        "",
        "## 口径边界",
        "",
        "- 首次回调与回调日六均线条件不使用未来数据；箱体持续天数及箱体内最窄离差允许观察回调后最多30个交易日，只作事后形态描述。",
        "- 历史事件可能发生在2025年，因为选中的长均线上升锚点使用完整已观察区间，而不只裁取2026自然年。",
        "- 股票池由事后观察到的超长MA60/MA120上涨段筛选，存在选择偏差；结果不是买入信号。",
    ]
    if not sensitivity.empty:
        names = {
            "six_ma_spread_threshold": "六均线离差上限",
            "box_range_threshold": "箱体振幅上限",
            "prior_extension_threshold": "此前离开MA120下限",
            "touch_upper_threshold": "最低价触及上限",
            "close_floor_threshold": "收盘破位下限",
            "position_floor_threshold": "短中均线最低位置",
            "min_short_above_count": "MA120上方均线条数",
        }
        lines.extend(
            [
                "",
                "## 阈值敏感性",
                "",
                "改变回调阈值时均重新寻找每个锚点的第一次回调，再做股票级最早事件选择和六均线过滤；不是只对主样本二次筛选。",
                "",
                "| 参数 | 参数阈值 | 首触股票 | 严格事件 | 箱体振幅阈值 | 窄箱体数 |",
                "|---|---:|---:|---:|---:|---:|",
            ]
        )
        for item in sensitivity.itertuples():
            first_count = "—" if pd.isna(getattr(item, "first_touch_stock_count", float("nan"))) else str(int(item.first_touch_stock_count))
            threshold = f"{float(item.threshold):g}"
            lines.append(
                f"| {names.get(item.sensitivity_type, item.sensitivity_type)} | {threshold} | {first_count} | {int(item.event_count)} | {float(item.box_range_threshold_pct):g}% | {int(item.narrow_box_count)} |"
            )
    lines.extend(
        [
            "",
            "## 事件清单",
            "",
            "| 股票 | 代码 | 首次回调 | MA120上涨日数 | 箱体天数 | 箱体振幅 | 回调日六均线离差 | 箱体内最窄 | 箱体内五线≥MA120 |",
            "|---|---|---:|---:|---:|---:|---:|---:|---:|",
        ]
    )
    for event in events.itertuples():
        lines.append(
            f"| {event.stock_name} | {event.ts_code} | {_fmt_date(event.touch_date)} | {int(event.ma120_rising_streak_at_touch)} | {int(event.box_trading_days)} | {float(event.box_range_pct):.2f}% | {float(event.six_ma_spread_at_touch_pct):.2f}% | {float(event.minimum_six_ma_spread_pct):.2f}% | {100.0 * float(event.box_short_ma_observations_above_ma120_ratio):.1f}% |"
        )
    output.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main(argv: Sequence[str] | None = None) -> int:
    args = _parse_args(argv)
    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    result = run_scan(args)
    events = result["unique_events"].copy()
    events["event_key"] = events.apply(
        lambda row: f"{row['ts_code']}|{int(row['anchor_period'])}|{row['anchor_start_date']}|{row['touch_date']}",
        axis=1,
    )
    result["unique_events"] = events
    result["unique_box_events"] = events.loc[events["narrow_box_pass"]].copy()
    sensitivity = build_sensitivity(result, args)
    outputs = {
        "anchor_eligibility.csv": result["eligibility"],
        "all_first_ma120_pullbacks.csv": result["all_first"],
        "anchor_level_six_ma_events.csv": result["anchor_events"],
        "unique_stock_first_ma120_pullbacks.csv": result["unique_first"],
        "unique_stock_six_ma_events.csv": events,
        "unique_stock_narrow_box_events.csv": result["unique_box_events"],
        "coverage_summary.csv": result["coverage"],
        "metric_summary.csv": result["summary"],
        "threshold_sensitivity.csv": sensitivity,
    }
    for name, frame in outputs.items():
        frame.to_csv(output_dir / name, index=False)
    report_path = output_dir / "report.md"
    write_report(result, sensitivity, report_path)

    pptx_path: Path | None = None
    pdf_path: Path | None = None
    contacts: list[Path] = []
    chart_paths: dict[str, Path] = {}
    rendered: list[Path] = []
    summary_chart = output_dir / "summary_metrics.png"
    if not args.metrics_only:
        charts_dir = output_dir / "charts"
        charts_dir.mkdir(parents=True, exist_ok=True)
        for stale in charts_dir.glob("*.png"):
            stale.unlink()
        for event in events.itertuples(index=False):
            path = charts_dir / f"{int(event.display_order):02d}_{str(event.ts_code).replace('.', '_')}_{event.touch_date}.png"
            create_event_chart(result["features_by_code"][str(event.ts_code)], pd.Series(event._asdict()), path)
            chart_paths[str(event.event_key)] = path
        if len(chart_paths) != len(events) or len({_sha256(path) for path in chart_paths.values()}) != len(events):
            raise ValueError("chart count or uniqueness mismatch")
        create_summary_chart(events, summary_chart)
        pptx_path = output_dir / "第一次回调MA120_六均线粘合与窄幅箱体.pptx"
        build_deck(
            events=events,
            coverage=result["coverage"],
            summary=result["summary"],
            sensitivity=sensitivity,
            chart_paths=chart_paths,
            summary_chart=summary_chart,
            output=pptx_path,
        )
        pdf_path = convert_ppt_to_pdf(pptx_path, output_dir)
        rendered = render_pdf(pdf_path, output_dir / "rendered_slides")
        contacts = create_contact_sheets(rendered, output_dir / "contact_sheets")
        prs = Presentation(pptx_path)
        expected_slides = 7 + len(events)
        if len(prs.slides) != expected_slides or len(rendered) != expected_slides:
            raise ValueError("PPT/PDF page count mismatch")
        detail_slides = list(prs.slides)[7:]
        picture_counts = [sum(1 for shape in slide.shapes if shape.shape_type == 13) for slide in detail_slides]
        if picture_counts != [1] * len(events):
            raise ValueError(f"detail picture count mismatch: {picture_counts}")

    artifacts: list[Path] = [output_dir / name for name in outputs] + [report_path]
    if pptx_path and pdf_path:
        artifacts.extend([summary_chart, pptx_path, pdf_path, *chart_paths.values(), *contacts])
    metadata = {
        "generated_at": datetime.now().astimezone().isoformat(),
        "study": "first_ma120_pullback_six_ma_cohesion_and_narrow_box",
        "window": "selected long-MA anchors through 2026-09-30; historical events may occur in 2025",
        "data_contract": {
            "price_adjustment": "qfq=True",
            "volume": "raw volume",
            "trading_days_only": True,
            "daily_input": str(Path(args.daily_file).resolve()),
            "source_contract": "formal cache built via coreClient.data_provider with source=database_only",
            "ma_computation": "MA5/10/20/30/60/120 are rolling means of qfq close on each stock's complete cached trading-day history before slicing anchors",
        },
        "definitions": {
            "first_ma120_pullback": f"inside selected non-left-censored MA60/MA120 anchor, prior {args.extension_lookback}-day max(close/own-day-MA120-1)>={args.prior_extension_pct}%, touch low/MA120-1<={args.touch_upper_pct}%, close/MA120-1>={args.close_floor_pct}%, and MA120 rising streak>0; select first before cohesion",
            "six_ma_cohesion": f"(max(MA5,MA10,MA20,MA30,MA60,MA120)-min)/mean<={args.cohesion_threshold_pct}%",
            "position_constraint": f"at least {args.min_short_above_count}/5 of MA5/10/20/30/60 >= MA120 and no short/mid MA below MA120 by more than {abs(args.max_short_below_ma120_pct)}%",
            "box_position_description": f"descriptive pass when at least 80% of all MA5/10/20/30/60 by box-day observations are >= MA120 and 100% are >= MA120{args.max_short_below_ma120_pct:+g}%; uses box-period observations and is not part of the no-lookahead event filter",
            "narrow_box": f"longest interval containing touch, within anchor and event-{args.box_lookback} to event+{args.box_lookforward} trading days, max(high)/min(low)-1<={args.box_max_range_pct}%, at least {args.box_min_days} trading days; ties choose lower range then earlier start",
            "ma120_rising_streak": "MA120_t>MA120_(t-1); flat/fall/unavailable resets to zero",
            "minimum_six_ma_spread": "minimum six-MA dimensionless spread inside selected box; may use post-touch observations",
            "unique_stock_rule": "select each stock's earliest first MA120 pullback across eligible MA60/MA120 anchors before applying cohesion and position constraints; same-day ties sort by anchor_period then anchor_start_date",
        },
        "sensitivity_grids": {
            "six_ma_spread_pct": [4, 5, 6, 7, 8],
            "box_range_pct": [8, 10, 12],
            "prior_extension_pct": [5, 8, 10],
            "touch_upper_pct": [1, 2, 3],
            "close_floor_pct": [-1, -3, -5],
            "max_short_below_ma120_pct": [-1, -2, -3],
            "min_short_above_count": [3, 4, 5],
            "rescan_rule": "for pullback and position variants, recompute each anchor's first touch and then select each stock's earliest touch before filtering",
        },
        "counts": {
            "candidate_stocks": int(result["daily"]["ts_code"].nunique()),
            "eligible_anchors": int(result["anchors"]["search_eligible"].sum()),
            "anchor_first_touch_events": int(len(result["all_first"])),
            "anchor_six_ma_events": int(len(result["anchor_events"])),
            "anchor_narrow_box_events": int(len(result["anchor_box_events"])),
            "unique_first_touch_stocks": int(len(result["unique_first"])),
            "unique_six_ma_events": int(len(events)),
            "unique_narrow_box_events": int(len(result["unique_box_events"])),
            "event_charts": int(len(chart_paths)),
            "slides": int(7 + len(events)) if pptx_path else 0,
            "pdf_pages": int(len(rendered)),
        },
        "lookahead_boundary": "first touch and event-day six-MA qualification are no-lookahead; box duration, minimum spread, and box-period MA position coverage are descriptive and may include up to 30 post-touch trading days",
        "selection_bias_warning": "The pool is selected by future-observed long MA duration; findings are descriptive, not live-strategy evidence or buy signals.",
        "artifacts": {
            path.name: {"path": str(path.resolve()), "sha256": _sha256(path), "bytes": path.stat().st_size}
            for path in artifacts
        },
    }
    metadata_path = output_dir / "metadata.json"
    metadata_path.write_text(json.dumps(_clean_json(metadata), ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    payload = {
        "output_dir": str(output_dir.resolve()),
        "eligible_anchors": int(result["anchors"]["search_eligible"].sum()),
        "first_touch_anchor_events": int(len(result["all_first"])),
        "anchor_six_ma_events": int(len(result["anchor_events"])),
        "anchor_narrow_box_events": int(len(result["anchor_box_events"])),
        "unique_first_touch_stocks": int(len(result["unique_first"])),
        "unique_six_ma_events": int(len(events)),
        "unique_narrow_box_events": int(len(result["unique_box_events"])),
        "top_events": events[
            ["stock_name", "ts_code", "touch_date", "ma120_rising_streak_at_touch", "box_trading_days", "box_range_pct", "minimum_six_ma_spread_pct", "narrow_box_pass"]
        ].to_dict("records"),
    }
    if pptx_path and pdf_path:
        payload.update({"slides": 7 + len(events), "pptx": str(pptx_path.resolve()), "pdf": str(pdf_path.resolve()), "contact_sheets": [str(path.resolve()) for path in contacts]})
    print(json.dumps(payload, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

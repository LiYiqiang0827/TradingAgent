from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_ma60_ma120_top100_kline_ppt import (
    _parse_args,
    build_combined_presentation,
    display_start_for_spans,
    episode_spans_from_row,
    materialize_reviewed_classifications,
    price_basis_label,
    render_dual_period_chart,
)


def test_parse_args_allows_explicit_qfq_charts() -> None:
    assert _parse_args([]).qfq is False
    assert _parse_args(["--qfq"]).qfq is True


def test_price_basis_label_distinguishes_qfq_from_unadjusted() -> None:
    assert price_basis_label(True) == "前复权"
    assert price_basis_label(False) == "未复权"


def test_display_start_for_spans_keeps_context_and_analysis_floor() -> None:
    assert display_start_for_spans(
        [{"period": 120, "start_date": "20220101", "end_date": "20260930", "target_days": 181}]
    ) == pd.Timestamp("2023-01-01")
    assert display_start_for_spans(
        [{"period": 60, "start_date": "20260601", "end_date": "20260930", "target_days": 80}]
    ) == pd.Timestamp("2025-10-01")
    assert display_start_for_spans(
        [{"period": 60, "start_date": "20251223.0", "end_date": "20260930.0", "target_days": 80}]
    ) == pd.Timestamp("2025-09-24")


def test_episode_spans_from_row_returns_only_present_periods() -> None:
    row = pd.Series(
        {
            "in_ma60_top100": True,
            "ma60_start_date": "20260101",
            "ma60_end_date": "20260201",
            "ma60_target_days": 20,
            "in_ma120_top100": False,
            "ma120_start_date": pd.NA,
            "ma120_end_date": pd.NA,
            "ma120_target_days": pd.NA,
        }
    )

    assert episode_spans_from_row(row) == [
        {
            "period": 60,
            "start_date": "20260101",
            "end_date": "20260201",
            "target_days": 20,
        }
    ]


def test_materialize_reviewed_classifications_requires_exact_coverage() -> None:
    evidence = [
        {
            "ts_code": "A",
            "kpl_list_themes": [{"theme": "芯片", "count": 2}],
            "concepts": [],
            "limit_performance_themes": [],
            "industry": "半导体",
        },
        {
            "ts_code": "B",
            "kpl_list_themes": [],
            "concepts": [],
            "limit_performance_themes": [],
            "industry": "机械基件",
        },
    ]
    reviewed = [
        {"ts_code": "A", "primary_theme": "芯片", "broad_category": "半导体/芯片", "secondary_themes": [], "confidence": "high", "rationale": "KPL证据"},
        {"ts_code": "B", "primary_theme": "机械基件（行业兜底）", "broad_category": "机器人/高端装备", "secondary_themes": [], "confidence": "low", "rationale": "行业兜底"},
    ]

    result = materialize_reviewed_classifications(evidence, reviewed)

    assert [item["ts_code"] for item in result] == ["A", "B"]
    assert result[0]["evidence_labels"] == ["芯片"]
    assert result[1]["evidence_note"] == "行业兜底"
    with pytest.raises(ValueError, match="coverage"):
        materialize_reviewed_classifications(evidence, reviewed[:1])


def test_render_dual_period_chart_draws_price_and_volume(tmp_path: Path) -> None:
    dates = pd.bdate_range("2025-09-01", periods=190)
    close = pd.Series(range(100, 290), dtype=float)
    frame = pd.DataFrame(
        {
            "trade_date": dates.strftime("%Y%m%d"),
            "open": close - 0.4,
            "high": close + 1.0,
            "low": close - 1.0,
            "close": close,
            "vol": 10000 + close * 20,
            "ma20": close.rolling(20).mean(),
            "ma30": close.rolling(30).mean(),
            "ma60": close.rolling(60).mean(),
            "ma120": close.rolling(120).mean(),
        }
    )
    output = tmp_path / "dual.png"

    receipt = render_dual_period_chart(
        frame,
        output,
        stock_name="测试股份",
        ts_code="000001.SZ",
        episode_spans=[
            {"period": 60, "start_date": dates[80].strftime("%Y%m%d"), "end_date": dates[160].strftime("%Y%m%d"), "target_days": 81},
            {"period": 120, "start_date": dates[30].strftime("%Y%m%d"), "end_date": dates[-1].strftime("%Y%m%d"), "target_days": 160},
        ],
        display_start=dates[0],
        display_end=dates[-1],
        price_basis="前复权",
    )

    with Image.open(output) as image:
        assert image.width >= 1800
        assert image.height >= 900
    assert receipt["bars"] == 190
    assert receipt["periods"] == [60, 120]
    assert receipt["price_basis"] == "前复权"


def test_build_combined_presentation_adds_one_stock_slide_per_unique_code(tmp_path: Path) -> None:
    image_a = tmp_path / "a.png"
    image_b = tmp_path / "b.png"
    Image.new("RGB", (1800, 900), "white").save(image_a)
    Image.new("RGB", (1800, 900), "#f5f5f5").save(image_b)
    stocks = pd.DataFrame(
        [
            {
                "observation_rank": 1, "ts_code": "A", "stock_name": "甲公司", "industry": "半导体",
                "membership": "MA60+MA120", "in_ma60_top100": True, "in_ma120_top100": True,
                "ma60_rank": 1, "ma60_start_date": "20260101", "ma60_end_date": "20260930", "ma60_target_days": 181,
                "ma60_carried_in": False, "ma60_right_censored": True,
                "ma120_rank": 2, "ma120_start_date": "20250101", "ma120_end_date": "20260930", "ma120_target_days": 181,
                "ma120_carried_in": True, "ma120_right_censored": True,
                "primary_theme": "半导体设备", "broad_category": "半导体/芯片", "secondary_themes": ["先进封装"],
                "evidence_labels": ["芯片"], "confidence": "high", "evidence_note": "本地证据",
            },
            {
                "observation_rank": 2, "ts_code": "B", "stock_name": "乙公司", "industry": "银行",
                "membership": "MA120", "in_ma60_top100": False, "in_ma120_top100": True,
                "ma60_rank": pd.NA, "ma60_start_date": pd.NA, "ma60_end_date": pd.NA, "ma60_target_days": pd.NA,
                "ma60_carried_in": pd.NA, "ma60_right_censored": pd.NA,
                "ma120_rank": 1, "ma120_start_date": "20240101", "ma120_end_date": "20260930", "ma120_target_days": 181,
                "ma120_carried_in": True, "ma120_right_censored": True,
                "primary_theme": "银行", "broad_category": "金融", "secondary_themes": ["高股息"],
                "evidence_labels": ["银行"], "confidence": "medium", "evidence_note": "本地证据",
            },
        ]
    )
    output = tmp_path / "combined.pptx"

    receipt = build_combined_presentation(
        stocks,
        {"A": image_a, "B": image_b},
        output,
        data_end="20260930",
        ma60_cutoff=128,
        ma120_cutoff=181,
        price_basis="前复权",
    )

    reopened = Presentation(output)
    assert receipt == {
        "slides": 11,
        "stock_slides": 2,
        "category_slides": 2,
        "inventory_slides": 1,
        "ranking_slides": 2,
    }
    assert len(reopened.slides) == 11
    all_text = "\n".join(
        shape.text
        for slide in reopened.slides
        for shape in slide.shapes
        if hasattr(shape, "text_frame")
    )
    assert "未复权" not in all_text
    assert "排名与展示均使用前复权" in all_text

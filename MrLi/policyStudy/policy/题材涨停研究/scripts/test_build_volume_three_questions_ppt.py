from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from PIL import Image
from pptx import Presentation

sys.path.insert(0, str(Path(__file__).resolve().parent))

from build_volume_three_questions_ppt import (
    build_presentation,
    derive_key_findings,
    extract_equal_length_volume_ratios,
    humanize_divergence_spec,
    select_cooldown_sensitivity,
    write_markdown_report,
)


def _synthetic_inputs() -> dict[str, pd.DataFrame]:
    long_summary = pd.DataFrame(
        [
            {
                "anchor_period": 60,
                "cohort": "all_observed",
                "baseline": "equal_length",
                "comparable_count": 100,
                "ratio_mean": 1.9,
                "ratio_median": 1.75,
                "ratio_p25": 1.3,
                "ratio_p75": 2.1,
                "amplified_gt_1_count": 93,
                "amplified_gt_1_pct": 93.0,
                "amplified_ge_1_2_count": 84,
                "amplified_ge_1_2_pct": 84.0,
                "amplified_ge_1_5_count": 64,
                "amplified_ge_1_5_pct": 64.0,
                "amplified_ge_2_count": 31,
                "amplified_ge_2_pct": 31.0,
            },
            {
                "anchor_period": 60,
                "cohort": "all_observed",
                "baseline": "pre20",
                "comparable_count": 100,
                "ratio_mean": 2.5,
                "ratio_median": 2.0,
                "ratio_p25": 1.5,
                "ratio_p75": 3.2,
                "amplified_gt_1_count": 94,
                "amplified_gt_1_pct": 94.0,
                "amplified_ge_1_2_count": 89,
                "amplified_ge_1_2_pct": 89.0,
                "amplified_ge_1_5_count": 76,
                "amplified_ge_1_5_pct": 76.0,
                "amplified_ge_2_count": 51,
                "amplified_ge_2_pct": 51.0,
            },
            {
                "anchor_period": 120,
                "cohort": "all_observed",
                "baseline": "equal_length",
                "comparable_count": 120,
                "ratio_mean": 2.3,
                "ratio_median": 1.89,
                "ratio_p25": 1.5,
                "ratio_p75": 3.0,
                "amplified_gt_1_count": 112,
                "amplified_gt_1_pct": 93.3,
                "amplified_ge_1_2_count": 103,
                "amplified_ge_1_2_pct": 85.8,
                "amplified_ge_1_5_count": 90,
                "amplified_ge_1_5_pct": 75.0,
                "amplified_ge_2_count": 54,
                "amplified_ge_2_pct": 45.0,
            },
            {
                "anchor_period": 120,
                "cohort": "all_observed",
                "baseline": "pre20",
                "comparable_count": 124,
                "ratio_mean": 2.5,
                "ratio_median": 2.19,
                "ratio_p25": 1.34,
                "ratio_p75": 3.16,
                "amplified_gt_1_count": 115,
                "amplified_gt_1_pct": 92.7,
                "amplified_ge_1_2_count": 101,
                "amplified_ge_1_2_pct": 81.5,
                "amplified_ge_1_5_count": 88,
                "amplified_ge_1_5_pct": 71.0,
                "amplified_ge_2_count": 67,
                "amplified_ge_2_pct": 54.0,
            },
        ]
    )
    segment_rows = []
    anchor_rows = []
    for period, segment_pct, all_pct, median_pct in (
        (60, 69.0, 13.1, 70.0),
        (120, 65.1, 0.0, 66.1),
    ):
        for days in (1, 5, 10):
            segment_rows.append(
                {
                    "anchor_period": period,
                    "comparison": "full_episode_equal_prior",
                    "min_segment_days": days,
                    "comparable_count": 700,
                    "ratio_median": 1.2,
                    "amplified_gt_1_count": 480,
                    "amplified_gt_1_pct": segment_pct + days,
                }
            )
            segment_rows.append(
                {
                    "anchor_period": period,
                    "comparison": "clipped_vs_anchor_mean",
                    "min_segment_days": days,
                    "comparable_count": 700,
                    "ratio_median": 0.98,
                    "amplified_gt_1_count": 320,
                    "amplified_gt_1_pct": 45.0 + days,
                }
            )
            anchor_rows.append(
                {
                    "anchor_period": period,
                    "comparison": "full_episode_equal_prior",
                    "min_episode_days": days,
                    "anchor_rows": 100,
                    "evaluable_anchor_count": 100,
                    "all_segments_amplified_gt_1_count": int(all_pct),
                    "all_segments_amplified_gt_1_pct": all_pct + (days - 1),
                    "median_anchor_amplified_segment_pct": median_pct + (days - 1),
                }
            )
    effectiveness = pd.DataFrame(
        [
            {
                "anchor_period": period,
                "scope": scope,
                "event_class": event_class,
                "event_count": 100,
                "forward_comparable_count": 100,
                "median_forward_return_pct": 5.0,
                "median_forward_max_drawdown_pct": -7.0,
                "drawdown_ge_10pct_rate": rate,
                "top_like_forward_rate": rate / 2,
                "near_anchor_peak_10d_rate": near,
            }
            for period, scope, event_class, rate, near in [
                (60, "all_candidate_days", "non_divergent_new_high", 42.0, 22.6),
                (60, "all_candidate_days", "divergence", 38.3, 18.4),
                (60, "cooldown_10d", "non_divergent_new_high", 37.3, 20.3),
                (60, "cooldown_10d", "divergence", 39.1, 18.9),
                (120, "all_candidate_days", "non_divergent_new_high", 41.4, float("nan")),
                (120, "all_candidate_days", "divergence", 34.6, float("nan")),
                (120, "cooldown_10d", "non_divergent_new_high", 38.1, float("nan")),
                (120, "cooldown_10d", "divergence", 34.0, float("nan")),
            ]
        ]
    )
    lifts = pd.DataFrame(
        [
            {
                "anchor_period": period,
                "scope": scope,
                "drawdown_ge_10pct_rate_lift": lift,
                "top_like_forward_rate_lift": 0.9,
                "near_anchor_peak_10d_rate_lift": near_lift,
            }
            for period, scope, lift, near_lift in [
                (60, "all_candidate_days", 0.91, 0.81),
                (60, "cooldown_10d", 1.05, 0.93),
                (120, "all_candidate_days", 0.83, float("nan")),
                (120, "cooldown_10d", 0.89, float("nan")),
            ]
        ]
    )
    coverage = pd.DataFrame(
        [
            {
                "anchor_period": 60,
                "scope": "all_candidate_days",
                "completed_anchor_count": 101,
                "anchors_with_any_divergence": 87,
                "anchors_with_any_divergence_pct": 86.1,
                "anchors_with_divergence_within_10d_of_top": 37,
                "top_coverage_rate_pct": 36.6,
                "near_top_signal_lead_median_days": 4.0,
            }
        ]
    )
    within = pd.DataFrame(
        [
            {
                "anchor_period": 60,
                "scope": "all_candidate_days",
                "paired_anchor_count": 90,
                "divergence_rate_higher_count": 41,
                "divergence_rate_lower_count": 47,
                "median_drawdown_rate_difference_pct": -6.2,
                "median_forward_return_difference_pct": 4.3,
            },
            {
                "anchor_period": 120,
                "scope": "all_candidate_days",
                "paired_anchor_count": 123,
                "divergence_rate_higher_count": 39,
                "divergence_rate_lower_count": 81,
                "median_drawdown_rate_difference_pct": -10.7,
                "median_forward_return_difference_pct": 2.8,
            },
        ]
    )
    return {
        "long_summary": long_summary,
        "segment_summary": pd.DataFrame(segment_rows),
        "anchor_summary": pd.DataFrame(anchor_rows),
        "effectiveness": effectiveness,
        "lifts": lifts,
        "coverage": coverage,
        "within": within,
    }


def test_extract_equal_length_ratios_uses_the_production_detail_column() -> None:
    detail = pd.DataFrame(
        {
            "anchor_period": [60, 60, 120],
            "anchor_left_censored": [False, True, False],
            "equal_prior_volume_ratio": [1.5, 9.9, 2.0],
        }
    )

    ratios = extract_equal_length_volume_ratios(detail, 60)

    assert ratios.tolist() == [1.5]


def test_select_cooldown_sensitivity_accepts_aggregate_production_schema() -> None:
    sensitivity = pd.DataFrame(
        {
            "spec_name": ["a", "a", "b"],
            "anchor_period": [60, 60, 120],
            "scope": ["all_candidate_days", "cooldown_10d", "cooldown_10d"],
            "drawdown_rate_lift": [0.8, 0.9, 1.0],
        }
    )

    selected = select_cooldown_sensitivity(sensitivity)

    assert selected["spec_name"].tolist() == ["a", "b"]
    assert selected["drawdown_rate_lift"].tolist() == [0.9, 1.0]


def test_humanize_divergence_spec_replaces_internal_codes() -> None:
    assert humanize_divergence_spec("close_l60_v5_t0.7") == "缩量≥30%"
    assert humanize_divergence_spec("close_l60_v5_t0.8") == "缩量≥20%"
    assert humanize_divergence_spec("close_l60_v5_t0.9") == "缩量≥10%"
    assert humanize_divergence_spec("close_l40_v5_t0.8") == "回看40日"
    assert humanize_divergence_spec("close_l120_v5_t0.8") == "回看120日"
    assert humanize_divergence_spec("close_l60_v10_t0.8") == "10日均量"
    assert humanize_divergence_spec("high_l60_v5_t0.8") == "最高价新高"


def test_derive_key_findings_keeps_the_three_questions_separate() -> None:
    findings = derive_key_findings(**_synthetic_inputs())

    assert findings["q1"][60]["equal_length_median_ratio"] == 1.75
    assert findings["q1"][120]["amplified_gt_1_pct"] == 93.3
    assert findings["q2"][60][1]["all_segments_amplified_pct"] == 13.1
    assert findings["q2"][120][1]["segment_amplified_pct"] == 66.1
    assert findings["q3"][60]["all_candidate_days"]["drawdown_rate_lift"] == 0.91
    assert findings["q3"][120]["all_candidate_days"]["paired_lower_count"] == 81
    assert findings["q3_top_coverage"]["top_coverage_rate_pct"] == 36.6


def test_build_presentation_and_markdown_are_complete(tmp_path: Path) -> None:
    findings = derive_key_findings(**_synthetic_inputs())
    image = tmp_path / "chart.png"
    Image.new("RGB", (1200, 700), "white").save(image)
    chart_paths = {
        key: image
        for key in (
            "q1_thresholds",
            "q1_distribution",
            "q2_segments",
            "q2_all_segments",
            "q2_benchmarks",
            "q3_forward",
            "q3_paired",
            "q3_sensitivity",
            "q3_examples",
        )
    }
    pptx_path = tmp_path / "report.pptx"
    markdown_path = tmp_path / "report.md"

    stats = build_presentation(findings, chart_paths, pptx_path)
    write_markdown_report(findings, markdown_path)

    deck = Presentation(pptx_path)
    assert stats["slide_count"] == 15
    assert len(deck.slides) == 15
    assert "三个成交量问题" in "\n".join(
        shape.text for shape in deck.slides[0].shapes if hasattr(shape, "text")
    )
    markdown = markdown_path.read_text(encoding="utf-8")
    assert "问题一" in markdown
    assert "问题二" in markdown
    assert "问题三" in markdown
    appendix_text = "\n".join(
        shape.text for shape in deck.slides[14].shapes if hasattr(shape, "text")
    )
    design_text = "\n".join(
        shape.text for shape in deck.slides[2].shapes if hasattr(shape, "text")
    )
    assert "带入锚点按完整已观察区间分析" in design_text
    assert "带入锚点" in markdown
    assert "明细、汇总和元数据保存在对应研究目录" in appendix_text
    assert "同一正式输出目录" not in appendix_text
